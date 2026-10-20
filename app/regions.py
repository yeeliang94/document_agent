"""Sparse, resumable XLSX indexing and source-backed region accounting (schema 2).

The index suggests boundaries. Only the owning agent establishes meaning and relevance.
"""

import hashlib
import posixpath
import time
import zipfile
from collections import defaultdict
from typing import Literal
from xml.etree import ElementTree as ET

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
from pydantic import BaseModel, Field

from .files import atomic_bytes, within
from .state import check_page, identity, page_result

INDEX_VERSION = 1
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


def accounting_version(state):
    return state.get('schema', 'accounting', {}).get('version', 1)


def workbook_parts(archive):
    rels = {r.attrib['Id']: r.attrib['Target'] for r in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
    sheets = ET.fromstring(archive.read('xl/workbook.xml')).find(NS + 'sheets')
    return [
        (
            s.attrib['name'],
            s.attrib.get('state', 'visible'),
            posixpath.normpath('xl/' + rels[s.attrib[REL + 'id']].lstrip('/').removeprefix('xl/')),
        )
        for s in sheets
    ]


def declared_tables(archive, sheet_part):
    folder, name = posixpath.split(sheet_part)
    relpart = folder + '/_rels/' + name + '.rels'
    if relpart not in archive.namelist():
        return []
    tables = []
    for relation in ET.fromstring(archive.read(relpart)):
        if relation.attrib.get('Type', '').endswith('/table'):
            target = posixpath.normpath(posixpath.join(folder, relation.attrib['Target'])).lstrip('/')
            table = ET.fromstring(archive.read(target))
            tables.append(
                {
                    'name': table.attrib.get('displayName'),
                    'range': table.attrib['ref'],
                    'totals_rows': table.attrib.get('totalsRowCount', '0'),
                }
            )
    return tables


def sheet_notes(archive, sheet_part):
    folder, name = posixpath.split(sheet_part)
    relpart = folder + '/_rels/' + name + '.rels'
    if relpart not in archive.namelist():
        return {}
    notes = {}
    for relation in ET.fromstring(archive.read(relpart)):
        if relation.attrib.get('Type', '').endswith('/comments'):
            target = posixpath.normpath(posixpath.join(folder, relation.attrib['Target'])).lstrip('/')
            for comment in ET.fromstring(archive.read(target)).iter(NS + 'comment'):
                text = comment.find(NS + 'text')
                notes[comment.attrib['ref']] = ''.join(text.itertext())[:2000] if text is not None else ''
                if len(notes) > 10000:
                    raise NotImplementedError(
                        'More than 10000 worksheet comments require a narrower source; discovery is incomplete'
                    )
    return notes


def sheet_layout(source, sheet):
    """Notes and merged ranges for one sheet, streamed so a bounded read never loads the whole workbook."""
    with zipfile.ZipFile(source) as archive:
        part = next((part for title, _, part in workbook_parts(archive) if title == sheet), None)
        if part is None:
            return {}, []
        try:
            notes = sheet_notes(archive, part)
        except NotImplementedError:
            notes = {}
        merges = []
        with archive.open(part) as stream:
            for _, element in ET.iterparse(stream, events=('end',)):
                if element.tag == NS + 'mergeCell':
                    merges.append(element.attrib['ref'])
                elif element.tag == NS + 'row':
                    element.clear()
    return notes, merges


def shared_strings(archive):
    if 'xl/sharedStrings.xml' not in archive.namelist():
        return []
    strings = []
    with archive.open('xl/sharedStrings.xml') as stream:
        for _, element in ET.iterparse(stream, events=('end',)):
            if element.tag == NS + 'si':
                strings.append(''.join(element.itertext()))
                element.clear()
    return strings


def indexed_cells(state, source_id):
    return state.get('cell_index', source_id, [])


def propose_regions(state, source_id, cells, entry, path, sheet):
    by_row = defaultdict(list)
    for cell in cells:
        by_row[cell['row']].append(cell['column'])
    groups, active = [], []
    for row, columns in sorted(by_row.items()):
        runs = []
        for column in sorted(set(columns)):
            if runs and column <= runs[-1][1] + 1:
                runs[-1][1] = column
            else:
                runs.append([column, column])
        current = []
        for left, right in runs:
            matches = [g for g in active if g[3] >= row - 2 and left <= g[2] and right >= g[0]]
            if matches:
                group = matches[0]
                group[0], group[2], group[3] = min(group[0], left), max(group[2], right), row
            else:
                group = [left, row, right, row]
                groups.append(group)
            current.append(group)
        active = current
    for left, top, right, bottom in groups:
        bounds = f'{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}'
        key = identity(source_id, bounds, INDEX_VERSION)
        samples = [c for c in cells if left <= c['column'] <= right and top <= c['row'] <= min(bottom, top + 2)][:12]
        state.put(
            'region',
            key,
            {
                'source': source_id,
                'path': path,
                'sheet': sheet,
                'bounds': bounds,
                'source_hash': entry['sha256'],
                'index_version': INDEX_VERSION,
                'proposed_type': 'table' if bottom > top else 'unclassified_band',
                'header_rows': [top],
                'samples': samples,
                'uncertainty': 'Candidate only; verify headers, totals, blank gaps and adjacent regions.',
            },
        )


def discover(state, path, sheet=None, start_row=1, max_cells=1_000_000, seconds=30, stop=lambda: None):
    file_id, entry = state.file_entry(path)
    if not entry:
        raise ValueError('Source must appear in inventory')
    source = within(state.root, path)
    if hashlib.sha256(source.read_bytes()).hexdigest() != entry['sha256']:
        raise ValueError('Source changed since intake')
    if start_row < 1:
        raise ValueError('start_row must be positive')
    summaries = []
    started = time.monotonic()
    with zipfile.ZipFile(source) as archive:
        strings = None
        found = False
        for title, visibility, part in workbook_parts(archive):
            key = identity(file_id, title)
            old = state.get('inventory', key, {})
            info = old | {
                'kind': 'sheet',
                'parent': file_id,
                'path': path,
                'sheet': title,
                'visibility': visibility,
                'status': 'inspected',
                'index_version': INDEX_VERSION,
                'tables': declared_tables(archive, part),
            }
            if sheet and title != sheet:
                if not old:
                    state.put('inventory', key, info | {'status': 'discovered', 'scan_complete': False, 'next_row': 1})
                continue
            found = True
            if old.get('scan_complete') and old.get('index_version') == INDEX_VERSION:
                summaries.append({'id': key, **old})
                continue
            cursor = old.get('next_row', 1)
            if start_row > cursor:
                raise ValueError(f'Continue {title} at row {cursor}; skipping discovery is not allowed')
            # A repeated first page continues from the saved cursor rather than incrementing counters.
            cursor = max(cursor, start_row)
            cached = state.get('index_cache', identity(entry['sha256'], title, INDEX_VERSION))
            if cached:
                cells = cached['cells']
                info = info | cached['metadata']
                state.put('cell_index', key, cells)
                state.put('inventory', key, info)
                propose_regions(state, key, cells, entry, path, title)
                inventory_images(state, archive, part, key, path, title, entry['sha256'])
                summaries.append({'id': key, **info})
                continue
            if strings is None:
                strings = shared_strings(archive)
            cells = indexed_cells(state, key)
            notes = sheet_notes(archive, part)
            visited, complete, next_row = 0, True, None
            merges, validations, drawings = [], [], []
            with archive.open(part) as stream:
                for _, element in ET.iterparse(stream, events=('end',)):
                    if element.tag == NS + 'row':
                        number = int(element.attrib['r'])
                        if number < cursor:
                            element.clear()
                            continue
                        stop()
                        row = []
                        for cell in element.findall(NS + 'c'):
                            value = cell.find(NS + 'v')
                            inline = cell.find(NS + 'is')
                            formula = cell.find(NS + 'f')
                            if (
                                value is None
                                and inline is None
                                and formula is None
                                and cell.attrib['r'] not in notes
                                and not (entry.get('role') == 'template' and cell.attrib.get('s'))
                            ):
                                continue
                            raw = (
                                value.text
                                if value is not None
                                else ''.join(inline.itertext())
                                if inline is not None
                                else None
                            )
                            if cell.attrib.get('t') == 's' and raw is not None:
                                raw = strings[int(raw)]
                            r, c = coordinate_to_tuple(cell.attrib['r'])
                            row.append(
                                {
                                    'cell': cell.attrib['r'],
                                    'row': r,
                                    'column': c,
                                    'value': raw,
                                    'formula': formula.text if formula is not None else None,
                                    'type': cell.attrib.get('t', 'n'),
                                    'style': cell.attrib.get('s'),
                                    'note': notes.get(cell.attrib['r']),
                                }
                            )
                        if visited and (visited + len(row) > max_cells or time.monotonic() - started >= seconds):
                            complete, next_row = False, number
                            break
                        if len(row) > max_cells:
                            raise NotImplementedError(
                                'A single row exceeds the cell work-unit limit; narrow or convert the source'
                            )
                        cells.extend(row)
                        visited += len(row)
                        element.clear()
                    elif element.tag == NS + 'mergeCell':
                        merges.append(element.attrib['ref'])
                    elif element.tag == NS + 'dataValidation':
                        validations.append(
                            {'range': element.attrib['sqref'], 'formula': element.findtext(NS + 'formula1')}
                        )
                    elif element.tag == NS + 'drawing':
                        drawings.append(element.attrib[REL + 'id'])
            if complete:
                indexed = {cell['cell'] for cell in cells}
                for coordinate, note in notes.items():
                    if coordinate not in indexed:
                        row, column = coordinate_to_tuple(coordinate)
                        cells.append(
                            {
                                'cell': coordinate,
                                'row': row,
                                'column': column,
                                'value': None,
                                'formula': None,
                                'type': 'n',
                                'style': None,
                                'note': note,
                            }
                        )
            state.put('cell_index', key, cells)
            info.update(
                scan_complete=complete,
                next_row=next_row,
                scanned_cells=len(cells),
                observed_rows=max((c['row'] for c in cells), default=0),
                observed_columns=max((c['column'] for c in cells), default=0),
                merged_ranges=merges,
                validation_lists=validations,
                drawings=drawings,
                preview=cells[:12],
                continuation_reason=None if complete else 'cell_or_time_budget',
            )
            state.put('inventory', key, info)
            if complete:
                state.put(
                    'index_cache',
                    identity(entry['sha256'], title, INDEX_VERSION),
                    {
                        'cells': cells,
                        'metadata': {
                            k: info[k]
                            for k in (
                                'scan_complete',
                                'next_row',
                                'scanned_cells',
                                'observed_rows',
                                'observed_columns',
                                'merged_ranges',
                                'validation_lists',
                                'drawings',
                                'preview',
                                'continuation_reason',
                            )
                        },
                    },
                )
                propose_regions(state, key, cells, entry, path, title)
                inventory_images(state, archive, part, key, path, title, entry['sha256'])
            summaries.append({'id': key, **info})
        if sheet and not found:
            raise ValueError('Unknown worksheet: ' + sheet)
    state.put('inventory', file_id, entry | {'status': 'inspected'})
    return {
        'file_id': file_id,
        'sheets': summaries,
        'accounting_version': 2,
        'instruction': 'Query inspect_regions and search_cells; inspect candidates, then save source-backed roles. Discovery creates no business records.',
    }


def search(state, path, query, sheet=None, offset=0, limit=25):
    check_page(offset, limit)
    if not query.strip() or len(query) > 200:
        raise ValueError('Provide a text query of 1 to 200 characters')
    file_id, entry = state.file_entry(path)
    if not entry:
        raise ValueError('Unknown source')
    inventory = state.all('inventory')
    query = query.casefold().strip()
    matches, total = [], 0
    for key, source in inventory:
        if source.get('parent') != file_id or (sheet and source['sheet'] != sheet):
            continue
        cells = indexed_cells(state, key)
        selected = []
        for cell in cells:
            if query in (str(cell['value']) + ' ' + (cell.get('note') or '')).casefold():
                if offset <= total < offset + limit:
                    selected.append(cell)
                total += 1
        if selected:
            coordinates = {(c['row'], c['column']): c for c in cells}
            for cell in selected:
                nearby = [
                    coordinates[row, column]
                    for row in range(cell['row'] - 1, cell['row'] + 2)
                    for column in range(cell['column'] - 2, cell['column'] + 3)
                    if (row, column) in coordinates
                ]
                matches.append({'sheet': source['sheet'], **cell, 'nearby': nearby})
    return page_result(matches, offset, total) | {
        'complete_index': all(s.get('scan_complete') for _, s in inventory if s.get('parent') == file_id)
    }


class RegionRole(BaseModel):
    region: str
    role: Literal['records', 'reference', 'instructions', 'outside_scope', 'unresolved', 'unsupported']
    rationale: str = Field(min_length=1, max_length=2000)
    evidence: list[str] = Field(min_length=1, max_length=20)


def save_role(state, decision):
    region = state.get('region', decision.region)
    if not region or region.get('superseded'):
        raise ValueError(
            'Unknown or superseded region. Query inspect_regions for the exact region ID; available IDs: '
            + str(
                [
                    {'id': key, 'sheet': r['sheet'], 'bounds': r['bounds']}
                    for key, r in state.all('region')
                    if not r.get('superseded')
                ][:15]
            )
        )
    for key in decision.evidence:
        evidence = state.get('evidence', key)
        if not evidence or evidence['source_hash'] != region['source_hash'] or evidence.get('actor', 'main') != 'main':
            raise ValueError('Inspect source content and cite its evidence_id; filenames are insufficient')
    if decision.role != 'records':
        record_ids = {key for key, item in state.all('observation') if item.get('region') == decision.region}
        if any(record_ids.intersection(item['sources']) for _, item in state.all('assignment')):
            raise ValueError('Region has output contributions; reconcile these before changing its role')
    result = decision.model_dump() | {
        'source_hash': region['source_hash'],
        'instruction_revision': state.revision('instruction_revision'),
    }
    state.put('region_role', decision.region, result)
    return result


def inventory_images(state, archive, sheet_part, source_id, path, sheet, source_hash):
    folder, name = posixpath.split(sheet_part)
    relpart = folder + '/_rels/' + name + '.rels'
    if relpart not in archive.namelist():
        return
    for relationship in ET.fromstring(archive.read(relpart)):
        if not relationship.attrib.get('Type', '').endswith('/drawing'):
            continue
        drawing = posixpath.normpath(posixpath.join(folder, relationship.attrib['Target']))
        drawing = drawing.lstrip('/')
        if drawing not in archive.namelist():
            continue
        dfolder, dname = posixpath.split(drawing)
        drels = dfolder + '/_rels/' + dname + '.rels'
        links = (
            {r.attrib['Id']: r.attrib for r in ET.fromstring(archive.read(drels))}
            if drels in archive.namelist()
            else {}
        )
        for anchor in ET.fromstring(archive.read(drawing)):
            if not list(anchor.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip')):
                key = identity(source_id, drawing, ET.tostring(anchor).decode())
                state.put(
                    'visual',
                    key,
                    {
                        'path': path,
                        'sheet': sheet,
                        'source_hash': source_hash,
                        'artifact': None,
                        'anchor': ET.tostring(anchor).decode()[:3000],
                        'status': 'unsupported',
                        'instruction': 'Native chart/drawing has no supported worksheet renderer. Inspect underlying evidence to justify outside_scope, otherwise retain a material limitation.',
                    },
                )
            for blip in anchor.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
                rel = links.get(blip.attrib.get(REL + 'embed'), {})
                target = posixpath.normpath(posixpath.join(dfolder, rel.get('Target', ''))).lstrip('/')
                if target not in archive.namelist():
                    continue
                payload = archive.read(target)
                extension = posixpath.splitext(target)[1].lower()
                artifact = 'working/images/' + hashlib.sha256(payload).hexdigest() + extension
                atomic_bytes(state.root / artifact, payload)
                key = identity(source_id, target, ET.tostring(anchor).decode())
                state.put(
                    'visual',
                    key,
                    {
                        'path': path,
                        'sheet': sheet,
                        'source_hash': source_hash,
                        'artifact': artifact,
                        'anchor': ET.tostring(anchor).decode()[:3000],
                        'status': 'unresolved',
                        'instruction': 'Use view_page on artifact, then classify its relevance with source evidence.',
                    },
                )


def region_coverage(state):
    issues, counts = [], defaultdict(int)
    # Runs at every checkpoint: load each kind once instead of querying per record.
    revision = state.revision('instruction_revision')
    observations = state.all('observation')
    assignments = {source for _, a in state.all('assignment') for source in a['sources']}
    decisions = dict(state.all('decision'))
    inventory = dict(state.all('inventory'))
    regions = dict(state.all('region'))
    roles = dict(state.all('region_role'))
    mappings = dict(state.all('region_mapping'))
    evidence = dict(state.all('evidence'))
    files = {item['path']: (key, item) for key, item in inventory.items() if item['kind'] == 'file'}
    parents = {source.get('parent') for source in inventory.values()}
    template_ids = {key for key, source in inventory.items() if source.get('role') == 'template'}
    template_sheets = {key for key, source in inventory.items() if source.get('parent') in template_ids}
    excluded = {
        key
        for key, source in inventory.items()
        if source.get('status') == 'excluded'
        and evidence.get(source.get('evidence', ''), {}).get('source_hash') == files[source['path']][1]['sha256']
        and source.get('instruction_revision') == revision
    }
    excluded_sheets = {key for key, source in inventory.items() if key in excluded or source.get('parent') in excluded}
    current = []
    for key, item in observations:
        region_id = item.get('region')
        if region_id:
            region, role, mapping = regions.get(region_id, {}), roles.get(region_id, {}), mappings.get(region_id, {})
            active = (
                region
                and not region.get('superseded')
                and role.get('role') in ('records', 'reference')
                and not role.get('stale')
                and role.get('instruction_revision') == revision
                and mapping.get('role') == role.get('role')
                and not mapping.get('stale')
                and mapping.get('instruction_revision') == revision
                and item.get('mapping_version') == mapping.get('hash')
                and bool(item.get('reference_only')) == (role['role'] == 'reference')
            )
            if not active:
                if key in assignments:
                    issues.append({'id': key, 'reason': 'Output contribution uses an obsolete region interpretation'})
                continue
        current.append((key, item))
    extracted = {item['region'] for _, item in current if item.get('region') and not item.get('reference_only')}
    for key, region in regions.items():
        if region['source'] in excluded_sheets:
            continue
        if region.get('superseded') or region['source'] in template_sheets:
            continue
        role = roles.get(key, {})
        counts[role.get('role', 'unclassified')] += 1
        if (
            not role
            or role.get('stale')
            or role.get('instruction_revision') != revision
            or role['role'] in ('unresolved', 'unsupported')
        ):
            issues.append({'id': key, 'reason': 'Region scope remains unresolved'})
        if role.get('role') == 'records' and key not in extracted:
            issues.append({'id': key, 'reason': 'Applicable region has not been extracted'})
    record_counts = {'pending': 0, 'included': 0, 'excluded': 0, 'unresolved': 0}
    for key, item in current:
        if item.get('reference_only'):
            continue
        if item.get('source') in excluded_sheets or item.get('file') in excluded:
            record_counts['excluded'] += 1
            continue
        decision = decisions.get(key, {})
        status = decision.get('status', 'pending') if decision.get('instruction_revision', 0) == revision else 'pending'
        if key in assignments and status not in ('excluded', 'unresolved'):
            status = 'included'
        record_counts[status] += 1
        if status in ('pending', 'unresolved'):
            issues.append({'id': key, 'reason': 'Applicable record has no output assignment or supported exclusion'})
    for key, visual in state.all('visual'):
        if files.get(visual['path'], (None, None))[0] in excluded:
            continue
        if visual.get('status') != 'resolved' or visual.get('instruction_revision', revision) != revision:
            issues.append({'id': key, 'reason': 'Embedded visual content remains unresolved'})
    for key, source in inventory.items():
        if key in excluded_sheets:
            continue
        if source.get('role') == 'template':
            continue
        parent = inventory.get(source.get('parent'), {}) if source.get('parent') else {}
        if parent.get('role') == 'template':
            continue
        if source['kind'] == 'sheet' and not source.get('scan_complete') and source.get('status') != 'excluded':
            issues.append({'id': key, 'reason': 'Sheet discovery is incomplete'})
        if (
            source['kind'] == 'file'
            and source['path'].lower().endswith('.xlsx')
            and source.get('status') != 'excluded'
            and key not in parents
        ):
            issues.append({'id': key, 'reason': 'Workbook sheets have not been discovered'})
        if (
            source['kind'] == 'file'
            and not source['path'].lower().endswith('.xlsx')
            and (source.get('status') != 'reviewed' or not state.get('read', source['path'], {}).get('complete'))
        ):
            issues.append({'id': key, 'reason': 'Document has no reviewed disposition'})
    active_bounds = defaultdict(list)
    for region in regions.values():
        if not region.get('superseded'):
            active_bounds[region['source']].append(range_boundaries(region['bounds']))
    for source_id, source in inventory.items():
        if (
            source['kind'] != 'sheet'
            or source_id in template_sheets
            or source_id in excluded_sheets
            or not source.get('scan_complete')
        ):
            continue
        for cell in indexed_cells(state, source_id):
            if not any(
                left <= cell['column'] <= right and top <= cell['row'] <= bottom
                for left, top, right, bottom in active_bounds[source_id]
            ):
                issues.append({'id': source_id, 'reason': 'Unclassified occupied source area: ' + cell['cell']})
                break
    contributions = {}
    for key, item in current:
        if key not in assignments or item.get('reference_only'):
            continue
        source = files[item['path']][1]
        for cell in item.get('cells', []):
            pointer = identity(source['sha256'], item['sheet'], cell)
            if pointer in contributions and contributions[pointer] != key:
                issues.append(
                    {'id': key, 'reason': 'Overlapping or duplicate business contributions require reconciliation'}
                )
            contributions[pointer] = key
    return {'regions': dict(counts), 'records': record_counts, 'issues': issues}
