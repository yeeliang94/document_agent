"""Bounded native workbook reads with lossless access to oversized cell text."""

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple

from .regions import sheet_layout
from .state import encode


def unsized(sheet):
    if sheet.max_row is None or sheet.max_column is None:
        sheet.calculate_dimension(force=True)
    return sheet


def read_sheet(source, sheet, start_row, end_row, start_column, end_column, formulas, limit):
    if (
        start_row < 1
        or start_column < 1
        or end_row < start_row
        or end_column < start_column
        or end_column > 16384
        or end_row > 1048576
        or (end_row - start_row + 1) * (end_column - start_column + 1) > 10000
    ):
        raise ValueError('Use a valid range of at most 10000 cells')
    # Reserve room for application evidence, capability metadata and reading warnings.
    budget = limit - 650
    workbook = load_workbook(source, read_only=True, data_only=not formulas)
    try:
        sh = unsized(workbook[sheet])
        notes, merges = sheet_layout(source, sheet)
        last = min(end_row, sh.max_row)
        result = {
            'sheet': sheet,
            'rows': [],
            'empty_cells_omitted': 0,
            'merged_ranges': [],
            'last_row': last,
            'sheet_rows': sh.max_row,
            'more_rows': end_row < sh.max_row,
            'next_row': end_row + 1 if end_row < sh.max_row else None,
            'stopped_at_size_limit': False,
        }
        # Merge metadata must not force otherwise bounded row content through generic shortening.
        for merge in merges[:200]:
            if len(encode(result['merged_ranges'] + [merge]).encode()) > budget // 4:
                break
            result['merged_ranges'].append(merge)
        if len(result['merged_ranges']) < len(merges):
            result['merged_ranges_omitted'] = len(merges) - len(result['merged_ranges'])
        result['formulas'] = formulas
        # Counters and cursor digits may grow after this estimate.
        size = len(encode(result).encode()) + 32
        for number, row in enumerate(
            sh.iter_rows(min_row=start_row, max_row=last, min_col=start_column, max_col=end_column), start_row
        ):
            cells, empty = [], 0
            for column, cell in enumerate(row, start_column):
                coordinate = f'{get_column_letter(column)}{number}'
                note = notes.get(coordinate)
                if cell.value is None and not note:
                    empty += 1
                    continue
                item = {
                    'cell': coordinate,
                    'value': cell.value.isoformat() if hasattr(cell.value, 'isoformat') else cell.value,
                }
                if cell.data_type not in ('s', 'n'):
                    item['data_type'] = cell.data_type
                if (cell.number_format or 'General') != 'General':
                    item['number_format'] = cell.number_format
                if note:
                    item['note'] = note
                    if len(note) == 2000:
                        item['note_preview'] = True
                cells.append(item)
            if not cells:
                result['empty_cells_omitted'] += empty
                continue
            entry = {'row': number, 'cells': cells}
            cost = len(encode(entry).encode()) + 2
            if size + cost > budget:
                result.update(last_row=number - 1, next_row=number, more_rows=True, stopped_at_size_limit=True)
                if not result['rows']:
                    result['instruction'] = (
                        'Narrow columns or page value/note text with read_cell. Finish this row before advancing.'
                    )
                break
            result['rows'].append(entry)
            result['empty_cells_omitted'] += empty
            size += cost
        return result
    finally:
        workbook.close()


def read_cell(source, sheet, cell, field, offset, formulas, limit):
    if offset < 0 or field not in ('value', 'note'):
        raise ValueError('Use a nonnegative character offset and field value or note')
    try:
        row, column = coordinate_to_tuple(cell.upper())
    except (ValueError, KeyError) as error:
        raise ValueError('Use one Excel cell coordinate') from error
    if not 1 <= row <= 1048576 or not 1 <= column <= 16384:
        raise ValueError('Use one valid Excel cell coordinate')
    coordinate = f'{get_column_letter(column)}{row}'
    if field == 'note':
        notes, _ = sheet_layout(source, sheet, note_cell=coordinate)
        text = notes.get(coordinate) or ''
    else:
        workbook = load_workbook(source, read_only=True, data_only=not formulas)
        try:
            value = workbook[sheet].cell(row, column).value
            text = '' if value is None else value.isoformat() if hasattr(value, 'isoformat') else str(value)
        finally:
            workbook.close()
    chunk = text[offset : offset + 16000]
    while True:
        end = offset + len(chunk)
        result = {
            'sheet': sheet,
            'cell': coordinate,
            'field': field,
            'text': chunk,
            'offset': offset,
            'next_offset': end if end < len(text) else None,
        }
        if len(encode(result).encode()) <= limit - 650:
            return result
        if not chunk:
            raise ValueError('Cell metadata exceeds the result budget; use a larger tool-result limit')
        chunk = chunk[: len(chunk) // 2]
