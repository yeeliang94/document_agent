"""Typed bulk operations with durable source accounting and atomic workbook commits."""

import copy
import hashlib
import io
import operator
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Literal

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from .files import atomic_bytes, within
from .state import identity

Scalar = str | int | float | bool | None


class Column(BaseModel):
    name: str
    column: int = Field(ge=1, le=16384)
    kind: Literal['raw', 'identifier', 'text', 'decimal', 'date'] = 'raw'


class Predicate(BaseModel):
    field: str
    operator: Literal['eq', 'ne', 'lt', 'le', 'gt', 'ge', 'contains', 'blank'] = 'eq'
    value: Scalar = None


class Transformation(BaseModel):
    operation: Literal['filter', 'select', 'rename', 'convert', 'lookup', 'join', 'group', 'deduplicate', 'arithmetic']
    fields: list[str] = Field(default_factory=list)
    field: str = ''
    target: str = ''
    kind: Literal['raw', 'identifier', 'text', 'decimal', 'date'] = 'raw'
    predicate: Predicate | None = None
    reason: str = ''
    lookup: dict[str, Scalar] = Field(default_factory=dict)
    dataset: str = ''
    right_keys: list[str] = Field(default_factory=list)
    cardinality: Literal['one_to_one', 'many_to_one'] = 'many_to_one'
    aggregate: Literal['sum', 'count'] = 'sum'
    arithmetic: Literal['add', 'subtract', 'multiply', 'divide'] = 'add'
    operand: Scalar = None


class DestinationColumn(BaseModel):
    field: str
    column: int = Field(ge=1, le=16384)
    number_format: str | None = None


class Destination(BaseModel):
    sheet: str
    start_row: int = Field(ge=1, le=1048576)
    end_row: int = Field(ge=1, le=1048576)
    columns: list[DestinationColumn] = Field(min_length=1)
    required: list[str] = Field(default_factory=list)
    unique: list[str] = Field(default_factory=list)
    allowed: dict[str, list[Scalar]] = Field(default_factory=dict)
    totals: dict[str, str] = Field(default_factory=dict)


def convert(value, kind, number_format='General'):
    if value is None:
        return None
    if kind == 'identifier':
        if isinstance(value, (int, float)) and number_format and set(number_format) == {'0'}:
            return str(int(value)).zfill(len(number_format))
        return str(value)
    if kind == 'text':
        return str(value)
    if kind == 'decimal':
        number = Decimal(str(value))
        if not number.is_finite():
            raise ValueError('Nonfinite decimal')
        return str(number)
    if kind == 'date':
        return value.isoformat() if isinstance(value, (date, datetime)) else date.fromisoformat(str(value)).isoformat()
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def dataset_rows(state, dataset):
    info = state.get('dataset', dataset)
    if info is None:
        raise ValueError('Unknown dataset')
    return info, state.get('rows', dataset, [])


def lineage(state, dataset):
    """A dataset and every dataset it was derived or joined from."""
    seen, pending = set(), [dataset]
    while pending:
        key = pending.pop()
        if not key or key in seen:
            continue
        seen.add(key)
        pending.append((state.get('dataset', key) or {}).get('parent'))
        pending.extend(
            op.get('dataset')
            for op in (state.get('mapping', key) or {}).get('operations', [])
            if op.get('operation') == 'join'
        )
    return seen


def require_current(state, dataset):
    """Refuse a dataset that user instructions have invalidated, directly or through its sources."""
    revision = state.revision('instruction_revision')
    for key in lineage(state, dataset):
        info = state.get('dataset', key) or {}
        if info.get('instruction_revision', revision) != revision:
            raise ValueError('Instructions changed; re-extract/re-transform the dataset before writing')
        if info.get('stale'):
            raise ValueError(f'Dataset {key} was invalidated by a user instruction ({info["stale"]}); rebuild it')


def fresh_id(state, key):
    """The dataset ID for these inputs, skipping earlier builds that an instruction invalidated."""
    candidate, attempt = key, 0
    while (state.get('dataset', candidate) or {}).get('stale'):
        attempt += 1
        candidate = identity(key, attempt)
    return candidate


def settle(state, operation_id, journal):
    """Record a written journal's assignments and dispositions. Retries and recovery share this."""
    mapping = journal['mapping']
    state.delete_many('assignment', journal.get('removed_assignments', []))
    state.put_many('assignment', [(item['id'], item) for item in journal['assignments']])
    origins = lineage(state, mapping['dataset'])
    state.delete_many(
        'decision',
        [
            source
            for source in journal.get('resolved_sources', [])
            if state.get('decision', source, {}).get('evidence') in origins
        ],
    )
    for source, reason in journal.get('exclusions', []):
        # A disposition made under newer instructions outranks this older write.
        if state.get('decision', source, {}).get('instruction_revision', 0) <= mapping['instruction_revision']:
            state.put(
                'decision',
                source,
                {
                    'status': 'excluded',
                    'reason': reason,
                    'evidence': mapping['dataset'],
                    'instruction_revision': mapping['instruction_revision'],
                },
            )
    state.put('mapping', operation_id, mapping)
    state.put('exceptions', operation_id, journal.get('rejected', []))
    journal['status'] = 'completed'
    state.put('operation', operation_id, journal)


def extract_dataset(state, path, sheet, start_row, end_row, columns, header_rows, stop=lambda: None):
    if start_row < 1 or end_row < start_row or end_row > 1048576 or not columns:
        raise ValueError('Specify a valid region and at least one column')
    if len({c.name for c in columns}) != len(columns):
        raise ValueError('Column names must be unique')
    file_id, entry = state.file_entry(path)
    if not file_id:
        raise ValueError('Unknown source path')
    source_id = identity(file_id, sheet)
    source = state.get('inventory', source_id)
    if not source or not source.get('scan_complete'):
        raise ValueError('Complete inspect_source for this sheet before extracting records')
    if hashlib.sha256(within(state.root, path).read_bytes()).hexdigest() != entry['sha256']:
        raise ValueError('Source changed since intake; submit a new input version')
    dataset = fresh_id(
        state,
        identity(
            file_id,
            sheet,
            start_row,
            end_row,
            [c.model_dump() for c in columns],
            header_rows,
            state.revision('instruction_revision'),
        ),
    )
    if state.get('dataset', dataset):
        return state.get('dataset', dataset)
    formulas = load_workbook(within(state.root, path), read_only=True, data_only=False)
    values = load_workbook(within(state.root, path), read_only=True, data_only=True)
    rows, errors = [], []
    profiles = {column.name: {'types': {}, 'missing': 0} for column in columns}
    header_signatures = set()
    try:
        left, right = min(c.column for c in columns), max(c.column for c in columns)
        raw_rows = formulas[sheet].iter_rows(min_row=start_row, max_row=end_row, min_col=left, max_col=right)
        value_rows = values[sheet].iter_rows(min_row=start_row, max_row=end_row, min_col=left, max_col=right)
        observed = {key for key, _ in state.all('observation')}
        for number, (raw, cached) in enumerate(zip(raw_rows, value_rows, strict=True), start_row):
            if (number - start_row) % 250 == 0:
                stop()
            source_record = identity(file_id, sheet, number)
            if source_record not in observed:
                continue
            if number in header_rows:
                header_signatures.add(tuple(str(cell.value) for cell in raw))
                continue
            record = {'id': source_record, 'sources': [source_record], 'values': {}, 'evidence': {}, 'issues': []}
            if tuple(str(cell.value) for cell in raw) in header_signatures:
                record['issues'].append('Repeated header inside the declared data region')
            labels = {str(cell.value).strip().casefold() for cell in raw if isinstance(cell.value, str)}
            if labels & {'total', 'subtotal', 'grand total', 'grandtotal', 'sub total'}:
                record['issues'].append('Possible subtotal/total row; investigate or explicitly exclude')
            for column in columns:
                cell, cache = raw[column.column - left], cached[column.column - left]
                value = cache.value if cell.data_type == 'f' else cell.value
                profile = profiles[column.name]
                kind = type(value).__name__
                profile['types'][kind] = profile['types'].get(kind, 0) + 1
                profile['missing'] += value is None
                record['evidence'][column.name] = {
                    'path': path,
                    'sheet': sheet,
                    'cell': f'{get_column_letter(column.column)}{number}',
                    'raw': convert(cell.value, 'raw'),
                    'number_format': cell.number_format,
                    'formula': cell.data_type == 'f',
                }
                try:
                    if cell.data_type == 'f' and value is None:
                        raise ValueError('Formula has no cached result; calculation is required')
                    record['values'][column.name] = convert(value, column.kind, cell.number_format)
                except (ValueError, TypeError, InvalidOperation) as error:
                    record['values'][column.name] = None
                    record['issues'].append(f'{column.name}: {error}')
            if record['issues']:
                errors.append({'id': record['id'], 'issues': record['issues']})
            rows.append(record)
    finally:
        formulas.close()
        values.close()
    state.put('rows', dataset, rows)
    info = {
        'id': dataset,
        'path': path,
        'sheet': sheet,
        'source': source_id,
        'rows': len(rows),
        'columns': [c.name for c in columns],
        'kinds': {c.name: c.kind for c in columns},
        'exceptions': len(errors),
        'profiles': profiles,
        'instruction_revision': state.revision('instruction_revision'),
    }
    state.put('dataset', dataset, info)
    for number in header_rows:
        key = identity(file_id, sheet, number)
        if state.get('observation', key):
            state.put(
                'decision',
                key,
                {
                    'status': 'excluded',
                    'reason': 'Declared table header',
                    'evidence': f'{path}:{sheet}:{number}',
                    'instruction_revision': info['instruction_revision'],
                },
            )
    state.put('exceptions', dataset, errors)
    return info


def same_value(value, expected, kind):
    # Decimal fields are stored as text, so 4200 must equal '4200'. Identifiers keep exact text.
    if kind == 'decimal' and value is not None and expected is not None and not isinstance(expected, bool):
        try:
            return Decimal(str(value)) == Decimal(str(expected))
        except InvalidOperation:
            pass
    return value == expected


def matches(values, predicate, kinds=None):
    value, expected = values.get(predicate.field), predicate.value
    kind = (kinds or {}).get(predicate.field)
    if predicate.operator == 'blank':
        return value is None or value == ''
    if predicate.operator == 'contains':
        return str(expected) in str(value) if value is not None else False
    if predicate.operator == 'eq':
        return same_value(value, expected, kind)
    if predicate.operator == 'ne':
        return not same_value(value, expected, kind)
    if value is None or expected is None:
        return False
    try:
        left, right = Decimal(str(value)), Decimal(str(expected))
    except InvalidOperation:
        left, right = str(value), str(expected)
    return {
        'lt': lambda: left < right,
        'le': lambda: left <= right,
        'gt': lambda: left > right,
        'ge': lambda: left >= right,
    }[predicate.operator]()


def transform_dataset(state, dataset, operations, stop=lambda: None):
    info, original = dataset_rows(state, dataset)
    require_current(state, dataset)
    for op in operations:
        if op.operation == 'join':
            require_current(state, op.dataset)
    revision = state.revision('instruction_revision')
    result_id = fresh_id(state, identity(dataset, [op.model_dump() for op in operations], revision))
    if state.get('dataset', result_id):
        return state.get('dataset', result_id)
    rows = copy.deepcopy(original)
    kinds = dict(info.get('kinds', {}))
    exclusions = list(state.get('exclusions', dataset, []))
    for op in operations:
        stop()
        if op.operation == 'filter' and op.predicate and op.predicate.field not in kinds:
            raise ValueError('Unknown filter field')
        if op.operation == 'rename':
            kinds[op.target] = kinds.pop(op.field, 'raw')
        elif op.operation == 'convert':
            kinds[op.target or op.field] = op.kind
        elif op.operation == 'arithmetic':
            kinds[op.target or op.field] = 'decimal'
        elif op.operation == 'group':
            kinds = {field: kinds.get(field, 'raw') for field in op.fields} | {op.target: 'decimal'}
        elif op.operation == 'select':
            kinds = {field: kinds.get(field, 'raw') for field in op.fields}
        elif op.operation == 'lookup':
            kinds[op.target or op.field] = 'raw'
        elif op.operation == 'join':
            joined, _ = dataset_rows(state, op.dataset)
            kinds.update({(op.target or 'joined_') + field: kind for field, kind in joined.get('kinds', {}).items()})
        if op.operation == 'filter':
            if not op.predicate or not op.reason:
                raise ValueError('Filtering requires a predicate and an exclusion reason')
            kept = []
            for row in rows:
                if matches(row['values'], op.predicate, kinds):
                    kept.append(row)
                else:
                    exclusions.extend((source, op.reason) for source in row['sources'])
            rows = kept
        elif op.operation == 'join':
            if not op.fields or len(op.fields) != len(op.right_keys):
                raise ValueError('Specify equally sized explicit left and right join keys')
            _, other = dataset_rows(state, op.dataset)
            index = {}
            for row in other:
                key = tuple(row['values'].get(k) for k in op.right_keys)
                index.setdefault(key, []).append(row)
            left_counts = {}
            for row in rows:
                key = tuple(row['values'].get(k) for k in op.fields)
                left_counts[key] = left_counts.get(key, 0) + 1
            for row in rows:
                key = tuple(row['values'].get(k) for k in op.fields)
                found = index.get(key, [])
                if (
                    any(v in (None, '') for v in key)
                    or len(found) != 1
                    or (op.cardinality == 'one_to_one' and left_counts[key] != 1)
                ):
                    row['issues'].append('Missing or ambiguous join match')
                    continue
                match = found[0]
                row['sources'] = sorted(set(row['sources'] + match['sources']))
                row['issues'].extend(match['issues'])
                for field, value in match['values'].items():
                    target = (op.target or 'joined_') + field
                    if target in row['values']:
                        raise ValueError('Joined field collides with an existing field')
                    row['values'][target] = value
                    row['evidence'][target] = match['evidence'].get(field)
        elif op.operation in ('group', 'deduplicate'):
            if not op.fields:
                raise ValueError('Explicit grouping/deduplication keys are required')
            groups = {}
            for row in rows:
                key = tuple(row['values'].get(k) for k in op.fields)
                if any(v in (None, '') for v in key):
                    row['issues'].append('Missing grouping key')
                    key = ('unresolved', row['id'])
                groups.setdefault(key, []).append(row)
            merged = []
            for group in groups.values():
                row = copy.deepcopy(group[0])
                row['sources'] = sorted({source for item in group for source in item['sources']})
                row['id'] = identity(result_id, row['sources'])
                row['issues'] = [issue for item in group for issue in item['issues']]
                if op.operation == 'deduplicate':
                    if any(item['values'] != group[0]['values'] for item in group[1:]):
                        row['issues'].append('Conflicting records share a deduplication key')
                else:
                    try:
                        total = (
                            len(group)
                            if op.aggregate == 'count'
                            else sum((Decimal(str(item['values'][op.field])) for item in group), Decimal(0))
                        )
                        row['values'] = {field: row['values'][field] for field in op.fields} | {op.target: str(total)}
                    except (InvalidOperation, KeyError, TypeError) as error:
                        row['issues'].append(f'Invalid aggregate: {error}')
                merged.append(row)
            rows = merged
        else:
            for number, row in enumerate(rows):
                if number % 250 == 0:
                    stop()
                values = row['values']
                try:
                    if op.operation == 'select':
                        row['values'] = {key: values[key] for key in op.fields}
                    elif op.operation == 'rename':
                        if op.target in values and op.target != op.field:
                            raise ValueError('Rename target already exists')
                        values[op.target] = values.pop(op.field)
                        row['evidence'][op.target] = row['evidence'].pop(op.field, None)
                    elif op.operation == 'convert':
                        evidence = row['evidence'].get(op.field) or {}
                        values[op.target or op.field] = convert(
                            values[op.field], op.kind, evidence.get('number_format')
                        )
                    elif op.operation == 'lookup':
                        values[op.target or op.field] = op.lookup[str(values[op.field])]
                    elif op.operation == 'arithmetic':
                        left, right = Decimal(str(values[op.field])), Decimal(str(op.operand))
                        value = {
                            'add': operator.add,
                            'subtract': operator.sub,
                            'multiply': operator.mul,
                            'divide': operator.truediv,
                        }[op.arithmetic](left, right)
                        if not value.is_finite():
                            raise ValueError('Nonfinite arithmetic result')
                        values[op.target or op.field] = str(value)
                except (KeyError, ValueError, TypeError, ArithmeticError) as error:
                    row['issues'].append(f'{op.operation}: {error}')
    # A many-to-one join shares right-hand sources; one that still backs a surviving row is not excluded.
    surviving = {source for row in rows for source in row['sources']}
    exclusions = [(source, reason) for source, reason in exclusions if source not in surviving]
    state.put('rows', result_id, rows)
    result = {
        **info,
        'id': result_id,
        'parent': dataset,
        'rows': len(rows),
        'exceptions': sum(bool(row['issues']) for row in rows),
        'instruction_revision': revision,
        'columns': list(rows[0]['values']) if rows else info['columns'],
        'kinds': kinds,
    }
    state.put('dataset', result_id, result)
    state.put(
        'mapping',
        result_id,
        {'dataset': dataset, 'operations': [op.model_dump() for op in operations], 'instruction_revision': revision},
    )
    # Exclusions are facts about this proposed transformation, applied only with its output.
    state.put('exclusions', result_id, exclusions)
    return result


def populate_template(state, dataset, path, destination, operation_id, stop=lambda: None):
    info, rows = dataset_rows(state, dataset)
    require_current(state, dataset)
    target = within(state.root, path)
    if not target.is_relative_to(state.root / 'outputs'):
        raise ValueError('Only output copies can be edited')
    path = target.relative_to(state.root).as_posix()
    args = {'dataset': dataset, 'path': path, 'destination': destination.model_dump()}
    fingerprint = identity(args)
    old = state.get('operation', operation_id)
    current_hash = hashlib.sha256(target.read_bytes()).hexdigest()
    if old and old['fingerprint'] != fingerprint:
        raise ValueError('Operation ID was already used with different arguments')
    if old and old.get('sha256') == current_hash:
        if old['status'] == 'completed':
            return old['result']
        # Also closes the crash window after file replacement but before assignment persistence.
        settle(state, operation_id, old)
        return old['result']
    if old and old['status'] == 'completed':
        raise ValueError('Output changed since this operation; use a new operation ID')
    if old and old.get('before_sha256') != current_hash:
        raise ValueError('Output changed during an interrupted operation; reconcile before retrying')
    if destination.end_row < destination.start_row:
        raise ValueError('Invalid writable destination region')
    columns = [column.column for column in destination.columns]
    if len(columns) != len(set(columns)):
        raise ValueError('Destination columns overlap')
    if destination.start_row + len(rows) - 1 > destination.end_row:
        raise ValueError('Dataset exceeds the declared writable destination region')
    workbook = load_workbook(target)
    assignments, rejected, seen = [], [], set()
    totals = {field: Decimal(0) for field in destination.totals}
    previous_assignments = [
        (key, item)
        for key, item in state.all('assignment')
        if item['path'] == path
        and item['sheet'] == destination.sheet
        and destination.start_row <= item['row'] <= destination.end_row
    ]
    old_mappings = dict(state.all('mapping'))
    for _, assignment in previous_assignments:
        old_columns = {column['column'] for column in old_mappings[assignment['operation']]['destination']['columns']}
        if not old_columns.issubset(set(columns)):
            raise ValueError(
                'Partial-column overwrite would discard tracked cells; include all previously mapped columns'
            )
    try:
        sh = workbook[destination.sheet]
        affected_rows = set(range(destination.start_row, destination.start_row + len(rows)))
        affected_rows.update(item['row'] for _, item in previous_assignments)
        for number in affected_rows:
            for column in destination.columns:
                cell = sh.cell(number, column.column)
                if cell.data_type == 'f':
                    raise ValueError(f'Preserve formula {sh.title}!{cell.coordinate}')
                cell.value = None
        for index, row in enumerate(rows):
            if index % 250 == 0:
                stop()
            values, issues = row['values'], list(row['issues'])
            issues.extend(f'Required field: {key}' for key in destination.required if values.get(key) in (None, ''))
            issues.extend(
                f'Invalid code: {key}' for key, allowed in destination.allowed.items() if values.get(key) not in allowed
            )
            key = tuple(values.get(field) for field in destination.unique)
            if destination.unique and (key in seen or any(value in (None, '') for value in key)):
                issues.append('Duplicate or missing destination key')
            seen.add(key)
            if issues:
                rejected.append({'id': row['id'], 'sources': row['sources'], 'issues': issues})
                continue
            number = destination.start_row + index
            for column in destination.columns:
                if column.field not in values:
                    raise ValueError(f'Missing mapped field: {column.field}')
                cell = sh.cell(number, column.column)
                if cell.data_type == 'f':
                    raise ValueError(f'Preserve formula {sh.title}!{cell.coordinate}')
                numeric = info.get('kinds', {}).get(column.field) == 'decimal'
                cell.value = (
                    Decimal(values[column.field])
                    if numeric and values[column.field] is not None
                    else values[column.field]
                )
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                if column.number_format:
                    cell.number_format = column.number_format
            for field in totals:
                totals[field] += Decimal(str(values[field]))
            assignments.append(
                {
                    'id': identity(path, sh.title, number),
                    'path': path,
                    'sheet': sh.title,
                    'row': number,
                    'sources': row['sources'],
                    'dataset': dataset,
                    'operation': operation_id,
                    'values': {column.field: values[column.field] for column in destination.columns},
                    'numeric_fields': [
                        column.field
                        for column in destination.columns
                        if info.get('kinds', {}).get(column.field) == 'decimal'
                    ],
                }
            )
        for field, expected in destination.totals.items():
            if totals[field] != Decimal(expected):
                raise ValueError(f'Total mismatch for {field}: {totals[field]} != {expected}')
        buffer = io.BytesIO()
        workbook.save(buffer)
    finally:
        workbook.close()
    stop()
    payload = buffer.getvalue()
    digest = hashlib.sha256(payload).hexdigest()
    result = {'path': path, 'sha256': digest, 'records_written': len(assignments), 'unresolved': len(rejected)}
    exclusions = state.get('exclusions', dataset, [])
    journal = {
        'status': 'prepared',
        'fingerprint': fingerprint,
        'sha256': digest,
        'before_sha256': current_hash,
        'mapping': args | {'instruction_revision': info['instruction_revision']},
        'assignments': assignments,
        'exclusions': exclusions,
        'rejected': rejected,
        'removed_assignments': [key for key, _ in previous_assignments],
        'resolved_sources': sorted(
            {
                source
                for row in rows
                if row.get('resolutions') and not row['issues']
                for source in row['sources']
                if any(source in assignment['sources'] for assignment in assignments)
            }
        ),
        'result': result,
    }
    state.put('operation', operation_id, journal)
    atomic_bytes(target, payload)
    settle(state, operation_id, journal)
    return result


def recover_operation(state, operation_id, abort=False, reason=''):
    """Settle an interrupted commit without replaying a write or applying new business instructions."""
    journal = state.get('operation', operation_id)
    if not journal:
        raise ValueError('Unknown operation ID')
    if journal['status'] != 'prepared':
        return {'operation_id': operation_id, 'status': journal['status']}
    target = within(state.root, journal['result']['path'])
    digest = hashlib.sha256(target.read_bytes()).hexdigest() if target.is_file() else None
    if abort:
        if not reason.strip():
            raise ValueError('An abort requires an explicit reconciliation reason')
        journal.update(status='aborted', reason=reason)
    elif digest == journal['before_sha256']:
        journal.update(status='aborted', reason='No file replacement occurred; output retained')
    elif digest == journal['sha256']:
        settle(state, operation_id, journal)
    else:
        raise ValueError(
            'Output differs from both journal revisions; reconcile it and explicitly abort the old operation'
        )
    state.put('operation', operation_id, journal)
    return {'operation_id': operation_id, 'status': journal['status'], 'path': journal['result']['path']}
