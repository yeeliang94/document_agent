"""Bounded structural discovery; previews never establish record completeness."""

import hashlib
import time
from pathlib import Path

from openpyxl import load_workbook

from .files import within
from .state import identity


def inspect_source(state, path, sheet=None, start_row=1, max_cells=1_000_000, seconds=30, stop=lambda: None):
    if start_row < 1:
        raise ValueError('start_row must be positive')
    file_id, entry = state.file_entry(path)
    if entry is None:
        raise ValueError('Source must appear in the job inventory')
    source = within(state.root, path)
    if hashlib.sha256(source.read_bytes()).hexdigest() != entry['sha256']:
        raise ValueError('Source changed since intake; submit a new input version')
    if Path(path).suffix.lower() != '.xlsx':
        result = {
            **entry,
            'status': 'inspected',
            'content_reviewed': False,
            'instruction': 'Read the document, then record an evidence-backed reviewed or excluded disposition.',
        }
        state.put('inventory', file_id, result)
        return {'id': file_id, **result}
    from .regions import accounting_version, discover

    if accounting_version(state) == 2:
        return discover(state, path, sheet, start_row, max_cells, seconds, stop)
    # Reuse complete native discovery for identical content, retaining each occurrence's identity.
    if sheet is None and start_row == 1:
        inventory = state.all('inventory')
        duplicate = next(
            (
                key
                for key, item in inventory
                if key != file_id
                and item['kind'] == 'file'
                and item.get('sha256') == entry['sha256']
                and item['status'] == 'inspected'
            ),
            None,
        )
        children = [(key, item) for key, item in inventory if item.get('parent') == duplicate] if duplicate else []
        if children and all(item.get('scan_complete') for _, item in children):
            sheets = []
            observations = state.all('observation')
            for old_key, item in children:
                key = identity(file_id, item['sheet'])
                mapped = item | {'parent': file_id, 'path': path, 'status': 'inspected'}
                state.put('inventory', key, mapped)
                state.put_many(
                    'observation',
                    [
                        (
                            identity(file_id, item['sheet'], row['row']),
                            row | {'source': key, 'file': file_id, 'path': path},
                        )
                        for _, row in observations
                        if row['source'] == old_key
                    ],
                )
                sheets.append({'id': key, **mapped})
            state.put('inventory', file_id, entry | {'status': 'inspected'})
            return {'file_id': file_id, 'sheets': sheets, 'reused_content_hash': entry['sha256']}
    cached = [
        dict(id=key, **item)
        for key, item in state.all('inventory')
        if item.get('parent') == file_id and (sheet is None or item.get('sheet') == sheet)
    ]
    if cached and all(item.get('scan_complete') for item in cached):
        return {'file_id': file_id, 'sheets': cached}
    started = time.monotonic()
    workbook = load_workbook(source, read_only=False, data_only=False)
    try:
        summaries = []
        for sh in workbook:
            key = identity(file_id, sh.title)
            old = state.get('inventory', key, {})
            info = {
                **old,
                'kind': 'sheet',
                'parent': file_id,
                'path': path,
                'sheet': sh.title,
                'visibility': sh.sheet_state,
                'reported_rows': sh.max_row,
                'reported_columns': sh.max_column,
                'merged_ranges': [str(value) for value in sh.merged_cells.ranges],
                'tables': list(sh.tables),
                'status': old.get('status', 'discovered'),
                'validation_lists': [
                    {'range': str(v.sqref), 'formula': v.formula1} for v in sh.data_validations.dataValidation
                ],
            }
            # Normal-mode cells are sparse: inflated formatted dimensions do not create a huge rectangle.
            if sheet is not None and sh.title != sheet:
                state.put('inventory', key, info)
                continue
            next_row = old.get('next_row', 1)
            if not old.get('scan_complete') and start_row > next_row:
                raise ValueError(f'Continue {sh.title} at row {next_row}; skipping discovery is not allowed')
            cells = sorted(
                (cell for cell in sh._cells.values() if cell.value is not None or cell.comment),
                key=lambda cell: (cell.row, cell.column),
            )
            by_row = {}
            for cell in cells:
                if cell.row >= start_row:
                    by_row.setdefault(cell.row, []).append(cell)
            visited, completed, preview = 0, True, []
            observations = []
            for row_number, row in by_row.items():
                stop()
                if visited and (visited + len(row) > max_cells or time.monotonic() - started >= seconds):
                    next_row, completed = row_number, False
                    break
                visited += len(row)
                record = identity(file_id, sh.title, row_number)
                observations.append(
                    (record, {'source': key, 'file': file_id, 'path': path, 'sheet': sh.title, 'row': row_number})
                )
                if len(preview) < 8:
                    preview.append(
                        {
                            'row': row_number,
                            'cells': [
                                {
                                    'cell': c.coordinate,
                                    'value': str(c.value)[:200],
                                    'formula': c.data_type == 'f',
                                    'format': c.number_format,
                                    'note': c.comment.text[:200] if c.comment else None,
                                }
                                for c in row[:40]
                            ],
                        }
                    )
            state.put_many('observation', observations)
            info.update(
                status='inspected',
                scan_complete=completed,
                next_row=None if completed else next_row,
                preview=preview,
                scanned_cells=old.get('scanned_cells', 0) + visited,
                observed_rows=max((c.row for c in cells), default=0),
                observed_columns=max((c.column for c in cells), default=0),
            )
            state.put('inventory', key, info)
            summaries.append({'id': key, **info})
        entry['status'] = 'inspected'
        state.put('inventory', file_id, entry)
        return {
            'file_id': file_id,
            'sheets': summaries,
            'instruction': 'Identify table/header regions. Extract applicable regions and explicitly account for other rows.',
        }
    finally:
        workbook.close()
