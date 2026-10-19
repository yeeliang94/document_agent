"""The agent-facing interface for inventory, datasets, coverage and delegation."""

import hashlib
import json
from decimal import Decimal, InvalidOperation
from typing import Literal

from openpyxl import load_workbook
from pydantic import BaseModel, Field
from pydantic_ai import RunContext

from .datasets import (
    Column,
    Destination,
    Predicate,
    Scalar,
    Transformation,
    convert,
    dataset_rows,
    extract_dataset,
    fresh_id,
    matches,
    populate_template,
    recover_operation,
    require_current,
    transform_dataset,
)
from .inventory import inspect_source
from .state import UNSETTLED_TASKS, identity, paginate
from .tools import JobDeps, registrar

Kind = Literal['raw', 'identifier', 'text', 'decimal', 'date']


class Decision(BaseModel):
    id: str
    scope: Literal['source', 'record']
    status: Literal['reviewed', 'excluded', 'unresolved']
    reason: str = Field(min_length=1)
    evidence: str = Field(min_length=1)


class RecordResolution(BaseModel):
    record_id: str
    values: dict[str, Scalar]
    resolved_issues: list[str] = Field(default_factory=list)
    supporting_records: list[str] = Field(default_factory=list)
    reason: str = Field(min_length=1)
    evidence: str = Field(min_length=1)


class DocumentObservation(BaseModel):
    key: str = Field(min_length=1)
    page: int = Field(default=1, ge=1)
    values: dict[str, Scalar]
    # Undeclared fields are decimal when every recorded value is a number, otherwise raw.
    kinds: dict[str, Kind] = Field(default_factory=dict)
    evidence: str = Field(min_length=1)


def assignment_datasets(assignment):
    """Every dataset whose values an assignment still holds."""
    if 'provenance' in assignment:
        return sorted({cell['dataset'] for cell in assignment['provenance'].values() if cell['dataset']})
    return [assignment['dataset']] if assignment.get('dataset') else []


def observed_kind(values):
    present = [value for value in values if value is not None]
    numeric = all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in present)
    return 'decimal' if present and numeric else 'raw'


def validate_job(deps, path):
    state = deps.state
    coverage = state.coverage()
    failures = [item['reason'] + ': ' + item['id'] for item in coverage['issues']]
    for entry in deps.current()['files']:
        source = deps.resolve(entry['path'])
        if not source.is_file() or hashlib.sha256(source.read_bytes()).hexdigest() != entry['sha256']:
            failures.append('Source changed or disappeared: ' + entry['path'])
    output = deps.resolve(path, output=True)
    path = output.relative_to(state.root).as_posix()
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    template_validation = state.root / 'working/validation.json'
    structural = json.loads(template_validation.read_text()) if template_validation.exists() else {}
    if not structural.get('passed') or structural.get('sha256') != digest:
        failures.append('Run validate_output on the current workbook revision first')
    workbook = load_workbook(output, data_only=False)
    mappings = dict(state.all('mapping'))
    decisions = dict(state.all('decision'))
    try:
        for _, assignment in state.all('assignment'):
            if assignment['path'] != path:
                failures.append('Assignment belongs to a different output: ' + assignment['path'])
                continue
            mapping = mappings.get(assignment['operation'])
            if (
                not mapping
                or mapping['instruction_revision'] != state.revision('instruction_revision')
                or mapping.get('stale')
            ):
                failures.append('Output mapping predates current instructions')
                continue
            if any(
                cell['instruction_revision'] != state.revision('instruction_revision')
                for cell in assignment.get('provenance', {}).values()
            ):
                failures.append('Output cells predate current instructions')
                continue
            try:
                for dataset in assignment_datasets(assignment):
                    require_current(state, dataset)
            except ValueError as error:
                failures.append(str(error))
                continue
            destination = mapping['destination']
            if assignment['sheet'] not in workbook.sheetnames:
                failures.append('Missing assignment sheet')
                continue
            for column in destination['columns']:
                cell = workbook[assignment['sheet']].cell(assignment['row'], column['column'])
                actual = cell.value
                expected = assignment['values'][column['field']]
                try:
                    equal = (
                        cell.data_type == 'n'
                        and not isinstance(actual, bool)
                        and Decimal(str(actual)) == Decimal(str(expected))
                        if column['field'] in assignment.get('numeric_fields', []) and expected is not None
                        else actual == expected
                    )
                except InvalidOperation:
                    equal = False
                if not equal:
                    failures.append(f'Changed assignment {assignment["sheet"]}!{assignment["row"]}:{column["column"]}')
            for source in assignment['sources']:
                if decisions.get(source, {}).get('status') == 'excluded':
                    failures.append('Excluded record still has an output assignment: ' + source)
        for _, task in state.all('task'):
            if task['status'] in UNSETTLED_TASKS:
                failures.append('Delegated task has not settled: ' + task['id'])
        for operation_id, operation in state.all('operation'):
            if operation.get('result', {}).get('path') == path and operation.get('status') == 'prepared':
                failures.append('Interrupted output operation needs recovery: ' + operation_id)
    finally:
        workbook.close()
    result = {
        'path': path,
        'sha256': digest,
        'passed': not failures,
        'failures': failures[:100],
        'failure_count': len(failures),
        'revision': state.revision(),
        'instruction_revision': state.revision('instruction_revision'),
        'coverage': {key: value for key, value in coverage.items() if key != 'issues'},
        'limitation': 'No native Excel recalculation; semantic correctness still requires source-backed rules.',
    }
    state.put('validation', 'current', result)
    return result


def register_harness_tools(agent, delegation=None, readonly=False):
    tool = registrar(agent, readonly, {'inspect_source', 'query_inventory'})

    @tool
    async def inspect_source_tool(
        ctx: RunContext[JobDeps], path: str, sheet: str | None = None, start_row: int = 1
    ) -> dict:
        """Map a source's structure and all occupied rows. Continue incomplete scans at next_row; previews are samples."""

        def inspect():
            ctx.deps.resolve(path)
            settings = ctx.deps.settings
            return inspect_source(
                ctx.deps.state,
                path,
                sheet,
                start_row,
                settings.inspection_cells if settings else 1_000_000,
                settings.inspection_seconds if settings else 30,
                ctx.deps.check_stop,
            )

        return await ctx.deps.action('inspect_source', {'path': path, 'sheet': sheet, 'start_row': start_row}, inspect)

    @tool
    async def query_inventory(ctx: RunContext[JobDeps], offset: int = 0, limit: int = 25) -> dict:
        """Retrieve paginated source inventory with structural discovery and disposition status."""

        def query():
            if ctx.deps.actor == 'main':
                return ctx.deps.state.page('inventory', offset, limit)
            allowed_paths = (*ctx.deps.scope, *ctx.deps.state.get('task', ctx.deps.actor, {}).get('artifacts', []))
            items = [
                {'id': key, **item} for key, item in ctx.deps.state.all('inventory') if item['path'] in allowed_paths
            ]
            return paginate(items, offset, limit)

        return await ctx.deps.action('query_inventory', {'offset': offset, 'limit': limit}, query)

    @tool
    async def extract_dataset_tool(
        ctx: RunContext[JobDeps],
        path: str,
        sheet: str,
        start_row: int,
        end_row: int,
        columns: list[Column],
        header_rows: list[int],
    ) -> dict:
        """Extract a complete specified table region locally. Include explicit header rows; preserve all source pointers."""
        return await ctx.deps.action(
            'extract_dataset',
            {'path': path, 'sheet': sheet},
            lambda: extract_dataset(
                ctx.deps.state, path, sheet, start_row, end_row, columns, header_rows, ctx.deps.check_stop
            ),
        )

    @tool
    async def query_dataset(
        ctx: RunContext[JobDeps],
        dataset: str,
        offset: int = 0,
        limit: int = 25,
        predicate: Predicate | None = None,
        exceptions_only: bool = False,
    ) -> dict:
        """Read bounded records with evidence; filtering scans the stored dataset, not just a preview."""

        def query():
            info, rows = dataset_rows(ctx.deps.state, dataset)
            selected = [
                row
                for row in rows
                if (not predicate or matches(row['values'], predicate, info.get('kinds')))
                and (not exceptions_only or row['issues'])
            ]
            page = paginate(selected, offset, limit)
            return {'dataset': info['id'], 'rows': page.pop('items'), **page}

        return await ctx.deps.action('query_dataset', {'dataset': dataset, 'offset': offset}, query)

    @tool
    async def transform_dataset_tool(ctx: RunContext[JobDeps], dataset: str, operations: list[Transformation]) -> dict:
        """Execute typed bulk filters, joins, conversions, grouping, lookups and arithmetic; no arbitrary code."""
        return await ctx.deps.action(
            'transform_dataset',
            {'dataset': dataset, 'operations': [o.model_dump() for o in operations]},
            lambda: transform_dataset(ctx.deps.state, dataset, operations, ctx.deps.check_stop),
        )

    @tool
    async def populate_template_tool(
        ctx: RunContext[JobDeps], dataset: str, path: str, destination: Destination, operation_id: str
    ) -> dict:
        """Bulk write a stored dataset into explicit writable template bounds. Stable operation IDs make retries recoverable."""

        def populate():
            ctx.deps.resolve(path, output=True)
            return populate_template(ctx.deps.state, dataset, path, destination, operation_id, ctx.deps.check_stop)

        return await ctx.deps.action(
            'populate_template', {'dataset': dataset, 'path': path, 'operation_id': operation_id}, populate
        )

    @tool
    async def record_decisions(ctx: RunContext[JobDeps], decisions: list[Decision]) -> dict:
        """Account for non-tabular documents/irrelevant rows with a reason and evidence. Included rows require actual output assignments."""

        def record():
            for decision in decisions:
                if decision.scope == 'source':
                    source = ctx.deps.state.get('inventory', decision.id)
                    if not source:
                        raise ValueError('Unknown inventory item')
                    if (
                        decision.status == 'reviewed'
                        and source['kind'] == 'file'
                        and not source['path'].lower().endswith('.xlsx')
                        and not ctx.deps.state.get('read', source['path'], {}).get('complete')
                    ):
                        raise ValueError('Read all document chunks/pages (or vision pages) before marking it reviewed')
                    if (
                        decision.status == 'reviewed'
                        and source.get('kind') == 'sheet'
                        and not source.get('scan_complete')
                    ):
                        raise ValueError('An incompletely discovered sheet cannot be marked reviewed')
                    ctx.deps.state.put(
                        'inventory',
                        decision.id,
                        source
                        | decision.model_dump(exclude={'id', 'scope'})
                        | {'instruction_revision': ctx.deps.state.revision('instruction_revision')},
                    )
                else:
                    if decision.status == 'reviewed' or not ctx.deps.state.get('observation', decision.id):
                        raise ValueError(
                            'Rows can be excluded or unresolved; inclusion is established by writing output'
                        )
                    ctx.deps.state.put(
                        'decision',
                        decision.id,
                        decision.model_dump(exclude={'id', 'scope'})
                        | {'instruction_revision': ctx.deps.state.revision('instruction_revision')},
                    )
            return ctx.deps.state.coverage()

        return await ctx.deps.action('record_decisions', {'count': len(decisions)}, record)

    @tool
    async def record_observations(ctx: RunContext[JobDeps], path: str, observations: list[DocumentObservation]) -> dict:
        """Persist source-backed records extracted from documents/images as a dataset; use native extraction for Excel."""

        def record():
            ctx.deps.resolve(path)
            state = ctx.deps.state
            source_id, source = state.file_entry(path)
            if not source or path.lower().endswith('.xlsx') or not observations:
                raise ValueError('Provide observations from an inventoried non-Excel source')
            if len({item.key for item in observations}) != len(observations):
                raise ValueError('Observation keys must be unique')
            read = state.get('read', path, {})
            records, facts = [], []
            fields = sorted({field for item in observations for field in item.values})
            declared = {}
            for item in observations:
                for field, kind in item.kinds.items():
                    if declared.setdefault(field, kind) != kind:
                        raise ValueError(f'Conflicting kinds declared for {field}')
            kinds = {
                field: declared.get(field) or observed_kind(item.values.get(field) for item in observations)
                for field in fields
            }
            for item in observations:
                if str(item.page) not in read.get('content', {}):
                    raise ValueError('Read or view the cited source page before recording observations')
                key = identity(source_id, 'document', item.key)
                old = state.get('observation', key)
                if old and old.get('original_values') != item.values:
                    raise ValueError('Observation already exists; use resolve_records for source-backed corrections')
                facts.append(
                    (
                        key,
                        {
                            'source': source_id,
                            'file': source_id,
                            'path': path,
                            'page': item.page,
                            'key': item.key,
                            'original_values': item.values,
                            'evidence': item.evidence,
                        },
                    )
                )
                values, issues = {}, []
                for field in fields:
                    try:
                        values[field] = convert(item.values.get(field), kinds[field])
                    except (ValueError, TypeError, InvalidOperation) as error:
                        values[field] = None
                        issues.append(f'{field}: {error}')
                records.append(
                    {
                        'id': key,
                        'sources': [key],
                        'values': values,
                        'evidence': {
                            field: {'path': path, 'page': item.page, 'reference': item.evidence}
                            for field in item.values
                        },
                        'issues': issues,
                    }
                )
            dataset = fresh_id(
                state,
                identity(
                    source_id, [item.model_dump() for item in observations], state.revision('instruction_revision')
                ),
            )
            info = {
                'id': dataset,
                'path': path,
                'rows': len(records),
                'columns': fields,
                'kinds': kinds,
                'exceptions': sum(bool(record['issues']) for record in records),
                'instruction_revision': state.revision('instruction_revision'),
            }
            state.put_many('observation', facts)
            state.put('rows', dataset, records)
            state.put('dataset', dataset, info)
            return info

        return await ctx.deps.action('record_observations', {'path': path, 'count': len(observations)}, record)

    @tool
    async def resolve_records(ctx: RunContext[JobDeps], dataset: str, resolutions: list[RecordResolution]) -> dict:
        """Create a revised dataset from source-backed semantic decisions. Explicitly name corrected values and resolved issues."""

        def resolve():
            state = ctx.deps.state
            info, rows = dataset_rows(state, dataset)
            require_current(state, dataset)
            indexed = {row['id']: row for row in rows}
            observations = dict(state.all('observation'))
            for resolution in resolutions:
                if resolution.record_id not in indexed:
                    raise ValueError('Unknown dataset record')
                row = indexed[resolution.record_id]
                if any(field not in row['values'] for field in resolution.values):
                    raise ValueError('Resolution references an unknown field')
                if any(issue not in row['issues'] for issue in resolution.resolved_issues):
                    raise ValueError('Resolution references an unknown issue')
                if any(key not in observations for key in resolution.supporting_records):
                    raise ValueError('Unknown supporting source record')
                row['values'].update(resolution.values)
                row['issues'] = [issue for issue in row['issues'] if issue not in resolution.resolved_issues]
                row['sources'] = sorted(set(row['sources'] + resolution.supporting_records))
                row.setdefault('resolutions', []).append(resolution.model_dump())
            key = fresh_id(
                state,
                identity(dataset, [item.model_dump() for item in resolutions], state.revision('instruction_revision')),
            )
            updated = info | {
                'id': key,
                'parent': dataset,
                'exceptions': sum(bool(row['issues']) for row in rows),
                'instruction_revision': state.revision('instruction_revision'),
                'resolutions': [item.model_dump() for item in resolutions],
            }
            state.put('rows', key, rows)
            state.put('dataset', key, updated)
            state.put('exclusions', key, state.get('exclusions', dataset, []))
            for resolution in resolutions:
                for source in indexed[resolution.record_id]['sources']:
                    state.put(
                        'decision',
                        source,
                        {
                            'status': 'unresolved',
                            'reason': 'Resolved dataset requires a fresh output assignment',
                            'evidence': key,
                            'instruction_revision': updated['instruction_revision'],
                        },
                    )
            return updated

        return await ctx.deps.action('resolve_records', {'dataset': dataset, 'count': len(resolutions)}, resolve)

    @tool
    async def recover_output_operation(
        ctx: RunContext[JobDeps], operation_id: str, abort: bool = False, reason: str = ''
    ) -> dict:
        """Settle a prepared output journal after restart, even if instructions changed. Abort never changes workbook bytes."""
        return await ctx.deps.action(
            'recover_output_operation',
            {'operation_id': operation_id, 'abort': abort},
            lambda: recover_operation(ctx.deps.state, operation_id, abort, reason),
        )

    @tool
    async def read_job_state(
        ctx: RunContext[JobDeps],
        kind: Literal[
            'snapshot', 'memory', 'mapping', 'dataset', 'decision', 'observation', 'exceptions', 'task', 'operation'
        ] = 'snapshot',
        key: str | None = None,
        offset: int = 0,
        limit: int = 25,
    ) -> dict:
        """Retrieve selected durable job facts after compaction; use IDs and pagination for details."""

        def read():
            if kind == 'snapshot':
                return ctx.deps.state.snapshot()
            if key:
                return {'id': key, 'value': ctx.deps.state.get(kind, key)}
            return ctx.deps.state.page(kind, offset, limit)

        return await ctx.deps.action('read_job_state', {'kind': kind, 'key': key, 'offset': offset}, read)

    @tool
    async def validate_job_tool(ctx: RunContext[JobDeps], path: str) -> dict:
        """Check source coverage and current output assignments. Run validate_output first; complete requires both checks."""
        return await ctx.deps.action('validate_job', {'path': path}, lambda: validate_job(ctx.deps, path))

    if readonly:
        return

    @agent.tool(sequential=True)
    async def apply_instructions(
        ctx: RunContext[JobDeps],
        instructions: list[int],
        reason: str,
        invalidate: list[str] | None = None,
        reset_all: bool = False,
    ) -> dict:
        """Settle new user instructions (numbers from read_job_state). Name the dataset, mapping, decision, inventory or task IDs they affect; reset_all only when every prior result is affected."""

        def apply():
            if not reason.strip():
                raise ValueError('Explain how the instructions affect existing work')
            return ctx.deps.state.apply_instructions(instructions, reason, invalidate or [], reset_all)

        return await ctx.deps.action(
            'apply_instructions',
            {'instructions': instructions, 'invalidate': invalidate or [], 'reset_all': reset_all},
            apply,
        )

    if delegation is None:
        return

    @agent.tool(sequential=True)
    async def delegate_task(
        ctx: RunContext[JobDeps],
        objective: str,
        paths: list[str],
        role: Literal['scout', 'investigate', 'review'] = 'scout',
    ) -> dict:
        """Start a bounded read-only sub-agent with its own context. It cannot write the output or spawn children."""
        return await ctx.deps.async_action(
            'delegate_task',
            {'objective': objective, 'paths': paths, 'role': role},
            delegation.delegate(ctx.deps, objective, paths, role),
        )

    @agent.tool(sequential=True)
    async def get_task_status(ctx: RunContext[JobDeps], task_id: str) -> dict:
        """Read a persisted delegated task's status, handover, evidence and limitations."""
        return await ctx.deps.async_action(
            'get_task_status', {'task_id': task_id}, delegation.status(ctx.deps, task_id)
        )

    @agent.tool(sequential=True)
    async def cancel_task(ctx: RunContext[JobDeps], task_id: str) -> dict:
        """Cancel an unnecessary delegated task, retaining its history and findings."""
        return await ctx.deps.async_action('cancel_task', {'task_id': task_id}, delegation.cancel(ctx.deps, task_id))
