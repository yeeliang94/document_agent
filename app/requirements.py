"""Source-backed destination requirements and checks against current output bytes."""

import itertools
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Literal

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from pydantic import BaseModel, Field, field_validator

from .files import within
from .state import identity


class RequiredField(BaseModel):
    column: int = Field(ge=1, le=16384)
    meaning: str = Field(min_length=1, max_length=500)
    required: bool = True
    kind: Literal['any', 'identifier', 'decimal', 'date'] = 'any'
    allowed_values: list[str] = Field(default_factory=list, max_length=100)
    unique: bool = False
    expected_total: str | None = Field(
        default=None, description='Expected numeric sum as a plain number such as "10000.00"; omit when unknown.'
    )

    @field_validator('expected_total')
    @classmethod
    def numeric_total(cls, expected):
        # A rule written as prose here made validation fail later as an "uncheckable" value.
        try:
            valid = expected is None or Decimal(expected).is_finite()
        except InvalidOperation:
            valid = False
        if not valid:
            raise ValueError(f'expected_total must be a plain number or omitted, not {expected!r}')
        return expected


class OutputArea(BaseModel):
    id: str = Field(min_length=1, max_length=100)
    sheet: str
    bounds: str
    role: Literal['required', 'optional', 'reference', 'preserved', 'unresolved']
    fields: list[RequiredField] = Field(default_factory=list, max_length=100)
    min_records: int = Field(default=0, ge=0)
    evidence: list[str] = Field(min_length=1, max_length=20)
    rationale: str = Field(min_length=1, max_length=2000)


class OutputSpecification(BaseModel):
    objective: str = Field(min_length=1, max_length=3000)
    path: str
    areas: list[OutputArea] = Field(min_length=1, max_length=50)
    material_questions: list[str] = Field(default_factory=list, max_length=20)


def contains(outer, inner):
    a, b, c, d = range_boundaries(outer)
    e, f, g, h = range_boundaries(inner)
    return a <= e <= g <= c and b <= f <= h <= d


def covered(region, areas):
    """Whether the areas together cover every cell of region (all as min_col, min_row, max_col, max_row).

    A detected block is often split by meaning, e.g. headers, data and formula columns, so one area need not hold it.
    """
    left, top, right, bottom = region
    # Rows between consecutive area edges share the same set of overlapping areas; check each band once.
    cuts = {top, bottom + 1}
    for _, area_top, _, area_bottom in areas:
        cuts.update(edge for edge in (area_top, area_bottom + 1) if top < edge <= bottom)
    cuts = sorted(cuts)
    for first, end in itertools.pairwise(cuts):
        spans = sorted(
            (max(area_left, left), min(area_right, right))
            for area_left, area_top, area_right, area_bottom in areas
            if area_top <= first and area_bottom >= end - 1 and area_left <= right and area_right >= left
        )
        reach = left
        for start, stop in spans:
            if start > reach:
                return False
            reach = max(reach, stop + 1)
        if reach <= right:
            return False
    return True


def save_requirements(deps, specification):
    target = within(deps.root, specification.path)
    if not target.is_relative_to(deps.root / 'outputs'):
        # Requirements describe the deliverable; naming the template is the natural slip when only one copy exists.
        copies = sorted((deps.root / 'outputs').glob('*.xlsx'))
        if len(copies) != 1:
            raise ValueError('Save requirements for the output copy: run create_output, then use its outputs/ path')
        target = copies[0]
    deps.resolve(target.relative_to(deps.root).as_posix(), output=True)
    # validate_job looks requirements up by the normalized output path.
    specification = specification.model_copy(update={'path': target.relative_to(deps.root).as_posix()})
    state = deps.state
    template = next(entry for _, entry in state.all('inventory') if entry.get('role') == 'template')
    if len({area.id for area in specification.areas}) != len(specification.areas):
        raise ValueError('Area IDs must be unique')
    from .regions import usable_evidence

    available = usable_evidence(state, template['sha256'], deps.actor)
    ignored = sorted({key for area in specification.areas for key in area.evidence if key not in available})
    for area in specification.areas:
        bounds = range_boundaries(area.bounds)
        if None in bounds or min(bounds) < 1 or bounds[2] > 16384 or bounds[3] > 1048576:
            raise ValueError('Use explicit valid Excel ranges')
        if (bounds[2] - bounds[0] + 1) * (bounds[3] - bounds[1] + 1) > 1_000_000:
            raise ValueError('Requirement areas must be bounded to one million cells')
        if any(not bounds[0] <= field.column <= bounds[2] for field in area.fields):
            raise ValueError('Field column is outside its area')
        # Each area needs one real inspection of the destination; extra mistyped citations are dropped.
        if not any(key in available for key in area.evidence):
            raise ValueError(
                f'Area {area.id}: cite an evidence_id from inspecting the destination. Available references: '
                + str([{'id': key, 'location': location} for key, location in available.items()][:8])
            )
        area.evidence = [key for key in area.evidence if key in available]
    result = specification.model_dump() | {
        'instruction_revision': state.revision('instruction_revision'),
        'template_hash': template['sha256'],
    }
    key = identity(specification.path)
    old = state.get('requirement', key)
    if old:
        state.put('requirement_history', identity(old), old)
    state.put('requirement', key, result)
    return {'id': key, **result} | ({'ignored_evidence': ignored} if ignored else {})


def output_row(sheet, original, row, left, right, assigned_columns):
    """Count non-empty assigned or changed rows, leaving unchanged template labels out.

    A header or label row copied from the template is not an output record.
    """
    cells = [(sheet.cell(row, col).value, col) for col in range(left, right + 1)]
    if all(value is None for value, _ in cells):
        return False
    return (
        any(left <= column <= right for column in assigned_columns)
        or original is None
        or any(value != original.cell(row, col).value for value, col in cells)
    )


def check_requirements(state, path):
    specification = state.get('requirement', identity(path))
    failures = []
    if not specification:
        return ['Save source-backed destination requirements before final validation']
    if specification.get('stale') or specification['instruction_revision'] != state.revision('instruction_revision'):
        failures.append('Destination requirements are stale')
    failures.extend('Unresolved requirement: ' + q for q in specification['material_questions'])
    template_ids = {key for key, source in state.all('inventory') if source.get('role') == 'template'}
    template_sheets = {key for key, source in state.all('inventory') if source.get('parent') in template_ids}
    if not template_sheets or any(not state.get('inventory', key).get('scan_complete') for key in template_sheets):
        failures.append('Destination-wide discovery has not completed')
    for _, region in state.all('region'):
        if region.get('superseded') or region['source'] not in template_sheets:
            continue
        areas = [
            range_boundaries(area['bounds'])
            for area in specification['areas']
            if area['sheet'] == region['sheet'] and area['role'] != 'unresolved'
        ]
        if not covered(range_boundaries(region['bounds']), areas):
            failures.append(
                'Unclassified destination area: '
                + region['sheet']
                + '!'
                + region['bounds']
                + ' (cover every cell with output areas; several adjoining areas are fine)'
            )
    workbook = load_workbook(within(state.root, path), data_only=False)
    template_entry = next((entry for _, entry in state.all('inventory') if entry.get('role') == 'template'), None)
    template = load_workbook(within(state.root, template_entry['path']), data_only=False) if template_entry else None
    try:
        mappings = dict(state.all('mapping'))
        assigned_columns = {}
        for _, assignment in state.all('assignment'):
            if assignment['path'] != path:
                continue
            destination = mappings.get(assignment['operation'], {}).get('destination', {})
            assigned_columns.setdefault((assignment['sheet'], assignment['row']), set()).update(
                column['column'] for column in destination.get('columns', [])
            )
        for area in specification['areas']:
            if area['role'] == 'unresolved':
                failures.append('Unresolved destination area: ' + area['id'])
                continue
            if area['sheet'] not in workbook.sheetnames:
                failures.append('Missing destination sheet: ' + area['sheet'])
                continue
            if area['role'] not in ('required', 'optional'):
                continue
            left, top, right, bottom = range_boundaries(area['bounds'])
            sheet = workbook[area['sheet']]
            original = template[area['sheet']] if template and area['sheet'] in template.sheetnames else None
            active = [
                row
                for row in range(top, bottom + 1)
                if output_row(sheet, original, row, left, right, assigned_columns.get((area['sheet'], row), ()))
            ]
            if len(active) < max(area['min_records'], 1 if area['role'] == 'required' else 0):
                failures.append('Too few output records: ' + area['id'])
            for field in area['fields']:
                values = [sheet.cell(row, field['column']).value for row in active]
                if field['required'] and any(v is None or v == '' for v in values):
                    failures.append('Missing required value: ' + field['meaning'])
                present = [v for v in values if v is not None and v != '']
                if field['unique'] and len(set(map(str, present))) != len(present):
                    failures.append('Duplicate values: ' + field['meaning'])
                if field['allowed_values'] and any(str(v) not in field['allowed_values'] for v in present):
                    failures.append('Disallowed value: ' + field['meaning'])
                if field['kind'] == 'identifier' and any(not isinstance(v, str) for v in present):
                    failures.append('Identifier must retain text representation: ' + field['meaning'])
                if field['kind'] == 'decimal' and any(
                    not isinstance(v, (int, float)) or isinstance(v, bool) for v in present
                ):
                    failures.append('Numeric field is not stored as an Excel number: ' + field['meaning'])
                if field['kind'] == 'date':
                    for value in present:
                        try:
                            if not isinstance(value, (date, datetime)):
                                date.fromisoformat(str(value))
                        except ValueError:
                            failures.append('Invalid effective date: ' + field['meaning'])
                            break
                if field['kind'] == 'decimal' or field['expected_total'] is not None:
                    try:
                        numbers = [Decimal(str(v)) for v in present]
                        if not all(v.is_finite() for v in numbers):
                            raise InvalidOperation
                        if field['expected_total'] is not None and sum(numbers) != Decimal(field['expected_total']):
                            failures.append('Total mismatch: ' + field['meaning'])
                    except (InvalidOperation, ValueError):
                        failures.append('Uncheckable numeric value or formula cache: ' + field['meaning'])
    finally:
        workbook.close()
        if template:
            template.close()
    return failures


def check_write_scope(state, path, sheet, first_row, last_row, columns):
    from .regions import accounting_version

    if accounting_version(state) != 2:
        return
    specification = state.get('requirement', identity(path), {})
    if (
        not specification
        or specification.get('stale')
        or specification['instruction_revision'] != state.revision('instruction_revision')
    ):
        raise ValueError('Save current output requirements before writing')
    areas = [
        (area['role'], range_boundaries(area['bounds'])) for area in specification['areas'] if area['sheet'] == sheet
    ]
    for column in columns:
        if any(
            role not in ('required', 'optional') and left <= column <= right and top <= last_row and first_row <= bottom
            for role, (left, top, right, bottom) in areas
        ):
            raise ValueError(f'Write intersects a protected destination area: {sheet}!{column}:{first_row}-{last_row}')
        if not any(
            role in ('required', 'optional') and left <= column <= right and top <= first_row <= last_row <= bottom
            for role, (left, top, right, bottom) in areas
        ):
            raise ValueError(f'Write falls outside established input areas: {sheet}!{column}:{first_row}-{last_row}')
