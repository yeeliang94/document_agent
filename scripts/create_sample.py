#!/usr/bin/env python3
"""Generate synthetic payroll inputs and a formula-bearing target workbook."""

import argparse
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook
from PIL import Image, ImageDraw


def create_sample(output: Path) -> Path:
    output.mkdir(parents=True, exist_ok=True)
    packet = output / 'payroll_packet'
    reports = packet / 'October 2026' / 'HR' / 'Reports'
    scans = packet / 'October 2026' / 'Supporting documents'
    reports.mkdir(parents=True, exist_ok=True)
    scans.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Worker Summary'
    sheet.append(['Employee ID', 'Name', 'Salary', 'Effective Date'])
    sheet.append(['EMP001', 'Alex Example', 4200, '2026-10-01'])
    sheet.append(['EMP002', 'Sam Sample', 5800, '2026-10-01'])
    notes = workbook.create_sheet('Instructions')
    notes.append(['Synthetic demo only. Currency: MYR. October 2026 monthly salary.'])
    workbook.save(reports / 'employee_report.xlsx')
    image = Image.new('RGB', (1200, 480), 'white')
    draw = ImageDraw.Draw(image)
    draw.multiline_text(
        (45, 45),
        'SYNTHETIC PAYROLL SUPPORTING NOTE\n\nEmployee ID: EMP002\nName: Sam Sample\nEffective: 2026-10-01\nMonthly salary confirmed: MYR 5800\n\nDemo data only. No real personal information.',
        fill='black',
        spacing=18,
        font_size=26,
    )
    image.save(scans / 'salary_confirmation.png')
    instructions = (
        'Prepare the October 2026 payroll workbook from employee_report.xlsx. '
        'Read all source tabs and the supporting image. Fill Employees!A2:D3 with the two '
        'employee records. Preserve the formulas in E2:E3, the hidden Lookup sheet, '
        'and workbook formatting. Keep Effective Date as the source ISO date text. '
        'Validate that both employees are included and salaries match. All data is synthetic.\n'
    )
    (packet / 'instructions.txt').write_text(instructions, encoding='utf-8')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Employees'
    sheet.append(['Employee ID', 'Name', 'Salary', 'Effective Date', 'Annual Salary'])
    sheet['E2'] = '=C2*12'
    sheet['E3'] = '=C3*12'
    sheet.freeze_panes = 'A2'
    for cell in sheet[1]:
        from openpyxl.styles import Font, PatternFill

        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='243B53')
    for col in ('A', 'B', 'C', 'D', 'E'):
        sheet.column_dimensions[col].width = 22
    lookup = workbook.create_sheet('Lookup')
    lookup.append(['Currency', 'MYR'])
    lookup.sheet_state = 'hidden'
    workbook.save(output / 'output_template.xlsx')
    (output / 'README.txt').write_text(
        '1. Select payroll.\n2. Upload payroll_packet.zip or the payroll_packet folder.\n3. Select output_template.xlsx as template.\n4. Paste instructions.txt into the run instructions.\n5. Start with a configured OpenAI key.\n\n'
        + instructions,
        encoding='utf-8',
    )
    with ZipFile(output / 'payroll_packet.zip', 'w') as archive:
        for file in sorted(packet.rglob('*')):
            if file.is_file():
                archive.write(file, file.relative_to(packet))
    return output


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        'output', nargs='?', type=Path, default=Path(__file__).resolve().parent.parent / 'data' / 'samples'
    )
    args = parser.parse_args()
    print(create_sample(args.output.resolve()))
