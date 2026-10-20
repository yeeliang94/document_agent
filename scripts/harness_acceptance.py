"""Synthetic 5,000 x 30 acceptance. --live explicitly invokes the configured running agent (paid)."""

import argparse
import io
import json
import sys
import time
from pathlib import Path

import httpx
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.config import Settings
from app.datasets import Column, Destination, DestinationColumn, dataset_rows, extract_dataset, populate_template
from app.files import add_file
from app.harness_tools import validate_job
from app.inventory import inspect_source
from app.store import Store
from app.tools import JobDeps, validate_workbook


def fixtures():
    source = Workbook()
    source.active.title = 'Employees'
    source.active.append(['Employee ID', *[f'Value {column}' for column in range(2, 31)], 'Check'])
    for row in range(1, 5001):
        # openpyxl stores no cached result, so the Check formula exercises missing-cache handling.
        source.active.append([row, *[row * column for column in range(2, 31)], f'=B{row + 1}*1'])
        source.active.cell(row + 1, 1).number_format = '000000'
    notes = source.create_sheet('Notes')
    notes.sheet_state = 'hidden'
    notes.append(['Prepared by payroll'])
    notes.append(['Not for upload'])
    template = Workbook()
    template.active.title = 'Output'
    template.active.append(['Employee ID', *[f'Value {column}' for column in range(2, 31)], 'Total'])
    template.active['AE2'] = '=SUM(B2:AD2)'
    template.create_sheet('Lookup').sheet_state = 'hidden'
    result = []
    for workbook in (source, template):
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        result.append(buffer.getvalue())
    return result


def verify(payload):
    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    try:
        sheet = workbook['Output']
        for row in range(1, 5001):
            assert sheet.cell(row + 1, 1).value == f'{row:06}', f'Identifier mismatch at record {row}'
            for column in range(2, 31):
                assert sheet.cell(row + 1, column).value == row * column, f'Value mismatch at {row}/{column}'
                assert sheet.cell(row + 1, column).data_type == 'n'
        assert sheet['AE2'].value == '=SUM(B2:AD2)'
        assert workbook['Lookup'].sheet_state == 'hidden'
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--live', action='store_true', help='Run paid model calls through the running application')
    parser.add_argument('--base-url', default='http://127.0.0.1:8003')
    parser.add_argument('--output', type=Path, default=ROOT / 'data/harness-tests')
    args = parser.parse_args()
    source, template = fixtures()
    instructions = (
        'Copy every Employees record into Output columns A:AD, preserving six-digit employee IDs and '
        'numeric amounts. Inspect all sheets, use bulk dataset operations, and preserve the formula '
        'in AE2 and hidden Lookup sheet. Account for the header row and the hidden Notes sheet; the Check '
        'column is not part of the output. You may delegate read-only scouting '
        'when useful. Validate workbook integrity and source coverage before completion.'
    )
    started = time.monotonic()
    args.output.mkdir(parents=True, exist_ok=True)
    if args.live:
        settings = Settings()
        headers = {'Authorization': 'Bearer ' + settings.access_token} if settings.access_token else {}
        with httpx.Client(base_url=args.base_url, headers=headers, timeout=180) as client:
            response = client.post(
                '/api/jobs',
                data={'module': 'payroll', 'period': '2026-10', 'instructions': instructions},
                files=[('files', ('nested/source.xlsx', source)), ('template', ('template.xlsx', template))],
            )
            response.raise_for_status()
            job_id = response.json()['id']
            client.post(f'/api/jobs/{job_id}/start').raise_for_status()
            deadline = time.monotonic() + 1900
            while time.monotonic() < deadline:
                job = client.get(f'/api/jobs/{job_id}').json()
                if job['status'] not in ('ready', 'queued', 'running', 'pausing'):
                    break
                time.sleep(2)
            else:
                raise RuntimeError('Live acceptance timed out; job remains available in application history')
            coverage = client.get(f'/api/jobs/{job_id}/coverage').json()
            report = {'job': job, 'coverage': coverage, 'live': True}
            folder = args.output / job_id
            folder.mkdir(exist_ok=True)
            (folder / 'report.json').write_text(json.dumps(report, indent=2))
            assert job['status'] == 'completed', f'Job stopped as {job["status"]}; inspect {folder}/report.json'
            artifact = next(item for item in job['artifacts'] if item['name'].endswith('.xlsx'))
            response = client.get(f'/api/jobs/{job_id}/artifacts/{artifact["name"]}')
            response.raise_for_status()
            payload = response.content
    else:
        store = Store(args.output)
        job = store.create('payroll', '2026-10', instructions, '')
        job_id, manifest = job['id'], []
        folder = store.path(job_id)
        add_file(folder, 'inputs', 'nested/source.xlsx', source, manifest)
        add_file(folder, 'template', 'template.xlsx', template, manifest)
        store.update(job_id, files=manifest)
        deps = JobDeps(store, job_id, Settings(data_dir=args.output))
        discovered = inspect_source(deps.state, 'inputs/nested/source.xlsx')
        hidden = next(sheet for sheet in discovered['sheets'] if sheet['sheet'] == 'Notes')
        assert hidden['visibility'] == 'hidden' and hidden['scan_complete'], 'Hidden source sheet was not discovered'
        checks = extract_dataset(
            deps.state, 'inputs/nested/source.xlsx', 'Employees', 1, 5001, [Column(name='check', column=31)], [1]
        )
        _, check_rows = dataset_rows(deps.state, checks['id'])
        assert checks['exceptions'] == 5000 and all(
            row['values']['check'] is None and any('no cached result' in issue for issue in row['issues'])
            for row in check_rows
        ), 'Formulas without cached results must stay exceptions'
        notes = [(key, item) for key, item in deps.state.all('observation') if item.get('sheet') == 'Notes']
        assert len(notes) == 2, 'Every populated hidden-sheet row must be observed'
        deps.state.put_many(
            'decision',
            [
                (key, {'status': 'excluded', 'reason': 'Preparer note', 'evidence': f'Notes!A{item["row"]}'})
                for key, item in notes
            ],
        )
        columns = [
            Column(name=f'f{column}', column=column, kind='identifier' if column == 1 else 'decimal')
            for column in range(1, 31)
        ]
        dataset = extract_dataset(deps.state, 'inputs/nested/source.xlsx', 'Employees', 1, 5001, columns, [1])
        path = 'outputs/output.xlsx'
        (folder / path).write_bytes(template)
        populate_template(
            deps.state,
            dataset['id'],
            path,
            Destination(
                sheet='Output',
                start_row=2,
                end_row=5001,
                columns=[DestinationColumn(field=f'f{column}', column=column) for column in range(1, 31)],
                required=['f1'],
                unique=['f1'],
            ),
            'acceptance',
        )
        payload = (folder / path).read_bytes()
        assert validate_workbook(deps, path)['passed']
        result = validate_job(deps, path)
        assert result['passed'], result
        report = {'job_id': job_id, 'coverage': result['coverage'], 'live': False}
    verify(payload)
    (folder / 'verified.xlsx').write_bytes(payload)
    report['elapsed_seconds'] = round(time.monotonic() - started, 2)
    report['verified_records'] = 5000
    (folder / 'report.json').write_text(json.dumps(report, indent=2))
    print(
        json.dumps(
            {
                'report': str(folder / 'report.json'),
                'elapsed_seconds': report['elapsed_seconds'],
                'verified_records': 5000,
            }
        )
    )


if __name__ == '__main__':
    main()
