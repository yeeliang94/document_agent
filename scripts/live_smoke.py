"""Explicit paid/live synthetic acceptance test against the running local MVP.

Run after ./start.sh: .venv/bin/python scripts/live_smoke.py
Uploads synthetic fixtures, runs gpt-6-luna twice, checks actual workbook bytes,
exercises a user question/answer, SSE replay and downloadable diagnostics.
"""

import argparse
import hashlib
import io
import json
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import httpx
from dotenv import dotenv_values
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.create_sample import create_sample


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run_scenario(client, base, packet, instructions, folder, question=False):
    with (packet / 'payroll_packet.zip').open('rb') as src, (packet / 'output_template.xlsx').open('rb') as template:
        response = client.post(
            base + '/api/jobs',
            data={'module': 'payroll', 'period': '2026-10', 'instructions': instructions},
            files=[
                ('files', ('payroll_packet.zip', src, 'application/zip')),
                (
                    'template',
                    (
                        'output_template.xlsx',
                        template,
                        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    ),
                ),
            ],
        )
    response.raise_for_status()
    job = response.json()
    job_id = job['id']
    print('Created', job_id, 'question=' + str(question), flush=True)
    before = {f['path']: f['sha256'] for f in job['files']}
    response = client.post(f'{base}/api/jobs/{job_id}/start')
    response.raise_for_status()
    deadline = time.monotonic() + 600
    cursor = 0
    events = []
    answered = False
    while time.monotonic() < deadline:
        response = client.get(f'{base}/api/jobs/{job_id}/events', params={'after': cursor})
        response.raise_for_status()
        new = response.json()
        events.extend(new)
        for event in new:
            cursor = max(cursor, event['seq'])
            if not event['name'].endswith('.started') and event['name'] != 'model.completed':
                print(job_id[:8], event['name'], str(event.get('message', ''))[:180], flush=True)
        response = client.get(f'{base}/api/jobs/{job_id}')
        response.raise_for_status()
        job = response.json()
        if job['status'] == 'waiting':
            if not question or answered:
                raise AssertionError('Unexpected user question: ' + str(job['question']))
            assert job['question'] and job['question']['text']
            print('Answering synthetic period question:', job['question']['text'], flush=True)
            response = client.post(
                f'{base}/api/jobs/{job_id}/answer',
                json={
                    'answer': 'Confirmed: use October 2026, reporting period 2026-10. Proceed with both synthetic employees and the supplied template.'
                },
            )
            response.raise_for_status()
            answered = True
        if job['status'] in ('completed', 'partial', 'failed', 'cancelled'):
            break
        time.sleep(1)
    else:
        raise TimeoutError('Live job timed out after 600 seconds')
    # Catch events published between the last event read and terminal state.
    tail = client.get(f'{base}/api/jobs/{job_id}/events', params={'after': cursor})
    tail.raise_for_status()
    events.extend(tail.json())
    if job['status'] == 'failed':
        raise AssertionError('Live agent failed: ' + str(job.get('error')))
    assert job['status'] == 'completed', job.get('summary')
    assert answered == question
    assert any(e['name'] == 'view_page' and e.get('status') == 'completed' for e in events), (
        'Agent did not inspect the scan'
    )
    assert any(e['name'] == 'write_cells' and e.get('status') == 'completed' for e in events)
    assert any(e['name'] == 'validate_output' and e.get('result', {}).get('passed') for e in events)
    artifact = next(a for a in job['artifacts'] if a['name'].endswith('.xlsx'))
    response = client.get(f'{base}/api/jobs/{job_id}/artifacts/{artifact["name"]}')
    response.raise_for_status()
    target = folder / (('question' if question else 'baseline') + '-completed.xlsx')
    target.write_bytes(response.content)
    w = load_workbook(io.BytesIO(response.content), data_only=False)
    expected = {'EMP001': ('Alex Example', 4200, '2026-10-01'), 'EMP002': ('Sam Sample', 5800, '2026-10-01')}
    actual = {
        w['Employees'].cell(row, 1).value: tuple(w['Employees'].cell(row, col).value for col in (2, 3, 4))
        for row in (2, 3)
    }
    assert actual == expected, (actual, expected)
    assert w['Employees']['E2'].value == '=C2*12' and w['Employees']['E3'].value == '=C3*12'
    assert w['Lookup'].sheet_state == 'hidden'
    assert w['Employees'].freeze_panes == 'A2'
    assert w['Employees']['A1'].font.bold and w['Employees']['A1'].fill.fgColor.rgb == '00243B53'
    w.close()
    jobroot = Path(dotenv_values(ROOT / '.env').get('DI_DATA_DIR') or ROOT / 'data') / 'jobs' / job_id
    if not jobroot.is_absolute():
        jobroot = ROOT / jobroot
    for path, digest in before.items():
        assert sha(jobroot / path) == digest, 'Source changed: ' + path
    logs = client.get(f'{base}/api/jobs/{job_id}/logs')
    logs.raise_for_status()
    entries = [json.loads(line) for line in logs.text.splitlines()]
    assert len(entries) == len(events)
    assert 'data:image/png;base64,' not in logs.text
    assert 'iVBORw0KGgo' not in logs.text, 'Image bytes leaked into diagnostic logs'
    (folder / (('question' if question else 'baseline') + '-events.jsonl')).write_text(logs.text, encoding='utf-8')
    stream = client.get(f'{base}/api/jobs/{job_id}/stream', params={'after': entries[-1]['seq']})
    stream.raise_for_status()
    assert 'event: state' in stream.text and job_id in stream.text
    assert job['usage']['requests'] > 0 and job['usage']['input_tokens'] > 0
    return {
        'job_id': job_id,
        'status': job['status'],
        'summary': job['summary'],
        'usage': job['usage'],
        'output': str(target),
        'checks': {
            'employee_values': True,
            'formulas': True,
            'hidden_sheet': True,
            'formatting': True,
            'source_immutability': True,
            'scan_read': True,
            'log_download': True,
            'image_payload_redaction': True,
            'SSE_terminal_state': True,
            'question_resume': answered,
        },
        'events': len(events),
        'tool_errors': [
            {'operation': e['name'], 'message': e['message']}
            for e in events
            if e.get('type') == 'tool' and e.get('status') == 'failed'
        ],
        'tool_operations': sorted(
            {e['name'] for e in events if e.get('type') == 'tool' and not e['name'].endswith('.started')}
        ),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--url', default='http://127.0.0.1:8003')
    parser.add_argument('--scenario', choices=['both', 'baseline', 'question'], default='both')
    args = parser.parse_args()
    env = dotenv_values(ROOT / '.env')
    headers = {}
    if env.get('DI_ACCESS_TOKEN'):
        headers['Authorization'] = 'Bearer ' + env['DI_ACCESS_TOKEN']
    folder = ROOT / 'data/live-tests' / datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')
    folder.mkdir(parents=True)
    report = {'started_at': datetime.now(timezone.utc).isoformat(), 'scenarios': []}
    try:
        with httpx.Client(timeout=45, headers=headers, trust_env=False) as client:
            cfg = client.get(args.url + '/api/config')
            cfg.raise_for_status()
            cfg = cfg.json()
            assert cfg['configured'] and cfg['model'] == 'gpt-6-luna', {
                'configured': cfg['configured'],
                'model': cfg['model'],
            }
            assert cfg['base_url'].rstrip('/') == 'https://api.openai.com/v1', 'Test expects direct OpenAI route'
            report['model'] = cfg['model']
            report['route'] = cfg['base_url']
            packet = create_sample(folder / 'fixtures')
            task = (packet / 'payroll_packet/instructions.txt').read_text()
            if args.scenario in ('both', 'baseline'):
                report['scenarios'].append(run_scenario(client, args.url, packet, task, folder))
            task += '\nAcceptance test: Before writing any output, ask_user to explicitly confirm the October 2026 reporting period, even though it appears in the source. Once the answer arrives, continue and complete the workbook. Read the supporting image in this run too.'
            if args.scenario in ('both', 'question'):
                report['scenarios'].append(run_scenario(client, args.url, packet, task, folder, question=True))
        report['passed'] = True
    except Exception as exc:
        report['passed'] = False
        report['error_type'] = type(exc).__name__
        report['error'] = str(exc)[:3000]
        print('FAILED:', type(exc).__name__, str(exc)[:1000], flush=True)
    report['finished_at'] = datetime.now(timezone.utc).isoformat()
    path = folder / 'report.json'
    path.write_text(json.dumps(report, indent=2, default=str), encoding='utf-8')
    print('Report:', path, flush=True)
    return 0 if report.get('passed') else 1


if __name__ == '__main__':
    raise SystemExit(main())
