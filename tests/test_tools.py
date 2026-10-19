import asyncio
import io
import re
from email.message import EmailMessage
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from pydantic_ai import Agent
from pydantic_ai.messages import (
    ModelRequest,
    ModelResponse,
    TextPart,
    ToolCallPart,
    ToolReturnPart,
    UserPromptPart,
)
from pydantic_ai.models.function import FunctionModel

from app.files import add_file
from app.harness_tools import register_harness_tools
from app.store import Store
from app.tools import JobDeps, job_lock, register_tools

from helpers import run_tool, workbook_bytes


def test_tool_error_is_recoverable_and_output_write_is_atomic(tmp_path):
    store = Store(tmp_path)
    job = store.create('payroll', '2026-10', '', '')
    workbook = Workbook()
    workbook.active.title = 'Employees'
    workbook.active['E2'] = '=C2*12'
    workbook.create_sheet('Lookup').sheet_state = 'hidden'
    buffer = io.BytesIO()
    workbook.save(buffer)
    manifest = []
    add_file(store.path(job['id']), 'template', 'template.xlsx', buffer.getvalue(), manifest)
    store.update(job['id'], files=manifest)
    sequence = [
        ('create_output', {'filename': 'completed.xlsx'}),
        (
            'write_cells',
            {
                'path': 'outputs/completed.xlsx',
                'changes': [
                    {'sheet': 'Employees', 'cell': 'A2', 'value': 'must not be partially saved'},
                    {'sheet': 'Employees', 'cell': 'E2', 'value': 'overwrite formula'},
                ],
            },
        ),
        ('validate_output', {'path': 'outputs/completed.xlsx'}),
        (
            'write_cells',
            {'path': 'outputs/completed.xlsx', 'changes': [{'sheet': 'Employees', 'cell': 'B2', 'value': 'Recovered'}]},
        ),
    ]
    called = []

    def respond(messages, info):
        if len(called) == len(sequence):
            return ModelResponse(parts=[TextPart('Done')])
        name, arguments = sequence[len(called)]
        called.append(name)
        return ModelResponse(parts=[ToolCallPart(name, arguments, tool_call_id=f'tool-{len(called)}')])

    agent = Agent(FunctionModel(respond), deps_type=JobDeps)
    register_tools(agent)
    result = asyncio.run(agent.run('Exercise recovery.', deps=JobDeps(store, job['id'])))
    assert result.output == 'Done'
    output = load_workbook(store.path(job['id']) / 'outputs/completed.xlsx')
    assert output['Employees']['A2'].value is None
    assert output['Employees']['B2'].value == 'Recovered'
    assert output['Employees']['E2'].value == '=C2*12'
    assert output['Lookup'].sheet_state == 'hidden'
    events = store.events(job['id'])
    failed = [event for event in events if event['name'] == 'write_cells' and event['status'] == 'failed']
    assert len(failed) == 1
    assert 'Preserve formula' in failed[0]['message']
    assert any(event['name'] == 'validate_output' and event['result']['passed'] for event in events)


def test_vision_tool_supplies_image_to_model_without_logging_pixels(tmp_path):
    from PIL import Image
    from pydantic_ai import BinaryContent

    store = Store(tmp_path)
    job = store.create('payroll', '2026-10', '', '')
    buffer = io.BytesIO()
    Image.new('RGB', (60, 40), 'white').save(buffer, format='PNG')
    manifest = []
    add_file(store.path(job['id']), 'inputs', 'nested/scan.png', buffer.getvalue(), manifest)
    store.update(job['id'], files=manifest)
    responses = 0
    received_pixels = []

    def respond(messages, info):
        nonlocal responses
        responses += 1
        if responses == 1:
            return ModelResponse(
                parts=[ToolCallPart('view_page', {'path': 'inputs/nested/scan.png'}, tool_call_id='vision-1')]
            )
        for message in messages:
            for part in message.parts:
                if isinstance(part, UserPromptPart) and isinstance(part.content, list):
                    received_pixels.extend(item for item in part.content if isinstance(item, BinaryContent))
        return ModelResponse(parts=[TextPart('Image read')])

    agent = Agent(FunctionModel(respond), deps_type=JobDeps)
    register_tools(agent)
    asyncio.run(agent.run('Read the scan.', deps=JobDeps(store, job['id'])))
    assert received_pixels and received_pixels[0].data == buffer.getvalue()
    assert received_pixels[0].media_type == 'image/png'
    events = store.events(job['id'])
    event = next(event for event in events if event['name'] == 'view_page')
    assert event['result']['image_artifact'].startswith('working/')
    assert 'base64' not in str(event)
    assert 'data' not in event['result']


def test_registered_tools_have_concrete_openai_parameter_schemas():
    """Catch unconstrained Any parameters rejected by the live Responses endpoint."""
    schemas = []

    def capture(messages, info):
        schemas.extend((tool.name, tool.parameters_json_schema) for tool in info.function_tools)
        return ModelResponse(parts=[TextPart('Inspected schemas')])

    agent = Agent(FunctionModel(capture), deps_type=JobDeps)
    register_tools(agent)
    register_harness_tools(agent)
    asyncio.run(agent.run('Inspect schemas only.'))

    def check(schema, path):
        assert any(key in schema for key in ('type', 'anyOf', 'oneOf', 'allOf', '$ref', 'enum')), (
            f'Unconstrained schema: {path}'
        )
        for key, value in schema.get('properties', {}).items():
            check(value, path + '.' + key)
        if 'items' in schema:
            check(schema['items'], path + '[]')
        if isinstance(schema.get('additionalProperties'), dict):
            check(schema['additionalProperties'], path + '.*')
        for key in ('anyOf', 'oneOf', 'allOf'):
            for i, value in enumerate(schema.get(key, [])):
                check(value, path + f'.{key}[{i}]')
        for key, value in schema.get('$defs', {}).items():
            check(value, path + '.$defs.' + key)

    for name, schema in schemas:
        check(schema, name)


def test_read_sheet_returns_coordinates_for_blank_readonly_cells(tmp_path):

    store = Store(tmp_path)
    job = store.create('payroll', '2026-10', '', '')
    wb = Workbook()
    wb.active.title = 'Source'
    wb.active.append(['Employee ID', 'Name'])
    wb.active.append(['EMP001', None])
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    manifest = []
    add_file(store.path(job['id']), 'inputs', 'source.xlsx', buf.getvalue(), manifest)
    store.update(job['id'], files=manifest)
    observed = []

    def respond(messages, info):
        for message in messages:
            if isinstance(message, ModelRequest):
                for part in message.parts:
                    if isinstance(part, ToolReturnPart) and part.tool_name == 'read_sheet':
                        observed.append(part.content)
        if not observed:
            return ModelResponse(
                parts=[
                    ToolCallPart(
                        'read_sheet',
                        {
                            'path': 'inputs/source.xlsx',
                            'sheet': 'Source',
                            'start_row': 1,
                            'end_row': 2,
                            'end_column': 4,
                        },
                        tool_call_id='sparse',
                    )
                ]
            )
        return ModelResponse(parts=[TextPart('Read')])

    agent = Agent(FunctionModel(respond), deps_type=JobDeps)
    register_tools(agent)
    asyncio.run(agent.run('Read the sparse sheet', deps=JobDeps(store, job['id'])))
    assert 'error' not in observed[-1], observed[-1]
    assert observed[-1]['rows'][1][3] == {'cell': 'D2', 'value': None}


def strip_dimension(data):
    src, out = ZipFile(io.BytesIO(data)), io.BytesIO()
    with ZipFile(out, 'w') as dst:
        for item in src.infolist():
            body = src.read(item)
            if item.filename.startswith('xl/worksheets/'):
                body = re.sub(rb'<dimension[^>]*/>', b'', body)
            dst.writestr(item, body)
    return out.getvalue()


def test_read_sheet_handles_workbooks_without_dimension(tmp_path):
    store = Store(tmp_path)
    job = store.create('payroll', '2026-10', '', '')
    wb = Workbook()
    wb.active.title = 'Source'
    wb.active.append(['Employee ID'])
    wb.active.append(['EMP001'])
    buf = io.BytesIO()
    wb.save(buf)
    add_file(store.path(job['id']), 'inputs', 'nodim.xlsx', strip_dimension(buf.getvalue()), [])
    result = run_tool(store, job['id'], 'read_sheet', {'path': 'inputs/nodim.xlsx', 'sheet': 'Source', 'end_column': 1})
    assert 'error' not in result, result
    assert result['sheet_rows'] == 2 and result['rows'][1][0]['value'] == 'EMP001'


def test_email_body_is_read_and_attachments_saved(tmp_path):
    store = Store(tmp_path)
    job = store.create('payroll', '2026-10', '', '')
    msg = EmailMessage()
    msg['From'], msg['Subject'] = 'hr@example.com', 'October changes'
    msg.set_content('EMP001 salary rises to 4500.')
    msg.add_attachment(b'EMP002,5800', maintype='text', subtype='csv', filename='changes.csv')
    add_file(store.path(job['id']), 'inputs', 'mail/changes.eml', msg.as_bytes(), [])
    result = run_tool(store, job['id'], 'read_document', {'path': 'inputs/mail/changes.eml'})
    assert 'Subject: October changes' in result['text'] and '4500' in result['text']
    assert 'base64' not in result['text']
    [attachment] = result['attachments']
    assert (store.path(job['id']) / attachment['path']).read_bytes() == b'EMP002,5800'


def test_extracted_attachments_are_distinct_and_stable(tmp_path):
    store = Store(tmp_path)
    job = store.create('payroll', '2026-10', '', '')
    root = store.path(job['id'])
    expected = {}
    paths = {}
    for department in ('HR', 'Finance'):
        mail = EmailMessage()
        mail.set_content('See attached payroll updates.')
        for index in (1, 2):
            mail.add_attachment(
                f'{department}-{index}'.encode(), maintype='text', subtype='csv', filename='changes.csv'
            )
        source = f'inputs/{department}/changes.eml'
        add_file(root, 'inputs', f'{department}/changes.eml', mail.as_bytes(), [])
        result = run_tool(store, job['id'], 'read_document', {'path': source})
        paths[source] = [item['path'] for item in result['attachments']]
        for index, path in enumerate(paths[source], 1):
            expected[path] = f'{department}-{index}'.encode()
    assert len(expected) == 4
    for path, content in expected.items():
        assert (root / path).read_bytes() == content
    for source, original_paths in paths.items():
        repeated = run_tool(store, job['id'], 'read_document', {'path': source})
        assert [item['path'] for item in repeated['attachments']] == original_paths


def test_validation_fails_when_template_features_are_lost(tmp_path):
    store = Store(tmp_path)
    job = store.create('payroll', '2026-10', '', '')
    out = io.BytesIO()
    with ZipFile(io.BytesIO(workbook_bytes())) as src, ZipFile(out, 'w') as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item))
        dst.writestr('xl/slicers/slicer1.xml', '<slicers/>')
    manifest = []
    add_file(store.path(job['id']), 'template', 'template.xlsx', out.getvalue(), manifest)
    store.update(job['id'], files=manifest)
    run_tool(store, job['id'], 'create_output', {'filename': 'out.xlsx'})
    run_tool(
        store,
        job['id'],
        'write_cells',
        {'path': 'outputs/out.xlsx', 'changes': [{'sheet': 'Employees', 'cell': 'A2', 'value': 'EMP001'}]},
    )
    result = run_tool(store, job['id'], 'validate_output', {'path': 'outputs/out.xlsx'})
    assert not result['passed']
    assert any('slicers' in f for f in result['failures'])


def chart_template(include_shape):
    wb = Workbook()
    sh = wb.active
    sh.title = 'Employees'
    for row in [('Employee', 'Salary'), ('EMP001', 4200)]:
        sh.append(row)
    chart = BarChart()
    chart.add_data(Reference(sh, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    sh.add_chart(chart, 'D1')
    data = io.BytesIO()
    wb.save(data)
    wb.close()
    if not include_shape:
        return data.getvalue()
    # Add a real DrawingML shape alongside the chart in the same drawing part.
    ns = 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'
    drawing = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    shape = ElementTree.fromstring(f'''<xdr:twoCellAnchor xmlns:xdr="{ns}" xmlns:a="{drawing}">
      <xdr:from><xdr:col>0</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>6</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>
      <xdr:to><xdr:col>2</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>9</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>
      <xdr:sp><xdr:nvSpPr><xdr:cNvPr id="99" name="Payroll note"/><xdr:cNvSpPr/></xdr:nvSpPr>
      <xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr>
      <xdr:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>Payroll note</a:t></a:r></a:p></xdr:txBody></xdr:sp>
      <xdr:clientData/></xdr:twoCellAnchor>''')
    result = io.BytesIO()
    with ZipFile(io.BytesIO(data.getvalue())) as src, ZipFile(result, 'w') as dst:
        for name in src.namelist():
            body = src.read(name)
            if name == 'xl/drawings/drawing1.xml':
                xml = ElementTree.fromstring(body)
                xml.append(shape)
                body = ElementTree.tostring(xml)
            dst.writestr(name, body)
    return result.getvalue()


@pytest.mark.parametrize('include_shape', [False, True])
def test_validation_detects_shape_loss_when_chart_survives(tmp_path, include_shape):
    store = Store(tmp_path)
    job = store.create('payroll', '2026-10', '', '')
    root = store.path(job['id'])
    manifest = []
    add_file(root, 'template', 'template.xlsx', chart_template(include_shape), manifest)
    store.update(job['id'], files=manifest)
    run_tool(store, job['id'], 'create_output', {'filename': 'out.xlsx'})
    # Untouched template copies preserve every feature, including shapes.
    assert run_tool(store, job['id'], 'validate_output', {'path': 'outputs/out.xlsx'})['passed']
    run_tool(
        store,
        job['id'],
        'write_cells',
        {'path': 'outputs/out.xlsx', 'changes': [{'sheet': 'Employees', 'cell': 'B2', 'value': 4500}]},
    )
    with ZipFile(root / 'outputs/out.xlsx') as saved:
        assert 'xl/charts/chart1.xml' in saved.namelist()
    result = run_tool(store, job['id'], 'validate_output', {'path': 'outputs/out.xlsx'})
    assert result['passed'] is not include_shape
    if include_shape:
        assert any('shapes' in failure for failure in result['failures'])


def test_tools_for_one_job_share_a_lock():
    assert job_lock('a' * 32) is job_lock('a' * 32)
    assert job_lock('a' * 32) is not job_lock('b' * 32)
