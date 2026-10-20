"""Regression coverage for output multiplicity, revised scope, bounded search and worker progress."""

import asyncio
from dataclasses import replace
from datetime import datetime, timezone
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pydantic_ai.messages import ModelResponse, ToolCallPart
from pydantic_ai.models.function import FunctionModel
from pydantic_ai.usage import RunUsage

from app import regions
from app.datasets import (
    Column,
    Destination,
    DestinationColumn,
    Predicate,
    Transformation,
    extract_dataset,
    populate_template,
    recover_operation,
    transform_dataset,
)
from app.harness_tools import validate_job
from app.inventory import inspect_source
from app.progress import observe
from app.regions import RegionRole, save_role, search
from app.requirements import OutputSpecification, check_requirements, save_requirements
from app.runtime import Runner
from app.state import JobState, identity
from app.supervisor import Supervisor
from app.tools import assign_cells, validate_workbook

from goal_helpers import establish_requirements, messy_case
from helpers import run_tool


def payment_case(tmp_path, prefilled=False):
    store, job, deps = messy_case(tmp_path, prefilled=prefilled)
    state = deps.state
    inspect_source(state, 'inputs/nested/October.xlsx')
    output = establish_requirements(deps)
    for key, region in state.all('region'):
        if region['path'].startswith('template/'):
            continue
        evidence = observe(deps, 'inspect_regions', {'path': region['path'], 'region': key}, {})[0]['key']
        save_role(
            state,
            RegionRole(
                region=key,
                role='records' if region['bounds'] == 'A14:C16' else 'reference',
                rationale='Current payments; remaining regions provide context',
                evidence=[evidence],
            ),
        )
    region = next(key for key, item in state.all('region_role') if item['role'] == 'records')
    dataset = extract_payments(state, region)
    return store, job, deps, output, region, dataset


def extract_payments(state, region, identifier_kind='identifier'):
    return extract_dataset(
        state,
        'inputs/nested/October.xlsx',
        'Payments',
        14,
        16,
        [
            Column(name='id', column=1, kind=identifier_kind),
            Column(name='amount', column=2, kind='decimal'),
            Column(name='code', column=3, kind='text'),
        ],
        [14],
        region_id=region,
    )


def payment_destination():
    return Destination(
        sheet='Output',
        start_row=2,
        end_row=3,
        columns=[
            DestinationColumn(field='id', column=1),
            DestinationColumn(field='amount', column=2),
            DestinationColumn(field='code', column=3),
        ],
        required=['id', 'amount', 'code'],
    )


def test_source_backed_records_matching_prefilled_template_pass_validation(tmp_path):
    _, _, deps, path, _, dataset = payment_case(tmp_path, prefilled=True)
    specification = deps.state.get('requirement', identity(path))
    # The prefilled data makes the detected block span A1:D3, including blank formula-column cells.
    next(area for area in specification['areas'] if area['id'] == 'formula')['bounds'] = 'D1:D3'
    save_requirements(deps, OutputSpecification.model_validate(specification))
    assert 'Too few output records: payments' in check_requirements(deps.state, path)
    result = populate_template(deps.state, dataset['id'], path, payment_destination(), 'matching-records')
    assert result['records_written'] == 2
    assert deps.state.coverage()['passed']
    assert validate_workbook(deps, path)['passed']
    checked = validate_job(deps, path)
    assert checked['passed'], checked['failures']


def unpivot_case(tmp_path):
    store, job, deps, path, _, dataset = payment_case(tmp_path)
    specification = deps.state.get('requirement', identity(path))
    specification['areas'][0]['bounds'] = 'A2:C5'
    specification['areas'][0]['fields'][2]['allowed_values'] = []
    save_requirements(deps, OutputSpecification.model_validate(specification))
    transformed = transform_dataset(
        deps.state,
        dataset['id'],
        [
            Transformation(operation='unpivot', fields=['amount', 'code']),
            Transformation(operation='convert', field='value', kind='decimal'),
        ],
    )
    destination = payment_destination()
    destination.end_row = 5
    destination.columns[1].field = 'value'
    destination.columns[2].field = 'category'
    destination.required = ['id', 'value']
    return store, job, deps, path, transformed, destination


def test_unpivot_rejected_siblings_block_coverage_and_final_validation_until_rewritten(tmp_path):
    _, _, deps, path, dataset, destination = unpivot_case(tmp_path)
    state = deps.state
    result = populate_template(state, dataset['id'], path, destination, 'partial')
    assert result['records_written'] == 2 and result['unresolved'] == 2
    assert validate_workbook(deps, path)['passed']
    coverage = JobState(deps.root).coverage()
    assert not coverage['passed']
    assert sum('Unresolved output record' in item['reason'] for item in coverage['issues']) == 2
    assert not validate_job(deps, path)['passed']

    # An explicit transformation removes irrelevant derived rows without excluding their shared sources.
    revised = transform_dataset(
        state,
        dataset['id'],
        [
            Transformation(
                operation='filter',
                predicate=Predicate(field='category', value='amount'),
                reason='Codes are reference data, not monetary payments',
            )
        ],
    )
    assert populate_template(state, revised['id'], path, destination, 'corrected')['unresolved'] == 0
    assert state.coverage()['passed']
    assert validate_workbook(deps, path)['passed']
    assert validate_job(deps, path)['passed']


def test_recovery_persists_rejected_output_record_obligations(tmp_path, monkeypatch):
    _, _, deps, path, dataset, destination = unpivot_case(tmp_path)
    original = deps.state.put_many

    def interrupted(kind, entries):
        if kind == 'assignment':
            raise OSError('Interrupted after workbook replacement')
        return original(kind, entries)

    monkeypatch.setattr(deps.state, 'put_many', interrupted)
    with pytest.raises(OSError, match='Interrupted'):
        populate_template(deps.state, dataset['id'], path, destination, 'interrupted')
    restored = JobState(deps.root)
    assert recover_operation(restored, 'interrupted')['status'] == 'completed'
    assert not restored.coverage()['passed']
    assert sum(bool(item['issues']) for _, item in restored.all('output_record')) == 2


@pytest.mark.parametrize('status', ['prepared', 'completed'])
def test_older_write_journals_restore_rejected_siblings_without_replaying_completed_assignments(tmp_path, status):
    _, _, deps, path, dataset, destination = unpivot_case(tmp_path)
    state = deps.state
    populate_template(state, dataset['id'], path, destination, 'saved')
    # Pre-fix journals stored assignments and rejections, with no per-output-row ledger.
    state.delete_many('output_record', [key for key, _ in state.all('output_record')])
    state.delete_many('output_accounting', [path])
    journal = state.get('operation', 'saved')
    journal['status'] = status
    state.put('operation', 'saved', journal)
    if status == 'prepared':
        state.delete_many('assignment', [key for key, _ in state.all('assignment')])
        assert recover_operation(state, 'saved')['status'] == 'completed'
    else:
        before = state.all('assignment')
        assert populate_template(state, dataset['id'], path, destination, 'saved')['unresolved'] == 2
        assert state.all('assignment') == before
    assert not state.coverage()['passed']
    assert sum(bool(item['issues']) for _, item in state.all('output_record')) == 2


def test_repeating_a_completed_write_keeps_validation_current(tmp_path):
    _, _, deps, path, _, dataset = payment_case(tmp_path)
    state = deps.state
    populate_template(state, dataset['id'], path, payment_destination(), 'payments')
    assert validate_workbook(deps, path)['passed']
    checked = validate_job(deps, path)
    assert checked['passed']
    populate_template(state, dataset['id'], path, payment_destination(), 'payments')
    assert state.revision() == checked['revision']


@pytest.mark.parametrize('legacy', [False, True])
def test_retrying_old_identical_bytes_does_not_resurrect_replaced_output_obligations(tmp_path, monkeypatch, legacy):
    # XLSX saves stamp wall-clock times; freeze both clocks so this collision is deterministic.
    fixed = datetime(2026, 10, 19, tzinfo=timezone.utc)
    monkeypatch.setattr(
        'openpyxl.writer.excel.datetime',
        SimpleNamespace(datetime=SimpleNamespace(now=lambda **kwargs: fixed), timezone=timezone),
    )
    monkeypatch.setattr('zipfile.time.localtime', lambda *args: fixed.timetuple())
    _, _, deps, path, _, dataset = payment_case(tmp_path)
    state = deps.state
    # Real templates commonly style blank input cells; retain their XML in both empty writes.
    workbook = load_workbook(deps.root / path)
    for row in workbook['Output'].iter_rows(min_row=2, max_row=3, min_col=1, max_col=3):
        for cell in row:
            cell.number_format = '@'
    workbook.save(deps.root / path)
    workbook.close()
    destination = payment_destination()
    destination.allowed = {'code': ['OTHER']}
    rejected = populate_template(state, dataset['id'], path, destination, 'rejected')
    assert rejected['unresolved'] == 2
    if legacy:
        state.delete_many('output_record', [key for key, _ in state.all('output_record')])
        state.delete_many('output_accounting', [path])
    excluded = transform_dataset(
        state,
        dataset['id'],
        [Transformation(operation='filter', predicate=Predicate(field='code', value='OTHER'), reason='Outside scope')],
    )
    replaced = populate_template(state, excluded['id'], path, destination, 'excluded')
    assert replaced['sha256'] == rejected['sha256']
    assert state.coverage()['passed']
    revision = state.revision()
    populate_template(state, dataset['id'], path, destination, 'rejected')
    assert state.coverage()['passed']
    assert state.revision() == revision


def test_partial_overwrite_cannot_erase_a_rejected_sibling(tmp_path):
    _, _, deps, path, dataset, destination = unpivot_case(tmp_path)
    populate_template(deps.state, dataset['id'], path, destination, 'partial')
    single_field = Destination(
        sheet='Output', start_row=3, end_row=3, columns=[DestinationColumn(field='value', column=2)]
    )
    selected = transform_dataset(
        deps.state,
        dataset['id'],
        [
            Transformation(operation='filter', predicate=Predicate(field='id', value='00101'), reason='Other employee'),
            Transformation(operation='filter', predicate=Predicate(field='category', value='amount'), reason='Code'),
        ],
    )
    with pytest.raises(ValueError, match='Partial-column overwrite'):
        populate_template(deps.state, selected['id'], path, single_field, 'bypass')
    assert not deps.state.coverage()['passed']


@pytest.mark.parametrize('role', ['outside_scope', 'instructions', 'reference'])
def test_reclassification_retires_unassigned_record_obligations(tmp_path, role):
    _, _, deps, _, region, _ = payment_case(tmp_path)
    old = deps.state.get('region_role', region)
    save_role(
        deps.state, RegionRole.model_validate(old | {'role': role, 'rationale': 'Corrected region interpretation'})
    )
    assert deps.state.coverage()['records']['pending'] == 0
    assert deps.state.coverage()['passed']


def test_mapping_revision_retires_old_rows_but_requires_replacing_stale_output(tmp_path):
    _, _, deps, path, region, first = payment_case(tmp_path)
    state = deps.state
    populate_template(state, first['id'], path, payment_destination(), 'original')
    revised = extract_payments(state, region, identifier_kind='text')
    assert revised['id'] != first['id']
    assert state.coverage()['records']['pending'] == 2
    assert any('obsolete region interpretation' in item['reason'] for item in state.coverage()['issues'])
    populate_template(state, revised['id'], path, payment_destination(), 'revised')
    assert state.coverage()['records'] == {'pending': 0, 'included': 2, 'excluded': 0, 'unresolved': 0}
    assert validate_workbook(deps, path)['passed']
    assert validate_job(deps, path)['passed']


def test_stale_individual_cell_contributions_are_not_hidden_by_retiring_observations(tmp_path):
    _, _, deps, path, region, first = payment_case(tmp_path)
    record = deps.state.get('rows', first['id'])[0]['id']
    assign_cells(deps.state, path, [{'sheet': 'Output', 'cell': 'A2', 'value': '00101', 'record': record}])
    with pytest.raises(ValueError, match='output contributions'):
        previous = deps.state.get('region_role', region)
        save_role(deps.state, RegionRole.model_validate(previous | {'role': 'outside_scope'}))
    extract_payments(deps.state, region, identifier_kind='text')
    assert any('obsolete region interpretation' in item['reason'] for item in deps.state.coverage()['issues'])


def test_reference_join_contributions_remain_valid_until_their_mapping_changes(tmp_path):
    _, _, deps, path, _, dataset = payment_case(tmp_path)
    state = deps.state
    reference = next(key for key, item in state.all('region') if item['bounds'] == 'G14:H16')

    def extract_codes(kind):
        return extract_dataset(
            state,
            'inputs/nested/October.xlsx',
            'Payments',
            14,
            16,
            [Column(name='code', column=7), Column(name='meaning', column=8, kind=kind)],
            [14],
            region_id=reference,
        )

    def joined(codes):
        return transform_dataset(
            state,
            dataset['id'],
            [Transformation(operation='join', dataset=codes['id'], fields=['code'], right_keys=['code'])],
        )

    original = joined(extract_codes('raw'))
    populate_template(state, original['id'], path, payment_destination(), 'joined')
    assert state.coverage()['passed']
    assert state.coverage()['records']['included'] == 2
    assert validate_workbook(deps, path)['passed']
    assert validate_job(deps, path)['passed']
    revised = joined(extract_codes('text'))
    assert any('obsolete region interpretation' in item['reason'] for item in state.coverage()['issues'])
    populate_template(state, revised['id'], path, payment_destination(), 'rejoined')
    assert state.coverage()['passed']
    assert validate_workbook(deps, path)['passed']
    assert validate_job(deps, path)['passed']


def test_superseded_template_region_does_not_block_split_requirements(tmp_path):
    store, job, deps, path, _, dataset = payment_case(tmp_path)
    populate_template(deps.state, dataset['id'], path, payment_destination(), 'payments')
    assert check_requirements(deps.state, path) == []
    parent = next(
        key
        for key, item in deps.state.all('region')
        if item['path'] == 'template/destination.xlsx' and item['bounds'] == 'A1:C1'
    )
    for bounds, supersede in [('A1:A1', [parent]), ('B1:C1', [])]:
        assert 'error' not in run_tool(
            store,
            job['id'],
            'revise_region',
            {'region': parent, 'bounds': bounds, 'supersede': supersede, 'reason': 'Classify separate header fields'},
        )
    specification = deps.state.get('requirement', identity(path))
    header = specification['areas'].pop(1)
    specification['areas'].extend(
        [header | {'id': 'id-header', 'bounds': 'A1:A1'}, header | {'id': 'other-headers', 'bounds': 'B1:C1'}]
    )
    save_requirements(deps, OutputSpecification.model_validate(specification))
    assert check_requirements(deps.state, path) == []
    assert validate_workbook(deps, path)['passed']
    assert validate_job(deps, path)['passed']


def test_common_search_is_linear_and_paginates_before_building_neighborhoods(tmp_path, monkeypatch):
    _, _, deps = messy_case(tmp_path)
    inspect_source(deps.state, 'inputs/nested/October.xlsx')
    accesses, size = 0, 4000

    class IndexedCell(dict):
        def __getitem__(self, key):
            nonlocal accesses
            accesses += 1
            assert accesses < size * 20, 'Search scanned the full index once per match'
            return super().__getitem__(key)

    cells = [IndexedCell(cell=f'A{row}', row=row, column=1, value='Payment') for row in range(1, size + 1)]
    indexed_cells = regions.indexed_cells

    def indexed(state, source):
        if state.get('inventory', source)['sheet'] == 'Payments':
            return cells
        return indexed_cells(state, source)

    monkeypatch.setattr(regions, 'indexed_cells', indexed)
    result = search(deps.state, 'inputs/nested/October.xlsx', 'payment', sheet='Payments', offset=2000, limit=1)
    assert result['total'] == size and result['next_offset'] == 2001
    assert result['items'][0]['cell'] == 'A2001'
    assert [cell['cell'] for cell in result['items'][0]['nearby']] == ['A2000', 'A2001', 'A2002']
    with pytest.raises(ValueError, match='nonnegative offset'):
        search(deps.state, 'inputs/nested/October.xlsx', 'payment', offset=-1)


def test_search_pages_across_sheets_without_losing_matches(tmp_path):
    _, _, deps = messy_case(tmp_path)
    path = 'inputs/nested/October.xlsx'
    inspect_source(deps.state, path)
    expected = search(deps.state, path, 'reg', limit=100)
    assert {item['sheet'] for item in expected['items']} == {'Payments', 'Codes'}
    actual, offset = [], 0
    while offset is not None:
        page = search(deps.state, path, 'reg', offset=offset, limit=1)
        assert page['total'] == expected['total'] and page['complete_index']
        actual.extend(page['items'])
        offset = page['next_offset']
    assert actual == expected['items']


@pytest.mark.parametrize('repeat', [False, True])
def test_worker_credits_scoped_reads_but_still_stops_repetition(tmp_path, repeat):
    store, job, deps = messy_case(tmp_path)
    deps.settings = replace(deps.settings, subagent_requests=20)
    calls = []

    def respond(messages, info):
        names = {tool.name for tool in info.function_tools}
        assert not {'write_cells', 'update_work_plan', 'save_interpretation'}.intersection(names)
        number = len(calls)
        calls.append(number)
        if number < 10:
            row = 14 if repeat else number + 14
            return ModelResponse(
                parts=[
                    ToolCallPart(
                        'read_sheet',
                        {
                            'path': 'inputs/nested/October.xlsx',
                            'sheet': 'Payments',
                            'start_row': row,
                            'end_row': row,
                            'start_column': 1,
                            'end_column': 3,
                        },
                        tool_call_id=f'read-{number}',
                    )
                ]
            )
        return ModelResponse(
            parts=[
                ToolCallPart('final_result', {'summary': 'Inspected ten assigned source ranges'}, tool_call_id='done')
            ]
        )

    runner = Runner(store, deps.settings, FunctionModel(respond))

    async def investigate():
        runner.delegation.begin(job['id'], RunUsage())
        task = await runner.delegation.delegate(
            deps, 'Inspect the payment sections', ['inputs/nested/October.xlsx'], 'scout'
        )
        await runner.delegation.tasks[(job['id'], task['id'])]
        result = deps.state.get('task', task['id'])
        if repeat:
            assert result['status'] == 'cancelled'
            assert len(calls) < 10
        else:
            assert result['status'] == 'completed', result
            assert result['handover']['summary'] == 'Inspected ten assigned source ranges'
            assert len(calls) == 11
            assert deps.state.get('supervision', task['id'])['recovery'] is None
            assert Supervisor(deps.state).turn()[1]['last_useful'] == 0
        assert not deps.state.all('assignment')

    asyncio.run(investigate())
