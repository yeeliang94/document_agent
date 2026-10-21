"""Coverage, bulk execution, bounded delegation and context recovery contracts."""

import asyncio
import hashlib
import io
import json
import threading
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pydantic_ai.messages import (
    ModelRequest,
    ModelResponse,
    RetryPromptPart,
    TextPart,
    ToolCallPart,
    ToolReturnPart,
    UserPromptPart,
)
from pydantic_ai.models import ModelRequestContext, ModelRequestParameters
from pydantic_ai.models.function import FunctionModel
from pydantic_ai.usage import RunUsage

from app.config import Settings
from app.context import ContextBudgetError, ContextPolicy, request_retry, safe_cut
from app.datasets import (
    Column,
    Destination,
    DestinationColumn,
    Transformation,
    dataset_rows,
    extract_dataset,
    populate_template,
    recover_operation,
    transform_dataset,
)
from app.files import add_file
from app.harness_tools import validate_job
from app.inventory import inspect_source
from app.main import create_app
from app.runtime import Runner
from app.state import JobState
from app.store import Store
from app.tools import JobDeps

from helpers import LOCAL, create_job, final, run_tool, scripted, wait_for_job


def fixture_job(tmp_path, count=3, width=3, dirty=False):
    store = Store(tmp_path)
    job = store.create('payroll', '2026-10', 'Copy source rows into the template', '')
    source = Workbook()
    sheet = source.active
    sheet.title = 'Data'
    sheet.append(['ID', 'Amount', *[f'Field {i}' for i in range(3, width + 1)]])
    for number in range(1, count + 1):
        sheet.append([number, number * 10, *[f'Value {number}/{column}' for column in range(3, width + 1)]])
        sheet.cell(number + 1, 1).number_format = '000000'
    if dirty:
        sheet.insert_rows(3)
        for column, value in enumerate(['ID', 'Amount', 'Field 3'], 1):
            sheet.cell(3, column, value)
        sheet.append(['Grand Total', 60, None])
    buffer = io.BytesIO()
    source.save(buffer)
    source.close()
    template = Workbook()
    template.active.title = 'Output'
    template.active.cell(2, width + 1, '=1+1')
    template.create_sheet('Lookup').sheet_state = 'hidden'
    destination = io.BytesIO()
    template.save(destination)
    template.close()
    manifest = []
    root = store.path(job['id'])
    add_file(root, 'inputs', 'nested/source.xlsx', buffer.getvalue(), manifest)
    add_file(root, 'template', 'template.xlsx', destination.getvalue(), manifest)
    store.update(job['id'], files=manifest)
    deps = JobDeps(store, job['id'], Settings(data_dir=tmp_path, max_subagents=2))
    state = deps.state
    return store, job, deps, state


def extract(state, count=3, width=3):
    inspect_source(state, 'inputs/nested/source.xlsx')
    columns = [Column(name='id', column=1, kind='identifier'), Column(name='amount', column=2, kind='decimal')]
    columns.extend(Column(name=f'f{i}', column=i) for i in range(3, width + 1))
    return extract_dataset(state, 'inputs/nested/source.xlsx', 'Data', 1, count + 1, columns, [1])


def destination(count=3, width=3):
    return Destination(
        sheet='Output',
        start_row=2,
        end_row=count + 1,
        columns=[DestinationColumn(field='id', column=1), DestinationColumn(field='amount', column=2)]
        + [DestinationColumn(field=f'f{i}', column=i) for i in range(3, width + 1)],
        required=['id', 'amount'],
        unique=['id'],
    )


def output_copy(deps):
    path = deps.root / 'outputs/out.xlsx'
    path.write_bytes((deps.root / 'template/template.xlsx').read_bytes())
    return 'outputs/out.xlsx'


def test_large_workbook_bulk_processing_preserves_types_formulas_and_all_rows(tmp_path):
    store, job, deps, state = fixture_job(tmp_path, count=5000, width=30)
    dataset = extract(state, count=5000, width=30)
    path = output_copy(deps)
    result = populate_template(state, dataset['id'], path, destination(5000, 30), 'bulk-1')
    assert result['records_written'] == 5000
    assert state.coverage()['records'] == {'pending': 0, 'included': 5000, 'excluded': 1, 'unresolved': 0}
    assert state.coverage()['passed']
    workbook = load_workbook(deps.root / path)
    assert workbook['Output']['A5001'].value == '005000'
    assert workbook['Output']['B5001'].value == 50000
    assert workbook['Output']['B5001'].data_type == 'n'
    assert workbook['Output']['AD5001'].value == 'Value 5000/30'
    assert workbook['Output']['AE2'].value == '=1+1'
    assert workbook['Lookup'].sheet_state == 'hidden'
    workbook.close()
    before = hashlib.sha256((deps.root / path).read_bytes()).hexdigest()
    assert populate_template(state, dataset['id'], path, destination(5000, 30), 'bulk-1') == result
    assert hashlib.sha256((deps.root / path).read_bytes()).hexdigest() == before
    assert run_tool(store, job['id'], 'validate_output', {'path': path})['passed']
    assert validate_job(deps, path)['passed']


def test_partial_discovery_requires_continuation_and_reinspection_is_safe(tmp_path):
    _, _, _, state = fixture_job(tmp_path)
    first = inspect_source(state, 'inputs/nested/source.xlsx', max_cells=3)
    assert not first['sheets'][0]['scan_complete']
    assert not state.coverage()['passed']
    with pytest.raises(ValueError, match='skipping discovery'):
        inspect_source(state, 'inputs/nested/source.xlsx', start_row=4)
    next_row = first['sheets'][0]['next_row']
    finished = inspect_source(state, 'inputs/nested/source.xlsx', start_row=next_row)
    assert finished['sheets'][0]['scan_complete']
    assert inspect_source(state, 'inputs/nested/source.xlsx')['sheets'][0]['scan_complete']


def test_scanned_rows_outside_extracted_region_remain_pending(tmp_path):
    _, _, deps, state = fixture_job(tmp_path)
    dataset = extract(state, count=2)
    populate_template(state, dataset['id'], output_copy(deps), destination(), 'partial')
    coverage = state.coverage()
    assert coverage['records']['pending'] == 1
    assert not coverage['passed']


def test_conflict_on_assigned_record_blocks_coverage_until_resolved_and_rewritten(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    path = output_copy(deps)
    populate_template(state, dataset['id'], path, destination(), 'first')
    _, rows = dataset_rows(state, dataset['id'])
    record = rows[0]['id']
    state.put('decision', record, {'status': 'unresolved', 'reason': 'Amount conflict'})
    assert state.coverage()['records']['unresolved'] == 1
    revised = run_tool(
        store,
        job['id'],
        'resolve_records',
        {
            'dataset': dataset['id'],
            'resolutions': [
                {
                    'record_id': record,
                    'values': {'amount': '12.50'},
                    'reason': 'Confirmed correction',
                    'evidence': 'Authoritative payroll instruction page 1',
                }
            ],
        },
    )
    assert not state.coverage()['passed']
    populate_template(state, revised['id'], path, destination(), 'corrected')
    assert state.coverage()['passed']
    assert load_workbook(deps.root / path)['Output']['B2'].value == 12.5


def test_filter_exclusions_commit_only_with_output_and_numeric_rename_stays_numeric(tmp_path):
    _, _, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    revised = transform_dataset(
        state,
        dataset['id'],
        [
            Transformation(operation='rename', field='amount', target='pay'),
            Transformation(
                operation='filter',
                predicate={'field': 'pay', 'operator': 'gt', 'value': 10},
                reason='Explicit period rule',
            ),
        ],
    )
    assert state.coverage()['records']['pending'] == 3
    target = destination()
    target.columns[1].field = 'pay'
    target.required = ['id', 'pay']
    path = output_copy(deps)
    populate_template(state, revised['id'], path, target, 'filtered')
    assert state.coverage()['passed']
    workbook = load_workbook(deps.root / path)
    assert workbook['Output']['B2'].value == 20
    assert workbook['Output']['B2'].data_type == 'n'


def test_ambiguous_join_is_reported_and_cannot_be_written(tmp_path):
    _, _, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    _, rows = dataset_rows(state, dataset['id'])
    state.put('dataset', 'ambiguous', {'id': 'ambiguous', 'rows': 2, 'columns': ['id'], 'kinds': {}})
    state.put('rows', 'ambiguous', [rows[0], rows[0]])
    revised = transform_dataset(
        state, dataset['id'], [Transformation(operation='join', dataset='ambiguous', fields=['id'], right_keys=['id'])]
    )
    assert revised['exceptions'] == 3
    result = populate_template(state, revised['id'], output_copy(deps), destination(), 'joined')
    assert result['records_written'] == 0
    assert not state.coverage()['passed']


def test_operation_recovers_after_file_commit_before_assignments(tmp_path, monkeypatch):
    _, _, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    path = output_copy(deps)
    put = state.put_many

    def fail_assignment(kind, entries):
        if kind == 'assignment':
            raise OSError('Simulated crash after output replacement')
        return put(kind, entries)

    monkeypatch.setattr(state, 'put_many', fail_assignment)
    with pytest.raises(OSError, match='Simulated crash'):
        populate_template(state, dataset['id'], path, destination(), 'recover')
    recovered = JobState(deps.root)
    assert recovered.get('operation', 'recover')['status'] == 'prepared'
    assert populate_template(recovered, dataset['id'], path, destination(), 'recover')['records_written'] == 3
    assert recovered.get('mapping', 'recover')
    assert recovered.coverage()['passed']
    with pytest.raises(ValueError, match='different arguments'):
        populate_template(
            recovered,
            dataset['id'],
            path,
            Destination(sheet='Output', start_row=3, end_row=5, columns=destination().columns),
            'recover',
        )


def test_instruction_invalidates_only_the_datasets_it_names(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    path = output_copy(deps)
    populate_template(state, dataset['id'], path, destination(), 'old')
    run_tool(store, job['id'], 'validate_output', {'path': path})
    assert validate_job(deps, path)['passed']
    state.instruction('Thanks')
    # An unapplied instruction blocks completion until the agent settles it.
    assert any('not been applied' in failure for failure in validate_job(deps, path)['failures'])
    state.apply_instructions([1], 'Acknowledgement only')
    run_tool(store, job['id'], 'validate_output', {'path': path})
    assert validate_job(deps, path)['passed']
    state.instruction('Exclude employee 000001')
    state.apply_instructions([2], 'Source rows change', invalidate=[dataset['id']])
    with pytest.raises(ValueError, match='invalidated'):
        populate_template(state, dataset['id'], path, destination(), 'stale')
    with pytest.raises(ValueError, match='invalidated'):
        transform_dataset(state, dataset['id'], [Transformation(operation='select', fields=['employee_id'])])
    assert not validate_job(deps, path)['passed']
    assert state.revision('instruction_revision') == 0
    with pytest.raises(ValueError, match='pending instruction'):
        state.apply_instructions([2], 'Already applied')


def test_reset_all_instruction_invalidates_every_earlier_result(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    path = output_copy(deps)
    populate_template(state, dataset['id'], path, destination(), 'old')
    state.instruction('Use a different payroll month')
    state.apply_instructions([1], 'Every mapping depends on the month', reset_all=True)
    with pytest.raises(ValueError, match='Instructions changed'):
        populate_template(state, dataset['id'], path, destination(), 'stale')
    assert not validate_job(deps, path)['passed']
    assert extract(state)['id'] != dataset['id']


def test_bounded_results_remain_retrievable_and_sources_cannot_be_overwritten(tmp_path):
    _, _, deps, _ = fixture_job(tmp_path)
    deps.settings = Settings(data_dir=tmp_path, tool_result_bytes=1024)
    result = asyncio.run(deps.action('large', {}, lambda: {'tail': 'x' * 5000 + 'END'}))
    assert len(json.dumps(result).encode()) < 1024
    assert json.loads((deps.root / result['artifact']).read_text())['tail'].endswith('END')
    with pytest.raises(ValueError, match='Only output'):
        deps.resolve('inputs/nested/source.xlsx', output=True)


def test_context_cut_handles_tool_retries_and_does_not_orphan_results():
    history = [
        ModelRequest(parts=[UserPromptPart('task')]),
        ModelResponse(parts=[ToolCallPart('read_document', {}, tool_call_id='bad')]),
        ModelRequest(parts=[RetryPromptPart('Invalid arguments', tool_name='read_document', tool_call_id='bad')]),
    ]
    for index in range(10):
        history.extend(
            [
                ModelResponse(parts=[ToolCallPart('read_document', {}, tool_call_id=str(index))]),
                ModelRequest(parts=[ToolReturnPart('read_document', 'ok', tool_call_id=str(index))]),
            ]
        )
    assert safe_cut(history, keep=1) >= 20
    history.append(ModelResponse(parts=[ToolCallPart('read_document', {}, tool_call_id='pending')]))
    assert safe_cut(history, keep=0) < len(history)


def test_delegated_scout_has_separate_history_readonly_tools_and_scoped_access(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)
    calls = []

    def respond(messages, info):
        names = {tool.name for tool in info.function_tools}
        assert 'delegate_task' not in names and 'write_cells' not in names
        assert 'inspect_source' in names
        calls.append(1)
        if len(calls) == 1:
            return ModelResponse(
                parts=[ToolCallPart('inspect_source', {'path': 'inputs/nested/source.xlsx'}, tool_call_id='scan')]
            )
        return ModelResponse(
            parts=[
                ToolCallPart(
                    'final_result',
                    {
                        'summary': 'Mapped Data sheet',
                        'evidence': ['inputs/nested/source.xlsx:Data'],
                        'uncertainties': [],
                        'coverage': 'assigned_scope_inspected',
                    },
                    tool_call_id='done',
                )
            ]
        )

    runner = Runner(store, deps.settings, FunctionModel(respond))

    async def run():
        runner.delegation.begin(job['id'], RunUsage())
        delegated = await runner.delegation.delegate(deps, 'Map the workbook', ['inputs/nested/source.xlsx'], 'scout')
        key = delegated['id']
        await runner.delegation.tasks[(job['id'], key)]
        result = await runner.delegation.status(deps, key)
        assert result['status'] == 'completed', result
        assert len(calls) == 2
        assert runner.delegation.usage[job['id']].requests == 2
        child = JobDeps(store, job['id'], deps.settings, key, ('inputs/nested/source.xlsx',))
        with pytest.raises(ValueError, match='assigned source'):
            child.resolve('template/template.xlsx')
        with pytest.raises(ValueError, match='assigned source'):
            child.resolve('outputs/out.xlsx', output=True)
        with pytest.raises(ValueError, match='cannot delegate'):
            await runner.delegation.delegate(child, 'Recursive', ['inputs/nested/source.xlsx'], 'scout')
        assert (deps.root / f'working/tasks/{key}/history.json').exists()
        assert not list((deps.root / 'outputs').glob('*.xlsx'))

    asyncio.run(run())


def test_delegation_limit_cancellation_and_shared_budget(tmp_path):
    store, job, deps, _ = fixture_job(tmp_path)
    entered = asyncio.Event()

    async def respond(messages, info):
        entered.set()
        await asyncio.Event().wait()

    settings = Settings(data_dir=tmp_path, max_subagents=1, max_requests=2)
    runner = Runner(store, settings, FunctionModel(respond))

    async def run():
        runner.delegation.begin(job['id'], RunUsage())
        task = await runner.delegation.delegate(deps, 'First', ['inputs/nested/source.xlsx'], 'scout')
        await entered.wait()
        assert 'error' in await runner.delegation.delegate(deps, 'Second', ['inputs/nested/source.xlsx'], 'scout')
        assert (await runner.delegation.cancel(deps, task['id']))['status'] == 'cancelled'
        runner.delegation.reserve_request(deps)
        from pydantic_ai.exceptions import UsageLimitExceeded

        with pytest.raises(UsageLimitExceeded):
            runner.delegation.reserve_request(deps)

    asyncio.run(run())


def test_budget_exhaustion_pauses_and_resume_gets_fresh_allowance(tmp_path):
    model, _ = scripted([('list_files', {}), final('partial', ['Intentional stop'])])
    with TestClient(
        create_app(Settings(data_dir=tmp_path, api_key='test', max_requests=1), model_override=model), base_url=LOCAL
    ) as client:
        job = create_job(client)
        client.post(f'/api/jobs/{job["id"]}/start')
        assert wait_for_job(client, job['id'], statuses=('paused', 'failed'))['status'] == 'paused'
        client.post(f'/api/jobs/{job["id"]}/resume')
        assert wait_for_job(client, job['id'])['status'] == 'partial'


def test_completion_without_coverage_validation_is_partial_and_endpoints_are_paginated(tmp_path):
    model, _ = scripted(
        [('create_output', {'filename': 'out.xlsx'}), ('validate_output', {'path': 'outputs/out.xlsx'}), final()]
    )
    with TestClient(
        create_app(Settings(data_dir=tmp_path, api_key='test'), model_override=model), base_url=LOCAL
    ) as client:
        job = create_job(client)
        client.post(f'/api/jobs/{job["id"]}/start')
        assert wait_for_job(client, job['id'])['status'] == 'partial'
        inventory = client.get(f'/api/jobs/{job["id"]}/inventory?limit=1').json()
        assert len(inventory['items']) == 1 and inventory['next_offset'] == 1
        assert client.get(f'/api/jobs/{job["id"]}/inventory?limit=1000').status_code == 400
        assert not client.get(f'/api/jobs/{job["id"]}/coverage').json()['passed']


def test_context_rebuild_retains_user_decisions_and_archived_evidence_on_summary_failure(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)
    state.instruction('Employee 000001 must be excluded; confirmed by the user.')
    deps.settings = Settings(data_dir=tmp_path, context_tokens=10000, response_tokens=100)

    def fail_summary(messages, info):
        raise RuntimeError('Synthetic summary failure')

    model = FunctionModel(fail_summary)
    runner = Runner(store, deps.settings, model)
    history = [ModelRequest(parts=[UserPromptPart('Original task')])]
    for number in range(12):
        history.extend(
            [
                ModelResponse(
                    parts=[
                        TextPart('Investigation notes ' * 80),
                        ToolCallPart('read_document', {}, tool_call_id=str(number)),
                    ]
                ),
                ModelRequest(
                    parts=[
                        ToolReturnPart(
                            'read_document', 'Evidence ' * 200 + f'UNIQUE_END_{number}', tool_call_id=str(number)
                        )
                    ]
                ),
            ]
        )
    ctx = SimpleNamespace(deps=deps, messages=history.copy(), usage=RunUsage())
    request = ModelRequestContext(
        model=model, messages=history, model_settings={}, model_request_parameters=ModelRequestParameters()
    )
    request_retry(state, 'main')  # Exercise actual forced replacement even with a provisioned tokenizer.
    result = asyncio.run(ContextPolicy(runner).before_model_request(ctx, request))
    serialized = str(result.messages)
    assert 'Employee 000001 must be excluded' in serialized
    assert len(result.messages) < len(history)
    assert any(event['name'] == 'context.summary_warning' for event in store.events(job['id']))
    assert list((deps.root / 'working/history-archives').glob('*.json'))
    archives = list((deps.root / 'working/tool-results').glob('*.json'))
    assert any('UNIQUE_END_0' in artifact.read_text() for artifact in archives)
    assert state.instructions()[0]['text'].startswith('Employee 000001')


def test_essential_instructions_over_budget_pause_without_dropping_them(tmp_path):
    store, _, deps, _ = fixture_job(tmp_path)
    deps.settings = Settings(data_dir=tmp_path, context_tokens=1000, response_tokens=100)
    model = FunctionModel(lambda messages, info: ModelResponse(parts=[TextPart('unused')]))
    runner = Runner(store, deps.settings, model)
    history = [ModelRequest(parts=[UserPromptPart('Essential user instruction ' * 1000)])]
    ctx = SimpleNamespace(deps=deps, messages=history.copy(), usage=RunUsage())
    request = ModelRequestContext(
        model=model, messages=history, model_settings={}, model_request_parameters=ModelRequestParameters()
    )
    with pytest.raises(ContextBudgetError):
        asyncio.run(ContextPolicy(runner).before_model_request(ctx, request))


def test_one_oversized_batch_of_reads_is_cleared_instead_of_pausing(tmp_path):
    store, job, deps, _ = fixture_job(tmp_path)
    deps.settings = Settings(data_dir=tmp_path, context_tokens=4000, response_tokens=100)
    model = FunctionModel(lambda messages, info: ModelResponse(parts=[TextPart('unused')]))
    runner = Runner(store, deps.settings, model)
    calls = [ToolCallPart('read_document', {}, tool_call_id=str(n)) for n in range(3)]
    history = [
        ModelRequest(parts=[UserPromptPart('Original task')]),
        ModelResponse(parts=calls),
        ModelRequest(
            parts=[
                ToolReturnPart('read_document', 'Evidence ' * 3000 + f'END_{n}', tool_call_id=str(n)) for n in range(3)
            ]
        ),
    ]
    ctx = SimpleNamespace(deps=deps, messages=history.copy(), usage=RunUsage())
    request = ModelRequestContext(
        model=model, messages=history, model_settings={}, model_request_parameters=ModelRequestParameters()
    )
    result = asyncio.run(ContextPolicy(runner).before_model_request(ctx, request))
    assert 'END_2' not in str(result.messages)
    archives = list((deps.root / 'working/tool-results').glob('*.json'))
    assert any('END_2' in artifact.read_text() for artifact in archives)


def test_changed_source_and_changed_output_fail_validation(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    path = output_copy(deps)
    populate_template(state, dataset['id'], path, destination(), 'first')
    run_tool(store, job['id'], 'validate_output', {'path': path})
    assert validate_job(deps, path)['passed']
    workbook = load_workbook(deps.root / path)
    workbook['Output']['B2'] = 999
    workbook.save(deps.root / path)
    workbook.close()
    assert not validate_job(deps, path)['passed']
    (deps.root / 'inputs/nested/source.xlsx').write_bytes(b'changed')
    assert any('Source changed' in item for item in validate_job(deps, path)['failures'])


def test_restart_marks_unfinished_child_interrupted_without_repeating_work(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)
    state.put('task', 'child', {'id': 'child', 'status': 'running', 'objective': 'Scout'})

    async def recover():
        runner = Runner(store, deps.settings)
        await runner.start()
        await runner.close()

    asyncio.run(recover())
    assert state.get('task', 'child')['status'] == 'interrupted'


def test_duplicate_file_occurrences_keep_distinct_coverage_when_parsing_is_reused(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)
    manifest = store.get(job['id'])['files']
    add_file(
        deps.root, 'inputs', 'another/source.xlsx', (deps.root / 'inputs/nested/source.xlsx').read_bytes(), manifest
    )
    store.update(job['id'], files=manifest)
    state.inventory(manifest)
    inspect_source(state, 'inputs/nested/source.xlsx')
    duplicate = inspect_source(state, 'inputs/another/source.xlsx')
    assert duplicate['reused_content_hash']
    assert state.coverage()['source_rows'] == 8
    assert len({key for key, _ in state.all('observation')}) == 8


def test_nested_zip_expands_with_cumulative_limits(tmp_path):
    from zipfile import ZipFile

    from app.files import add_zip

    nested, outer = io.BytesIO(), io.BytesIO()
    with ZipFile(nested, 'w') as archive:
        archive.writestr('team/source.txt', 'nested content')
    with ZipFile(outer, 'w') as archive:
        archive.writestr('batch.zip', nested.getvalue())
    manifest = []
    add_zip(tmp_path, 'packet.zip', outer.getvalue(), manifest, 10, 10000)
    assert manifest[0]['path'] == 'inputs/packet/batch/team/source.txt'
    with pytest.raises(ValueError):
        add_zip(tmp_path, 'other.zip', outer.getvalue(), [], 10, 5)


def test_source_exclusions_are_reconsidered_after_new_instructions(tmp_path):
    _, _, _, state = fixture_job(tmp_path)
    for key, item in state.all('inventory'):
        if item.get('role') != 'template':
            state.put(
                'inventory', key, item | {'status': 'excluded', 'reason': 'Outside scope', 'instruction_revision': 0}
            )
    assert state.coverage()['passed']
    state.instruction('Include the previously excluded workbook')
    assert not state.coverage()['passed']
    source = next(key for key, item in state.all('inventory') if item['path'] == 'inputs/nested/source.xlsx')
    state.apply_instructions([1], 'The workbook is now in scope', invalidate=[source])
    assert state.get('inventory', source)['status'] == 'inspected'
    assert not state.coverage()['passed']


def test_completed_operation_replay_cannot_erase_a_new_conflict(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    _, rows = dataset_rows(state, dataset['id'])
    record = rows[0]['id']
    revised = run_tool(
        store,
        job['id'],
        'resolve_records',
        {
            'dataset': dataset['id'],
            'resolutions': [
                {
                    'record_id': record,
                    'values': {'amount': '12.50'},
                    'reason': 'Confirmed',
                    'evidence': 'Instruction p1',
                }
            ],
        },
    )
    path = output_copy(deps)
    populate_template(state, revised['id'], path, destination(), 'corrected')
    state.put('decision', record, {'status': 'unresolved', 'reason': 'New conflict'})
    revision = state.revision()
    populate_template(state, revised['id'], path, destination(), 'corrected')
    assert state.revision() == revision
    assert not state.coverage()['passed']
    assert state.get('decision', record)['reason'] == 'New conflict'


def test_partial_column_overwrite_cannot_drop_untouched_cell_tracking(tmp_path):
    _, _, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    path = output_copy(deps)
    populate_template(state, dataset['id'], path, destination(), 'full')
    with pytest.raises(ValueError, match='Partial-column overwrite'):
        populate_template(
            state,
            dataset['id'],
            path,
            Destination(sheet='Output', start_row=2, end_row=4, columns=[DestinationColumn(field='id', column=1)]),
            'partial',
        )
    assert len(state.all('assignment')) == 3


def test_new_scout_after_validation_prevents_complete(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    path = output_copy(deps)
    populate_template(state, dataset['id'], path, destination(), 'full')
    run_tool(store, job['id'], 'validate_output', {'path': path})
    assert validate_job(deps, path)['passed']
    revision = state.revision()
    state.put('task', 'late', {'id': 'late', 'status': 'running'})
    assert state.revision() == revision
    Runner(store, deps.settings).complete(
        job['id'],
        SimpleNamespace(output=SimpleNamespace(summary='Done', status='complete', unresolved=[]), usage=RunUsage()),
    )
    assert store.get(job['id'])['status'] == 'partial'


def test_repeated_cancellation_does_not_release_a_mutating_thread(tmp_path):
    _, _, deps, _ = fixture_job(tmp_path)
    entered, release, finished = threading.Event(), threading.Event(), threading.Event()

    def work():
        entered.set()
        release.wait(5)
        finished.set()
        return {'ok': True}

    async def run():
        task = asyncio.create_task(deps.action('slow_write', {}, work))
        await asyncio.to_thread(entered.wait, 3)
        task.cancel()
        await asyncio.sleep(0.01)
        task.cancel()
        await asyncio.sleep(0.01)
        try:
            assert not task.done()
            assert not finished.is_set()
        finally:
            release.set()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert finished.is_set()

    asyncio.run(run())


def test_destination_discovery_does_not_become_source_records(tmp_path):
    _, _, deps, state = fixture_job(tmp_path)
    inspect_source(state, 'template/template.xlsx')
    dataset = extract(state)
    populate_template(state, dataset['id'], output_copy(deps), destination(), 'full')
    assert state.coverage()['source_rows'] == 4
    assert state.coverage()['passed']


def test_numeric_strings_fail_even_when_their_displayed_value_matches(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    path = output_copy(deps)
    populate_template(state, dataset['id'], path, destination(), 'full')
    run_tool(
        store, job['id'], 'write_cells', {'path': path, 'changes': [{'sheet': 'Output', 'cell': 'B2', 'value': '10'}]}
    )
    run_tool(store, job['id'], 'validate_output', {'path': path})
    assert not validate_job(deps, path)['passed']


def test_recovery_after_instruction_change_settles_old_commit_without_blessing_old_mapping(tmp_path, monkeypatch):
    _, _, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    path = output_copy(deps)
    put = state.put_many

    def crash(kind, entries):
        if kind == 'assignment':
            raise OSError('crash')
        return put(kind, entries)

    monkeypatch.setattr(state, 'put_many', crash)
    with pytest.raises(OSError):
        populate_template(state, dataset['id'], path, destination(), 'interrupted')
    restored = JobState(deps.root)
    restored.instruction('New scope excludes the first employee')
    assert recover_operation(restored, 'interrupted')['status'] == 'completed'
    assert restored.get('mapping', 'interrupted')['instruction_revision'] == 0
    assert not validate_job(deps, path)['passed']
    fresh = extract(restored)
    populate_template(restored, fresh['id'], path, destination(), 'revised')
    assert restored.get('operation', 'revised')['status'] == 'completed'


def test_document_coverage_only_credits_text_delivered_within_result_budget(tmp_path):
    model, _ = scripted([final('partial', ['Test'])])
    with TestClient(
        create_app(Settings(data_dir=tmp_path, api_key='test'), model_override=model), base_url=LOCAL
    ) as client:
        job = create_job(client, files=[('files', ('long.txt', ('漢字"' * 6000).encode(), 'text/plain'))])
        store = client.app.state.store
        state = JobState(store.path(job['id']))
        first = run_tool(store, job['id'], 'read_document', {'path': 'inputs/long.txt'})
        assert 'text' in first and len(json.dumps(first, ensure_ascii=False).encode()) <= 16000
        assert not state.get('read', 'inputs/long.txt')['complete']
        offset = first['next_offset']
        while offset is not None:
            chunk = run_tool(store, job['id'], 'read_document', {'path': 'inputs/long.txt', 'offset': offset})
            assert 'text' in chunk
            offset = chunk['next_offset']
        assert state.get('read', 'inputs/long.txt')['complete']


def test_document_observations_have_lineage_and_require_output_assignment(tmp_path):
    model, _ = scripted([final('partial', ['Test'])])
    with TestClient(
        create_app(Settings(data_dir=tmp_path, api_key='test'), model_override=model), base_url=LOCAL
    ) as client:
        job = create_job(client)
        store = client.app.state.store
        path = 'inputs/HR/Reports/employees.txt'
        run_tool(store, job['id'], 'read_document', {'path': path})
        dataset = run_tool(
            store,
            job['id'],
            'record_observations',
            {
                'path': path,
                'observations': [
                    {
                        'key': 'EMP001',
                        'values': {'id': 'EMP001', 'salary': '4200'},
                        'kinds': {'salary': 'decimal'},
                        'evidence': 'Employee record in the source text',
                    }
                ],
            },
        )
        state = JobState(store.path(job['id']))
        assert dataset['rows'] == 1
        assert state.coverage()['records']['pending'] == 1
        assert state.coverage()['source_rows'] == 0
        assert state.coverage()['source_observations'] == 1


def test_repeated_tool_failures_pause_instead_of_looping(tmp_path):
    steps = [('read_sheet', {'path': 'inputs/missing.xlsx', 'sheet': 'Missing'})] * 6
    model, _ = scripted(steps)
    with TestClient(
        create_app(Settings(data_dir=tmp_path, api_key='test'), model_override=model), base_url=LOCAL
    ) as client:
        job = create_job(client)
        client.post(f'/api/jobs/{job["id"]}/start')
        stopped = wait_for_job(client, job['id'], statuses=('paused', 'failed'))
        assert stopped['status'] == 'paused'
        assert stopped['request_count'] == 5


def test_repeated_headers_and_numeric_subtotals_are_reported_before_writing(tmp_path):
    _, _, deps, state = fixture_job(tmp_path, dirty=True)
    dataset = extract(state, count=5)
    assert dataset['exceptions'] == 2
    assert dataset['profiles']['amount']['types']['str'] == 1
    filtered = transform_dataset(
        state,
        dataset['id'],
        [
            Transformation(
                operation='filter', predicate={'field': 'id', 'operator': 'ne', 'value': 'ID'}, reason='Repeated header'
            ),
            Transformation(
                operation='filter',
                predicate={'field': 'id', 'operator': 'ne', 'value': 'Grand Total'},
                reason='Subtotal is not an employee',
            ),
        ],
    )
    result = populate_template(state, filtered['id'], output_copy(deps), destination(), 'cleaned')
    assert result['records_written'] == 3
    assert state.coverage()['passed']


def test_delegation_and_bulk_write_errors_return_to_the_agent(tmp_path):
    model, _ = scripted(
        [
            ('delegate_task', {'objective': 'Scout', 'paths': ['inputs/typo.xlsx']}),
            (
                'populate_template',
                {
                    'dataset': 'missing',
                    'path': 'template/template.xlsx',
                    'destination': destination().model_dump(),
                    'operation_id': 'bad',
                },
            ),
            final('partial', ['Recovered from tool errors']),
        ]
    )
    with TestClient(
        create_app(Settings(data_dir=tmp_path, api_key='test', max_subagents=2), model_override=model), base_url=LOCAL
    ) as client:
        job = create_job(client)
        client.post(f'/api/jobs/{job["id"]}/start')
        assert wait_for_job(client, job['id'])['status'] == 'partial'
        failed = {
            event['name'] for event in client.app.state.store.events(job['id']) if event.get('status') == 'failed'
        }
        assert {'delegate_task', 'populate_template'} <= failed


def test_pausing_a_job_interrupts_workers_instead_of_failing_them(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)

    def respond(messages, info):
        store.update(job['id'], status='pausing')
        return ModelResponse(parts=[ToolCallPart('list_files', {}, tool_call_id='list')])

    runner = Runner(store, deps.settings, FunctionModel(respond))

    async def run():
        runner.delegation.begin(job['id'], RunUsage())
        task = await runner.delegation.delegate(deps, 'Scout', ['inputs/nested/source.xlsx'], 'scout')
        await runner.delegation.tasks[(job['id'], task['id'])]
        return task['id']

    key = asyncio.run(run())
    assert state.get('task', key)['status'] == 'interrupted'
    assert not any(event['name'] == 'task.failed' for event in store.events(job['id']))


def test_job_record_changes_never_overwrite_a_concurrent_status_change(tmp_path):
    store, job, _, _ = fixture_job(tmp_path)
    writer = threading.Thread(target=lambda: store.update(job['id'], status='cancelled'))

    def attach(record):
        # The cancel arrives while this change is in progress; it must wait, then apply on top.
        writer.start()
        record['files'] = [*record['files'], {'path': 'inputs/extra.txt', 'size': 1, 'sha256': 'x'}]

    store.mutate(job['id'], attach)
    writer.join()
    final_job = store.get(job['id'])
    assert final_job['status'] == 'cancelled'
    assert any(entry['path'] == 'inputs/extra.txt' for entry in final_job['files'])


def test_context_retry_flag_belongs_to_one_actor_and_is_consumed_once(tmp_path):
    store, _, deps, state = fixture_job(tmp_path)
    model = FunctionModel(lambda messages, info: ModelResponse(parts=[TextPart('unused')]))
    runner = Runner(store, deps.settings, model)
    request_retry(state, 'main')
    state.put('task', 'worker', {'id': 'worker', 'status': 'running', 'instruction_revision': 0})
    worker = JobDeps(store, deps.job_id, deps.settings, 'worker', ('inputs/nested/source.xlsx',))
    history = [ModelRequest(parts=[UserPromptPart('Scout')])]

    def request():
        ctx = SimpleNamespace(deps=None, messages=history.copy(), usage=RunUsage())
        return ctx, ModelRequestContext(
            model=model, messages=history, model_settings={}, model_request_parameters=ModelRequestParameters()
        )

    ctx, context = request()
    ctx.deps = worker
    asyncio.run(ContextPolicy(runner).before_model_request(ctx, context))
    assert state.get('context', 'main:retry')['force']
    ctx, context = request()
    ctx.deps = deps
    asyncio.run(ContextPolicy(runner).before_model_request(ctx, context))
    assert not state.get('context', 'main:retry')['force']


def test_context_length_retry_does_not_send_the_request_twice(tmp_path):
    seen = []

    def respond(messages, info):
        seen.append(sum(isinstance(part, UserPromptPart) for message in messages for part in message.parts))
        if len(seen) == 1:
            raise RuntimeError('context_length_exceeded')
        return ModelResponse(
            parts=[ToolCallPart('final_result', {'summary': 'Done', 'status': 'partial', 'unresolved': ['x']})]
        )

    with TestClient(
        create_app(Settings(data_dir=tmp_path, api_key='test'), model_override=FunctionModel(respond)),
        base_url=LOCAL,
    ) as client:
        job = create_job(client)
        client.post(f'/api/jobs/{job["id"]}/start')
        assert wait_for_job(client, job['id'])['status'] == 'partial'
    history = json.loads((tmp_path / 'jobs' / job['id'] / 'working/history.json').read_text())
    kinds = [message['kind'] for message in history]
    assert seen == [1, 1]
    assert kinds[:2] == ['request', 'response'], kinds


def test_source_backed_cell_writes_count_as_output_assignments(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)
    inspect_source(state, 'inputs/nested/source.xlsx')
    rows = {item['row']: key for key, item in state.all('observation') if item.get('sheet') == 'Data'}
    path = output_copy(deps)
    changes = [
        change
        for row in (2, 3, 4)
        for change in (
            {'sheet': 'Output', 'cell': f'A{row}', 'value': f'{row - 1:06}', 'record': rows[row]},
            {'sheet': 'Output', 'cell': f'B{row}', 'value': (row - 1) * 10, 'record': rows[row]},
        )
    ]
    run_tool(store, job['id'], 'write_cells', {'path': path, 'changes': changes})
    state.put('decision', rows[1], {'status': 'excluded', 'reason': 'Header', 'evidence': 'Data!1'})
    assert state.coverage()['records'] == {'pending': 0, 'included': 3, 'excluded': 1, 'unresolved': 0}
    run_tool(store, job['id'], 'validate_output', {'path': path})
    assert validate_job(deps, path)['passed']
    run_tool(
        store, job['id'], 'write_cells', {'path': path, 'changes': [{'sheet': 'Output', 'cell': 'B2', 'value': 5}]}
    )
    run_tool(store, job['id'], 'validate_output', {'path': path})
    assert any('Changed assignment' in failure for failure in validate_job(deps, path)['failures'])
    assert 'error' in run_tool(
        store, job['id'], 'write_cells', {'path': path, 'changes': [{**changes[0], 'record': 'unknown'}]}
    )


def test_document_observations_infer_numeric_kinds_for_filtering_and_writing(tmp_path):
    model, _ = scripted([final('partial', ['Test'])])
    with TestClient(
        create_app(Settings(data_dir=tmp_path, api_key='test'), model_override=model), base_url=LOCAL
    ) as client:
        job = create_job(client)
        store = client.app.state.store
        path = 'inputs/HR/Reports/employees.txt'
        run_tool(store, job['id'], 'read_document', {'path': path})
        dataset = run_tool(
            store,
            job['id'],
            'record_observations',
            {
                'path': path,
                'observations': [
                    {'key': 'EMP001', 'values': {'id': 'EMP001', 'salary': 4200.5}, 'evidence': 'Line 1'},
                ],
            },
        )
        assert dataset['kinds'] == {'id': 'raw', 'salary': 'decimal'}
        state = store.state(job['id'])
        filtered = transform_dataset(
            state,
            dataset['id'],
            [
                Transformation(
                    operation='filter',
                    predicate={'field': 'salary', 'operator': 'gt', 'value': 1000},
                    reason='Below threshold',
                )
            ],
        )
        _, rows = dataset_rows(state, filtered['id'])
        assert rows[0]['values']['salary'] == '4200.5'


def test_alternating_unproductive_calls_pause_the_job(tmp_path):
    steps = [
        ('read_sheet', {'path': 'inputs/missing.xlsx', 'sheet': 'A'}),
        ('read_sheet', {'path': 'inputs/missing.xlsx', 'sheet': 'B'}),
    ] * 6
    model, _ = scripted(steps)
    with TestClient(
        create_app(Settings(data_dir=tmp_path, api_key='test'), model_override=model), base_url=LOCAL
    ) as client:
        job = create_job(client)
        client.post(f'/api/jobs/{job["id"]}/start')
        stopped = wait_for_job(client, job['id'], statuses=('paused', 'failed'))
        assert stopped['status'] == 'paused'
        assert 5 <= stopped['request_count'] <= 9


def test_worker_reads_do_not_establish_document_coverage_for_the_main_agent(tmp_path):
    model, _ = scripted([final('partial', ['Test'])])
    with TestClient(
        create_app(Settings(data_dir=tmp_path, api_key='test'), model_override=model), base_url=LOCAL
    ) as client:
        job = create_job(client)
        state = client.app.state.store.state(job['id'])
        path = 'inputs/HR/Reports/employees.txt'
        state.record_read(path, 1, 1, 0, 10, 10, actor='worker')
        assert state.get('read', path) is None
        assert state.get('read', 'worker:' + path)['complete']
        state.record_read(path, 1, 1, 0, 10, 10)
        assert state.get('read', path)['complete']


def test_invalidated_dataset_can_be_rebuilt_from_the_same_region(tmp_path):
    _, _, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    state.instruction('Amounts were corrected at source')
    state.apply_instructions([1], 'Amounts changed', invalidate=[dataset['id']])
    rebuilt = extract(state)
    assert rebuilt['id'] != dataset['id'] and not rebuilt.get('stale')
    path = output_copy(deps)
    assert populate_template(state, rebuilt['id'], path, destination(), 'rebuilt')['records_written'] == 3


def test_reset_rejects_transforms_built_on_superseded_datasets(tmp_path):
    _, _, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    state.instruction('Use a different month')
    state.apply_instructions([1], 'Everything changes', reset_all=True)
    with pytest.raises(ValueError, match='Instructions changed'):
        transform_dataset(state, dataset['id'], [Transformation(operation='select', fields=['id', 'amount'])])


def test_equality_filters_compare_decimal_amounts_as_numbers(tmp_path):
    _, _, _, state = fixture_job(tmp_path)
    dataset = extract(state)
    kept = transform_dataset(
        state,
        dataset['id'],
        [
            Transformation(
                operation='filter', predicate={'field': 'amount', 'operator': 'eq', 'value': 20}, reason='Other rows'
            )
        ],
    )
    _, rows = dataset_rows(state, kept['id'])
    assert [row['values']['amount'] for row in rows] == ['20']
    # Identifiers keep exact text matching, so leading zeros still matter.
    by_id = transform_dataset(
        state,
        dataset['id'],
        [Transformation(operation='filter', predicate={'field': 'id', 'operator': 'eq', 'value': 1}, reason='Other')],
    )
    assert by_id['rows'] == 0


def test_blank_optional_amounts_validate_as_unchanged(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    blanked = transform_dataset(
        state, dataset['id'], [Transformation(operation='arithmetic', field='amount', target='bonus', operand=0)]
    )
    _, rows = dataset_rows(state, blanked['id'])
    for row in rows:
        row['values']['bonus'] = None
    state.put('rows', blanked['id'], rows)
    path = output_copy(deps)
    target = destination()
    target.columns.append(DestinationColumn(field='bonus', column=5))
    populate_template(state, blanked['id'], path, target, 'blank')
    run_tool(store, job['id'], 'validate_output', {'path': path})
    assert validate_job(deps, path)['passed'], validate_job(deps, path)['failures']


def test_shared_join_source_is_not_excluded_while_a_surviving_row_uses_it(tmp_path):
    _, _, _, state = fixture_job(tmp_path)
    dataset = extract(state)
    grouped = transform_dataset(state, dataset['id'], [Transformation(operation='select', fields=['amount'])])
    rates = state.get('rows', grouped['id'])[:1]
    rates[0]['values'] = {'code': 'all', 'rate': '1'}
    state.put('rows', 'rates', rates)
    state.put(
        'dataset', 'rates', {'id': 'rates', 'kinds': {'code': 'raw', 'rate': 'decimal'}, 'instruction_revision': 0}
    )
    shared = rates[0]['sources'][0]
    joined = transform_dataset(
        state,
        dataset['id'],
        [
            Transformation(
                operation='lookup',
                field='id',
                target='code',
                lookup={'000001': 'all', '000002': 'all', '000003': 'all'},
            ),
            Transformation(operation='join', fields=['code'], right_keys=['code'], dataset='rates', target='r_'),
            Transformation(
                operation='filter', predicate={'field': 'id', 'operator': 'ne', 'value': '000002'}, reason='Leaver'
            ),
        ],
    )
    excluded = {source for source, _ in state.get('exclusions', joined['id'])}
    assert shared not in excluded
    assert excluded


def test_editing_one_cell_keeps_provenance_for_the_rest_of_the_row(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)
    dataset = extract(state)
    path = output_copy(deps)
    populate_template(state, dataset['id'], path, destination(), 'bulk')
    observation = next(key for key, item in state.all('observation') if item.get('row') == 2)
    run_tool(
        store,
        job['id'],
        'write_cells',
        {'path': path, 'changes': [{'sheet': 'Output', 'cell': 'B2', 'value': 10, 'record': observation}]},
    )
    assignment = state.get('assignment', next(key for key, item in state.all('assignment') if item['row'] == 2))
    assert assignment['provenance']['id']['dataset'] == dataset['id']
    assert assignment['provenance']['amount']['dataset'] is None
    state.instruction('Identifiers changed')
    state.apply_instructions([1], 'IDs are wrong', invalidate=[dataset['id']])
    run_tool(store, job['id'], 'validate_output', {'path': path})
    assert any('invalidated' in failure for failure in validate_job(deps, path)['failures'])


def test_resume_without_delegation_settles_interrupted_workers(tmp_path):
    store, job, deps, state = fixture_job(tmp_path)
    state.put('task', 'child', {'id': 'child', 'status': 'interrupted', 'objective': 'Scout'})
    model, _ = scripted([final('partial', ['Test'])])
    runner = Runner(store, Settings(data_dir=tmp_path, max_subagents=0), model)
    asyncio.run(runner.run(job['id']))
    assert state.get('task', 'child')['status'] == 'cancelled'
