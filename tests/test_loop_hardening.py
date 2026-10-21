"""Regressions for the payroll run that kept inspecting sparse sheets and never extracted a record.

The run's pattern: wide, mostly empty sheet reads overflowed the result limit, the agent paged through its own
saved results, plan edits citing fresh reads kept clearing the stall guard, and an untouched template copy named
"_completed.xlsx" looked like the deliverable.
"""

import asyncio
import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pydantic_ai import Agent
from pydantic_ai.messages import ModelRequest, ModelResponse, TextPart, ToolCallPart, ToolReturnPart, UserPromptPart
from pydantic_ai.models.function import FunctionModel

from app.config import Settings
from app.files import add_file
from app.harness_tools import register_harness_tools
from app.main import create_app
from app.state import JobState, encode, fit_result
from app.store import Store
from app.tools import JobDeps, register_tools

from helpers import LOCAL, create_job, wait_for_job

SHEETS = ('New Joiner', 'Earning_Deduction', 'Resignation')


def sparse_payroll():
    """Three wide sheets that are mostly blank, like the uploaded payroll source."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title in SHEETS:
        sheet = workbook.create_sheet(title)
        sheet.append(['Employee ID', 'Name', 'Amount'])
        for number in range(1, 6):
            sheet.append([f'00{number}', f'Employee {number}', number * 100])
        # A stray formatted cell far away stretches the sheet to 90 columns and 400 rows.
        sheet.cell(400, 90, ' ')
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def dense_sheet(rows=300):
    workbook = Workbook()
    workbook.active.title = 'Data'
    for number in range(1, rows + 1):
        workbook.active.append([f'EMP{number:05}', f'Employee name {number}', number * 10, 'Department A'])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def job_with(tmp_path, name, payload, settings):
    store = Store(tmp_path)
    job = store.create('payroll', '2026-10', 'Fill the template', '')
    manifest = []
    add_file(store.path(job['id']), 'inputs', name, payload, manifest)
    store.update(job['id'], files=manifest)
    return store, job, JobDeps(store, job['id'], settings)


def call(deps, name, args):
    """Run one registered tool through the real agent plumbing and return what the model would see."""
    seen = []

    def respond(messages, info):
        seen.extend(
            part.content
            for message in messages
            if isinstance(message, ModelRequest)
            for part in message.parts
            if isinstance(part, ToolReturnPart)
        )
        if seen:
            return ModelResponse(parts=[TextPart('ok')])
        return ModelResponse(parts=[ToolCallPart(name, args, tool_call_id='call')])

    agent = Agent(FunctionModel(respond), deps_type=JobDeps)
    register_tools(agent)
    register_harness_tools(agent)
    asyncio.run(agent.run('go', deps=deps))
    return seen[0]


def test_sparse_sheet_read_returns_only_content_and_fits_without_spilling(tmp_path):
    settings = Settings(data_dir=tmp_path)
    _, _, deps = job_with(tmp_path, 'payroll.xlsx', sparse_payroll(), settings)
    result = call(
        deps, 'read_sheet', {'path': 'inputs/payroll.xlsx', 'sheet': 'New Joiner', 'end_row': 100, 'end_column': 90}
    )
    assert 'error' not in result, result
    # The old reader sent 9,000 cell records here, 99.8% of them blank, and overflowed into a saved file.
    assert 'artifact' not in result and 'truncated' not in result
    assert len(encode(result).encode()) < 4000
    assert [row['row'] for row in result['rows']] == [1, 2, 3, 4, 5, 6]
    assert result['rows'][1]['cells'][0] == {'cell': 'A2', 'value': '001'}
    assert result['empty_cells_omitted'] == 100 * 90 - 18
    assert result['next_row'] == 101 and result['more_rows']


def test_large_read_stops_at_a_whole_row_and_continues_without_gaps(tmp_path):
    settings = Settings(data_dir=tmp_path, tool_result_bytes=6000)
    _, _, deps = job_with(tmp_path, 'dense.xlsx', dense_sheet(), settings)
    rows, start, reads = [], 1, 0
    while start:
        result = call(
            deps,
            'read_sheet',
            {'path': 'inputs/dense.xlsx', 'sheet': 'Data', 'start_row': start, 'end_row': 300, 'end_column': 4},
        )
        reads += 1
        assert 'artifact' not in result, 'a read must page itself instead of spilling to a file'
        assert len(encode(result).encode()) <= settings.tool_result_bytes
        rows.extend(row['row'] for row in result['rows'])
        start = result['next_row']
        assert reads < 20
    assert rows == list(range(1, 301)) and reads > 1


def test_oversized_results_keep_their_structure_and_say_to_narrow_the_request(tmp_path):
    big = {'items': [{'id': index, 'text': 'x' * 300} for index in range(200)], 'total': 200}
    fitted = fit_result(big, 4000)
    assert len(encode(fitted).encode()) <= 4000
    assert fitted['total'] == 200 and fitted['items'][0]['id'] == 0
    assert fitted['items'][-1] == {'omitted_items': 200 - (len(fitted['items']) - 1)}
    settings = Settings(data_dir=tmp_path, tool_result_bytes=4000)
    _, _, deps = job_with(tmp_path, 'dense.xlsx', dense_sheet(5), settings)
    shown = asyncio.run(deps.action('large', {}, lambda: big))
    assert len(encode(shown).encode()) <= 4000
    assert shown['items'][0]['id'] == 0
    assert 'narrower request' in shown['instruction'] and 'read_document' not in shown['instruction']


@pytest.mark.parametrize('limit', [1024, 4000, 16000])
@pytest.mark.parametrize('text', ['记录' * 200, '🙂' * 200, '\\"\n' * 200])
def test_oversized_wide_results_fit_with_unicode_escaping_and_the_envelope(tmp_path, limit, text):
    big = {f'字段{index}': text for index in range(200)}
    fitted = fit_result(big, limit)
    assert len(encode(fitted).encode()) <= limit
    assert 'preview' in fitted
    settings = Settings(data_dir=tmp_path, tool_result_bytes=limit)
    _, _, deps = job_with(tmp_path, 'dense.xlsx', dense_sheet(1), settings)
    shown = asyncio.run(deps.action('large', {}, lambda: big))
    assert len(encode(shown).encode()) <= limit
    assert shown['truncated'] and 'narrower request' in shown['instruction']
    assert json.loads((deps.root / shown['artifact']).read_text()) == big


def test_resume_card_points_to_all_instructions_with_ids_and_application_status(tmp_path):
    settings = Settings(data_dir=tmp_path)
    _, _, deps = job_with(tmp_path, 'dense.xlsx', dense_sheet(1), settings)
    for number in range(1, 16):
        deps.state.instruction(f'Instruction {number}: preserve this constraint')
    deps.state.apply_instructions([1], 'Applies to later work')
    card = call(deps, 'read_job_state', {'kind': 'snapshot'})
    assert card['user_instructions']['total'] == 15
    assert card['user_instructions']['pending'] == 14
    assert 'kind="instruction"' in card['user_instructions']['read']
    instructions, offset = [], 0
    while offset is not None:
        page = call(deps, 'read_job_state', {'kind': 'instruction', 'offset': offset, 'limit': 4})
        instructions.extend(page['items'])
        offset = page['next_offset']
    assert instructions == [
        {field: item[field] for field in ('seq', 'kind', 'applied')} for item in deps.state.instructions()
    ]
    assert instructions[0]['seq'] == 1 and instructions[0]['applied']
    pending = [item['seq'] for item in instructions if not item['applied']]
    result = call(deps, 'apply_instructions', {'instructions': pending, 'reason': 'Apply the remaining constraints'})
    assert result['applied'] == list(range(2, 16))


def test_full_instruction_can_be_retrieved_by_sequence_after_restart(tmp_path):
    settings = Settings(data_dir=tmp_path)
    store, job, deps = job_with(tmp_path, 'dense.xlsx', dense_sheet(1), settings)
    text = 'Keep the original constraint. ' * 180
    deps.state.instruction(text, 'answer')
    restored = JobDeps(store, job['id'], settings)
    selected = call(restored, 'read_job_state', {'kind': 'instruction', 'key': '1', 'limit': 16000})
    assert selected['value'] == {'seq': 1, 'kind': 'answer', 'text': text, 'applied': False}
    assert selected['next_offset'] is None
    assert call(restored, 'read_job_state', {'kind': 'instruction', 'key': '999'})['value'] is None


def test_instruction_text_paging_recovers_long_unicode_constraints_without_truncation(tmp_path):
    settings = Settings(data_dir=tmp_path, tool_result_bytes=4000)
    _, _, deps = job_with(tmp_path, 'dense.xlsx', dense_sheet(1), settings)
    text = ('Keep this constraint: 记录🙂\\"\n' * 800) + 'Required final constraint.'
    deps.state.instruction(text)
    deps.state.apply_instructions([1], 'This applied constraint still governs future work')
    chunks, offset = [], 0
    while offset is not None:
        result = call(deps, 'read_job_state', {'kind': 'instruction', 'key': '1', 'offset': offset, 'limit': 16000})
        assert len(encode(result).encode()) <= settings.tool_result_bytes
        assert 'truncated' not in result
        assert result['value']['seq'] == 1 and result['value']['applied']
        assert result['next_offset'] is None or result['next_offset'] > offset
        chunks.append(result['value']['text'])
        offset = result['next_offset']
    assert len(chunks) > 1 and ''.join(chunks) == text


def test_instruction_index_adapts_its_cursor_to_the_byte_limit(tmp_path):
    settings = Settings(data_dir=tmp_path, tool_result_bytes=1024)
    _, _, deps = job_with(tmp_path, 'dense.xlsx', dense_sheet(1), settings)
    for number in range(30):
        deps.state.instruction(f'Constraint {number}: ' + '记录🙂' * 2000)
    seqs, offset = [], 0
    while offset is not None:
        result = call(deps, 'read_job_state', {'kind': 'instruction', 'offset': offset, 'limit': 100})
        assert 'truncated' not in result
        assert len(encode(result).encode()) <= settings.tool_result_bytes
        assert result['next_offset'] is None or result['next_offset'] == offset + len(result['items'])
        seqs.extend(item['seq'] for item in result['items'])
        offset = result['next_offset']
    assert seqs == list(range(1, 31))


def test_resume_card_is_small_and_names_unviewed_scans(tmp_path):
    state = JobState(tmp_path)
    state.inventory([{'path': 'inputs/scan.pdf', 'sha256': 'a' * 64, 'size': 1, 'role': 'inputs'}])
    for number in range(400):
        state.put('approach', f'approach-{number}', {'tool': 'read_sheet', 'count': 2, 'outcome': 'completed ' * 50})
    state.put(
        'plan', 'current', {'objective': 'Prepare payroll ' * 200, 'focus': 'Read', 'items': [], 'next_actions': []}
    )
    state.record_read('inputs/scan.pdf', 1, 2, needs_vision=True)
    state.record_read('inputs/scan.pdf', 2, 2, vision=True)
    card = state.resume_card()
    assert len(encode(card).encode()) <= 4000
    assert card['unviewed_scanned_pages'] == [{'path': 'inputs/scan.pdf', 'page': 1}]


def test_unviewed_scanned_page_is_blocked_once_by_the_existing_document_rule(tmp_path):
    state = JobState(tmp_path)
    state.put('schema', 'accounting', {'version': 2})
    state.inventory([{'path': 'inputs/scan.pdf', 'sha256': 'a' * 64, 'size': 1, 'role': 'inputs'}])
    state.record_read('inputs/scan.pdf', 1, 1, needs_vision=True)
    reasons = [issue['reason'] for issue in state.coverage()['issues']]
    # The unread scan keeps the document unreviewed; no second, duplicate blocker is added for the same page.
    assert reasons == ['Document has no reviewed disposition']
    assert not state.get('read', 'inputs/scan.pdf')['complete']


def replay_inspection_loop(tmp_path, max_requests=40):
    """Replay the failing run's pattern against the real runtime and return (job, events, prompts).

    The scripted agent creates the output, then only reads: sheet pages, a plan rewrite every third turn citing
    its latest reads, and a scan of a new workbook whenever the harness asks it to recover. It never extracts.
    """
    prompts, job_ids, issued = [], [], []
    sources = [f'branch-{number}.xlsx' for number in range(1, 7)]

    def respond(messages, info):
        prompts.append(messages)
        turn = len(prompts)
        latest = ' '.join(
            part.content
            for part in messages[-1].parts
            if isinstance(part, UserPromptPart) and isinstance(part.content, str)
        )
        if turn == 1:
            return ModelResponse(parts=[ToolCallPart('create_output', {'filename': 'payroll_output.xlsx'}, 'out')])
        if 'Harness recovery' in latest and sources:
            # In the trace, recovery was "answered" by inspecting a further workbook's sheets.
            path = 'inputs/' + sources.pop(0)
            return ModelResponse(parts=[ToolCallPart('inspect_source', {'path': path}, f'scan-{turn}')])
        if turn % 3 == 0 and issued:
            state = JobState(tmp_path / 'jobs' / job_ids[0])
            evidence = [key for key, item in state.all('evidence') if item['location'] in issued[-2:]]
            plan = {
                'objective': 'Prepare payroll',
                'focus': 'Understand the sheets',
                'items': [
                    {'id': 'understand', 'objective': 'Read every sheet', 'status': 'active', 'evidence': evidence}
                ],
                'next_actions': ['Read the next sheet'],
            }
            return ModelResponse(parts=[ToolCallPart('update_work_plan', {'plan': plan}, f'plan-{turn}')])
        start = 1 + 40 * turn
        args = {'path': 'inputs/payroll.xlsx', 'sheet': SHEETS[turn % 3], 'start_row': start, 'end_row': start + 39}
        issued.append({**args, 'range': [start, start + 39, 1, 4], 'formulas': False})
        issued[-1].pop('start_row'), issued[-1].pop('end_row')
        return ModelResponse(parts=[ToolCallPart('read_sheet', {**args, 'end_column': 4}, f'read-{turn}')])

    settings = Settings(data_dir=tmp_path, api_key='test', max_requests=max_requests)
    sheet = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    uploads = [('files', (name, sparse_payroll(), sheet)) for name in ('payroll.xlsx', *sources)]
    with TestClient(create_app(settings, model_override=FunctionModel(respond)), base_url=LOCAL) as client:
        job = create_job(client, uploads, legacy=False)
        job_ids.append(job['id'])
        client.post(f'/api/jobs/{job["id"]}/start')
        stopped = wait_for_job(client, job['id'], statuses=('paused', 'failed', 'completed', 'partial'), timeout=180)
        events = client.get(f'/api/jobs/{job["id"]}/events', params={'after': 0}).json()
    return stopped, events, prompts, settings


def test_inspection_loop_stops_at_the_request_allowance_with_a_draft_output(tmp_path):
    stopped, events, prompts, settings = replay_inspection_loop(tmp_path, max_requests=20)
    assert stopped['status'] == 'paused', stopped.get('error')
    # No turn-by-turn steering: the hard allowance bounds an aimless run, and its draft is kept.
    assert stopped['request_count'] <= settings.max_requests
    assert 'Harness' not in str(prompts[-1])
    assert [(item['name'], item.get('status')) for item in stopped['artifacts']] == [('payroll_output.xlsx', 'draft')]


def test_output_filename_cannot_claim_a_status(tmp_path):
    settings = Settings(data_dir=tmp_path)
    store, job, deps = job_with(tmp_path, 'dense.xlsx', dense_sheet(3), settings)
    template = Workbook()
    buffer = io.BytesIO()
    template.save(buffer)
    manifest = store.get(job['id'])['files']
    add_file(store.path(job['id']), 'template', 'template.xlsx', buffer.getvalue(), manifest)
    store.update(job['id'], files=manifest)
    refused = call(deps, 'create_output', {'filename': 'payroll_completed.xlsx'})
    assert 'status' in refused['error']
    assert call(deps, 'create_output', {'filename': 'payroll_output.xlsx'})['path'] == 'outputs/payroll_output.xlsx'


def test_context_estimate_learns_from_the_providers_real_count(tmp_path):
    from types import SimpleNamespace

    from app.context import calibrate

    state = JobState(tmp_path)
    real = SimpleNamespace(usage=SimpleNamespace(input_tokens=11000), provider_name='openai')
    # The live run estimated about 45k tokens for requests the provider counted at about 11k.
    first = calibrate(state, {'estimated_tokens': 44000}, real)
    assert first['factor'] == 0.25
    # The next estimate already carries that factor, so the comparison must undo it first.
    second = calibrate(state, {'estimated_tokens': 11000, 'calibration': 0.25}, real)
    assert second['factor'] == 0.25 and second['samples'] == 2
    # A test model has no provider count and must not move the estimate.
    fake = SimpleNamespace(usage=SimpleNamespace(input_tokens=5), provider_name=None)
    assert calibrate(state, {'estimated_tokens': 44000}, fake) is None
    assert state.get('context', 'calibration')['samples'] == 2


def test_prose_totals_are_refused_with_a_plain_message():
    import pytest
    from pydantic import ValidationError

    from app.datasets import Destination
    from app.requirements import RequiredField

    rule = 'Both source salary values must match exactly'
    with pytest.raises(ValidationError, match='plain number'):
        Destination(
            sheet='Employees',
            start_row=2,
            end_row=3,
            columns=[{'field': 'Salary', 'column': 3}],
            totals={'Salary': rule},
        )
    with pytest.raises(ValidationError, match='plain number'):
        RequiredField(column=3, meaning='Monthly salary', kind='decimal', expected_total=rule)
    assert RequiredField(column=3, meaning='Monthly salary', expected_total='10000.00').expected_total == '10000.00'


def test_a_destination_block_may_be_covered_by_several_adjoining_areas():
    from openpyxl.utils.cell import range_boundaries

    from app.requirements import covered

    block = range_boundaries('A1:E3')
    headers, data, formulas = (range_boundaries(bounds) for bounds in ('A1:E1', 'A2:D3', 'E2:E3'))
    assert covered(block, [headers, data, formulas])
    assert not covered(block, [headers, data])  # E2:E3 left out
    assert covered(block, [range_boundaries('A1:E3')])


def test_relaxed_rules_accept_reasonable_work_but_keep_their_guarantees(tmp_path):
    import pytest

    from app.inventory import inspect_source
    from app.progress import observe
    from app.regions import RegionRole, save_role
    from app.requirements import OutputArea, OutputSpecification, check_requirements, save_requirements
    from app.workplan import WorkItem, WorkPlan, save_plan

    from goal_helpers import establish_requirements, messy_case

    _, _, deps = messy_case(tmp_path)
    state = deps.state
    path = establish_requirements(deps)

    # A header row left exactly as the template had it is not counted as an output record.
    evidence = next(
        key for key, item in state.all('evidence') if item['location']['path'] == 'template/destination.xlsx'
    )
    with_header = OutputSpecification(
        objective='Fill payments',
        path=path,
        areas=[
            OutputArea(
                id='payments',
                sheet='Output',
                bounds='A1:C3',
                role='required',
                min_records=1,
                evidence=[evidence],
                rationale='Header row plus records',
                fields=[{'column': 2, 'meaning': 'Amount', 'kind': 'decimal'}],
            ),
            OutputArea(
                id='formula',
                sheet='Output',
                bounds='D2:D2',
                role='preserved',
                evidence=[evidence],
                rationale='Template calculation',
            ),
        ],
    )
    save_requirements(deps, with_header)
    from openpyxl import load_workbook

    workbook = load_workbook(state.root / path)
    workbook['Output']['A2'], workbook['Output']['B2'], workbook['Output']['C2'] = '00101', 10, 'REG'
    workbook.save(state.root / path)
    failures = check_requirements(state, path)
    assert not any('Numeric field' in item or 'Uncheckable' in item for item in failures), failures

    # Naming the template in place of the single output copy is corrected rather than refused.
    saved = save_requirements(deps, with_header.model_copy(update={'path': 'template/destination.xlsx'}))
    assert saved['path'] == path
    # One valid citation backs an area; a mistyped extra one is dropped and reported.
    mixed = with_header.model_copy(
        update={
            'areas': [area.model_copy(update={'evidence': [evidence, 'not-a-real-id']}) for area in with_header.areas]
        }
    )
    assert save_requirements(deps, mixed)['ignored_evidence'] == ['not-a-real-id']
    with pytest.raises(ValueError, match='Available references'):
        save_requirements(
            deps,
            with_header.model_copy(
                update={
                    'areas': [area.model_copy(update={'evidence': ['not-a-real-id']}) for area in with_header.areas]
                }
            ),
        )

    # Interpretations follow the same rule.
    inspect_source(state, 'inputs/nested/October.xlsx')
    region = next(key for key, item in state.all('region') if item['bounds'] == 'A14:C16')
    source_evidence = observe(deps, 'inspect_regions', {'path': 'inputs/nested/October.xlsx', 'region': region}, {})
    role = RegionRole(
        region=region, role='records', rationale='Current payments', evidence=[source_evidence[0]['key'], 'typo']
    )
    assert save_role(state, role)['ignored_evidence'] == ['typo']

    # A plan is the agent's own notes: unknown references are reported, and nothing in it is refused.
    plan = WorkPlan(
        objective='Pay',
        focus='Extract',
        next_actions=['Extract the region'],
        items=[WorkItem(id='read', objective='Read', status='active', evidence=['validate_output'])],
    )
    assert save_plan(state, plan)['unrecognized_references'] == ['validate_output']
    done = plan.model_copy(update={'items': [WorkItem(id='read', objective='Read', status='done', evidence=['nope'])]})
    assert save_plan(state, done)['items'][0]['status'] == 'done'


def test_record_decisions_accepts_a_file_path_for_a_source(tmp_path):
    from helpers import create_job

    settings = Settings(data_dir=tmp_path, api_key='test')
    with TestClient(create_app(settings), base_url=LOCAL) as client:
        job = create_job(client)
        store = client.app.state.store
    deps = JobDeps(store, job['id'], settings)
    call(deps, 'read_document', {'path': 'inputs/HR/Reports/employees.txt'})
    decision = {
        'id': 'inputs/HR/Reports/employees.txt',
        'scope': 'source',
        'status': 'reviewed',
        'reason': 'Read in full',
        'evidence': 'inputs/HR/Reports/employees.txt',
    }
    result = call(deps, 'record_decisions', {'decisions': [decision]})
    assert 'error' not in result, result
    file_id, _ = deps.state.file_entry('inputs/HR/Reports/employees.txt')
    assert deps.state.get('inventory', file_id)['status'] == 'reviewed'
    bogus = call(deps, 'record_decisions', {'decisions': [decision | {'id': 'inputs/nothing.txt'}]})
    assert 'file path' in bogus['error']


def test_output_name_check_matches_whole_words_only(tmp_path):
    from app.tools import STATUS_WORDS

    assert STATUS_WORDS.search('payroll_completed.xlsx') and STATUS_WORDS.search('Final.xlsx')
    assert not STATUS_WORDS.search('abandoned_claims.xlsx')
    assert not STATUS_WORDS.search('incomplete_items.xlsx')
