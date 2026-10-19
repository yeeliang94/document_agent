"""Synthetic region-accounting workbooks shared by harness regression tests."""

import io

from openpyxl import Workbook

from app.config import Settings
from app.files import add_file, atomic_bytes
from app.inventory import inspect_source
from app.progress import observe
from app.requirements import OutputArea, OutputSpecification, RequiredField, save_requirements
from app.store import Store
from app.tools import JobDeps


def workbook_payload(workbook):
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def messy_case(tmp_path, top=14, reference_column=7, image=False):
    store = Store(tmp_path)
    job = store.create(
        'payroll',
        '2026-10',
        'Prepare current October payments, preserving IDs; use code references, exclude the historical section.',
        '',
    )
    source = Workbook()
    sh = source.active
    sh.title = 'Payments'
    sh['A1'] = 'October 2026 - synthetic company'
    sh['A3'] = 'Current payments below; reference codes on the right.'
    for i, row in enumerate([['Employee ID', 'Amount', 'Code'], ['00101', 10, 'REG'], ['00102', 20, 'REG']], top):
        for j, value in enumerate(row, 1):
            sh.cell(i, j, value)
    for i, row in enumerate([['Code', 'Meaning'], ['REG', 'Regular'], ['BON', 'Bonus']], top):
        for j, value in enumerate(row, reference_column):
            sh.cell(i, j, value)
    sh.cell(top + 10, 1, 'September 2025 historical payments')
    sh.cell(top + 11, 1, 'OLD01')
    sh.cell(top + 11, 2, 999)
    sh.cell(1048576, 200).number_format = '0.00'  # sparse formatting cannot inflate a scan
    codes = source.create_sheet('Codes')
    codes.sheet_state = 'hidden'
    codes.append(['Code', 'Meaning'])
    codes.append(['REG', 'Regular October payment'])
    if image:
        from openpyxl.drawing.image import Image as ExcelImage
        from PIL import Image

        pixels = io.BytesIO()
        Image.new('RGB', (100, 30), 'white').save(pixels, format='PNG')
        pixels.seek(0)
        sh.add_image(ExcelImage(pixels), 'J2')
    template = Workbook()
    template.active.title = 'Output'
    template.active.append(['Employee ID', 'Amount', 'Code'])
    template.active['D2'] = '=SUM(B2:B3)'
    template.create_sheet('Lookup').sheet_state = 'hidden'
    manifest = []
    root = store.path(job['id'])
    add_file(root, 'inputs', 'nested/October.xlsx', workbook_payload(source), manifest)
    add_file(root, 'template', 'destination.xlsx', workbook_payload(template), manifest)
    store.update(job['id'], files=manifest)
    deps = JobDeps(store, job['id'], Settings(data_dir=tmp_path))
    deps.state.put('schema', 'accounting', {'version': 2})
    return store, job, deps


def establish_requirements(deps):
    state = deps.state
    inspect_source(state, 'template/destination.xlsx')
    evidence = observe(
        deps, 'inspect_workbook', {'path': 'template/destination.xlsx'}, {'sheets': ['Output', 'Lookup']}
    )[0]['key']
    path = 'outputs/completed.xlsx'
    atomic_bytes(state.root / path, (state.root / 'template/destination.xlsx').read_bytes())
    areas = [
        OutputArea(
            id='payments',
            sheet='Output',
            bounds='A2:C3',
            role='required',
            min_records=2,
            evidence=[evidence],
            rationale='Two current source records',
            fields=[
                RequiredField(column=1, meaning='Employee identifier', kind='identifier', unique=True),
                RequiredField(column=2, meaning='Payment amount', kind='decimal', expected_total='30'),
                RequiredField(column=3, meaning='Payment code', allowed_values=['REG']),
            ],
        ),
        OutputArea(
            id='headers',
            sheet='Output',
            bounds='A1:C1',
            role='reference',
            evidence=[evidence],
            rationale='Destination field labels',
        ),
        OutputArea(
            id='formula',
            sheet='Output',
            bounds='D2:D2',
            role='preserved',
            evidence=[evidence],
            rationale='Preserve template calculation',
        ),
    ]
    save_requirements(deps, OutputSpecification(objective=deps.current()['instructions'], path=path, areas=areas))
    return path
