"""Ground-truth messy workbooks, general supervision, and killed parser regressions."""

import asyncio
import hashlib
import io
import sys
import time
from dataclasses import replace
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from app.config import Settings
from app.datasets import (
    Column,
    Destination,
    DestinationColumn,
    Transformation,
    extract_dataset,
    populate_template,
    transform_dataset,
)
from app.execution import execute_tool
from app.files import add_file, atomic_bytes
from app.harness_tools import validate_job
from app.inventory import inspect_source
from app.progress import observe
from app.regions import RegionRole, save_role, search
from app.requirements import OutputArea, OutputSpecification, RequiredField, check_requirements, save_requirements
from app.state import JobState, identity
from app.store import Store
from app.supervisor import CapabilityResult, Supervisor, normalized, repetition
from app.tools import JobDeps, job_lock, validate_workbook
from app.workplan import WorkItem, WorkPlan, save_plan

from helpers import run_tool


def workbook_payload(workbook):
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def messy_case(tmp_path, top=14, reference_column=7, image=False):
    store = Store(tmp_path)
    job = store.create(
        'payroll',
        '2026-10',
        'Prepare current October payments, preserving IDs; use code references, exclude the historical section.',
        '',
    )
    source = Workbook()
    sh = source.active
    sh.title = 'Payments'
    sh['A1'] = 'October 2026 - synthetic company'
    sh['A3'] = 'Current payments below; reference codes on the right.'
    for i, row in enumerate([['Employee ID', 'Amount', 'Code'], ['00101', 10, 'REG'], ['00102', 20, 'REG']], top):
        for j, value in enumerate(row, 1):
            sh.cell(i, j, value)
    for i, row in enumerate([['Code', 'Meaning'], ['REG', 'Regular'], ['BON', 'Bonus']], top):
        for j, value in enumerate(row, reference_column):
            sh.cell(i, j, value)
    sh.cell(top + 10, 1, 'September 2025 historical payments')
    sh.cell(top + 11, 1, 'OLD01')
    sh.cell(top + 11, 2, 999)
    sh.cell(1048576, 200).number_format = '0.00'  # sparse formatting cannot inflate a scan
    codes = source.create_sheet('Codes')
    codes.sheet_state = 'hidden'
    codes.append(['Code', 'Meaning'])
    codes.append(['REG', 'Regular October payment'])
    if image:
        from openpyxl.drawing.image import Image as ExcelImage
        from PIL import Image

        pixels = io.BytesIO()
        Image.new('RGB', (100, 30), 'white').save(pixels, format='PNG')
        pixels.seek(0)
        sh.add_image(ExcelImage(pixels), 'J2')
    template = Workbook()
    template.active.title = 'Output'
    template.active.append(['Employee ID', 'Amount', 'Code'])
    template.active['D2'] = '=SUM(B2:B3)'
    template.create_sheet('Lookup').sheet_state = 'hidden'
    manifest = []
    root = store.path(job['id'])
    add_file(root, 'inputs', 'nested/October.xlsx', workbook_payload(source), manifest)
    add_file(root, 'template', 'destination.xlsx', workbook_payload(template), manifest)
    store.update(job['id'], files=manifest)
    deps = JobDeps(store, job['id'], Settings(data_dir=tmp_path))
    deps.state.put('schema', 'accounting', {'version': 2})
    return store, job, deps


def establish_requirements(deps):
    state = deps.state
    inspect_source(state, 'template/destination.xlsx')
    evidence = observe(
        deps, 'inspect_workbook', {'path': 'template/destination.xlsx'}, {'sheets': ['Output', 'Lookup']}
    )[0]['key']
    path = 'outputs/completed.xlsx'
    atomic_bytes(state.root / path, (state.root / 'template/destination.xlsx').read_bytes())
    areas = [
        OutputArea(
            id='payments',
            sheet='Output',
            bounds='A2:C3',
            role='required',
            min_records=2,
            evidence=[evidence],
            rationale='Two current source records',
            fields=[
                RequiredField(column=1, meaning='Employee identifier', kind='identifier', unique=True),
                RequiredField(column=2, meaning='Payment amount', kind='decimal', expected_total='30'),
                RequiredField(column=3, meaning='Payment code', allowed_values=['REG']),
            ],
        ),
        OutputArea(
            id='headers',
            sheet='Output',
            bounds='A1:C1',
            role='reference',
            evidence=[evidence],
            rationale='Destination field labels',
        ),
        OutputArea(
            id='formula',
            sheet='Output',
            bounds='D2:D2',
            role='preserved',
            evidence=[evidence],
            rationale='Preserve template calculation',
        ),
    ]
    save_requirements(deps, OutputSpecification(objective=deps.current()['instructions'], path=path, areas=areas))
    return path


@pytest.mark.parametrize(('top', 'right'), [(14, 7), (37, 10)])
def test_messy_regions_to_exact_output_with_independent_reference_rows(tmp_path, top, right):
    store, job, deps = messy_case(tmp_path, top, right)
    state = deps.state
    first = inspect_source(state, 'inputs/nested/October.xlsx', max_cells=4)
    assert not first['sheets'][0]['scan_complete']
    assert not state.all('observation')
    while not state.get('inventory', identity(state.file_entry('inputs/nested/October.xlsx')[0], 'Payments'))[
        'scan_complete'
    ]:
        inspect_source(state, 'inputs/nested/October.xlsx', sheet='Payments', max_cells=12)
    inspect_source(state, 'inputs/nested/October.xlsx')
    assert search(state, 'inputs/nested/October.xlsx', 'Employee ID')['items'][0]['row'] == top
    assert any(item.get('visibility') == 'hidden' for _, item in state.all('inventory'))
    revision = state.revision()
    inspect_source(state, 'inputs/nested/October.xlsx')
    assert state.revision() == revision
    output = establish_requirements(deps)
    selected = None
    for key, region in state.all('region'):
        if region['path'].startswith('template/'):
            continue
        result = {'id': key}
        evidence = observe(deps, 'inspect_regions', {'path': region['path'], 'region': key}, result)[0]['key']
        role = 'reference'
        if region['bounds'] == f'A{top}:C{top + 2}':
            selected = key
            role = 'records'
        elif region['bounds'].startswith(f'A{top + 10}:'):
            role = 'outside_scope'
        elif region['bounds'].startswith(('A1:', 'A3:')) and region['sheet'] == 'Payments':
            role = 'instructions'
        save_role(
            state,
            RegionRole(
                region=key,
                role=role,
                rationale='Inspected exact current/historical labels and code definitions',
                evidence=[evidence],
            ),
        )
    assert selected
    dataset = extract_dataset(
        state,
        'inputs/nested/October.xlsx',
        'Payments',
        top,
        top + 2,
        [
            Column(name='id', column=1, kind='identifier'),
            Column(name='amount', column=2, kind='decimal'),
            Column(name='code', column=3, kind='text'),
        ],
        [top],
        region_id=selected,
    )
    assert dataset['rows'] == 2
    assert len(state.all('observation')) == 2
    destination = Destination(
        sheet='Output',
        start_row=2,
        end_row=3,
        columns=[
            DestinationColumn(field='id', column=1),
            DestinationColumn(field='amount', column=2),
            DestinationColumn(field='code', column=3),
        ],
        required=['id', 'amount', 'code'],
    )
    populate_template(state, dataset['id'], output, destination, 'current-payments')
    assert validate_workbook(deps, output)['passed']
    checked = validate_job(deps, output)
    assert checked['passed'], checked
    workbook = load_workbook(state.root / output)
    assert [[workbook['Output'].cell(row, column).value for column in range(1, 4)] for row in (2, 3)] == [
        ['00101', 10, 'REG'],
        ['00102', 20, 'REG'],
    ]
    assert workbook['Output']['D2'].value == '=SUM(B2:B3)'
    assert workbook['Lookup'].sheet_state == 'hidden'
    assert state.coverage()['records'] == {'included': 2, 'excluded': 0, 'pending': 0, 'unresolved': 0}
    assert run_tool(store, job['id'], 'inspect_regions', {'path': 'inputs/nested/October.xlsx'})['total'] >= 5


def test_unknown_scope_and_destination_areas_cannot_pass(tmp_path):
    _, _, deps = messy_case(tmp_path)
    inspect_source(deps.state, 'inputs/nested/October.xlsx')
    assert not deps.state.coverage()['passed']
    output = establish_requirements(deps)
    assert check_requirements(deps.state, output)  # empty required rows
    specification = deps.state.get('requirement', identity(output))
    specification['areas'] = [area for area in specification['areas'] if area['id'] != 'formula']
    deps.state.put('requirement', identity(output), specification)
    assert any('Unclassified destination' in message for message in check_requirements(deps.state, output))


def test_visual_pixels_have_scoped_artifacts_and_block_coverage(tmp_path):
    store, job, deps = messy_case(tmp_path, image=True)
    inspect_source(deps.state, 'inputs/nested/October.xlsx')
    visuals = deps.state.all('visual')
    assert len(visuals) == 1
    key, visual = visuals[0]
    assert (deps.root / visual['artifact']).is_file()
    result = run_tool(store, job['id'], 'view_page', {'path': visual['artifact']})
    evidence = result['evidence_id']
    resolved = run_tool(
        store,
        job['id'],
        'resolve_visual',
        {'visual': key, 'role': 'outside_scope', 'evidence': evidence, 'reason': 'Inspected blank synthetic image'},
    )
    assert resolved['status'] == 'resolved'


@pytest.mark.parametrize('size', [1, 2, 3, 4, 5])
def test_repeating_cycles_pause_despite_volatile_results(size):
    recent = []
    for turn in range(max(5, size * 4)):
        fingerprint = identity(
            'tool-' + str(turn % size), normalized({'revision': turn, 'scanned_cells': turn * 3, 'stable': 1})
        )
        recent, warning, pause = repetition(recent, fingerprint)
    assert warning and pause


def test_general_supervisor_dedup_flip_flops_and_unlinked_wandering(tmp_path):
    state = JobState(tmp_path)
    supervisor = Supervisor(state)
    # A text-document adapter has no Excel concepts.
    envelope = CapabilityResult(evidence=[{'location': 'paragraph:4', 'hash': 'v1'}])
    assert envelope.evidence
    assert supervisor.event('interpretation_settled', 'section', 'A')
    supervisor.turn()
    assert supervisor.event('interpretation_settled', 'section', 'B')
    supervisor.turn()
    assert not supervisor.event('interpretation_settled', 'section', 'A')
    decisions = []
    for number in range(14):
        supervisor.event('evidence_inspected', f'paragraph:{number}', 'v1')
        decisions.append(supervisor.turn()[0])
    assert 'recover' in decisions and 'pause' in decisions
    restored = Supervisor(JobState(tmp_path))
    assert restored.turn()[0] == 'pause'


def test_linked_investigation_and_plan_survive_restart(tmp_path):
    state = JobState(tmp_path)
    supervisor = Supervisor(state)
    for number in range(12):
        key = f'paragraph:{number}'
        supervisor.event('evidence_inspected', key, 'v1')
        supervisor.link([key])
        assert supervisor.turn()[0] == 'continue'
    state.put('evidence', 'paragraph:11', {'source_hash': 'v1'})
    plan = WorkPlan(
        objective='Summarize clauses',
        focus='Resolve conflict',
        items=[WorkItem(id='clause', objective='Compare evidence', evidence=['paragraph:11'])],
        next_actions=['Read the conflicting clause'],
    )
    save_plan(state, plan)
    restored = JobState(tmp_path)
    assert restored.snapshot()['plan']['next_actions'] == plan.next_actions
    with pytest.raises(ValidationError):
        WorkPlan(objective='Task', focus='Focus', items=[], next_actions=['action'] * 6)


def test_isolated_nonreturning_process_is_killed_before_lock_release(tmp_path, monkeypatch):
    _, _, deps = messy_case(tmp_path)
    deps.settings = replace(deps.settings, tool_timeout=0.2)
    original = asyncio.create_subprocess_exec
    processes = []

    async def hang(*args, **kwargs):
        process = await original(
            sys.executable, '-c', 'import sys,time; sys.stdin.buffer.read(); time.sleep(60)', **kwargs
        )
        processes.append(process)
        return process

    monkeypatch.setattr(asyncio, 'create_subprocess_exec', hang)
    started = time.monotonic()
    with pytest.raises(InterruptedError):
        asyncio.run(execute_tool(SimpleNamespace(deps=deps), 'inspect_source', {'path': 'inputs/nested/October.xlsx'}))
    assert time.monotonic() - started < 3
    assert processes[0].returncode is not None
    assert job_lock(deps.job_id).acquire(blocking=False)
    job_lock(deps.job_id).release()
    assert deps.current()['status'] == 'pausing'


def test_read_only_deadline_is_a_tool_error_not_a_pause(tmp_path, monkeypatch):
    _, _, deps = messy_case(tmp_path)
    deps.settings = replace(deps.settings, tool_timeout=0.2)
    original = asyncio.create_subprocess_exec

    async def hang(*args, **kwargs):
        return await original(
            sys.executable, '-c', 'import sys,time; sys.stdin.buffer.read(); time.sleep(60)', **kwargs
        )

    monkeypatch.setattr(asyncio, 'create_subprocess_exec', hang)
    status = deps.current()['status']
    with pytest.raises(TimeoutError, match='deadline'):
        asyncio.run(
            execute_tool(
                SimpleNamespace(deps=deps), 'read_sheet', {'path': 'inputs/nested/October.xlsx', 'sheet': 'Sheet'}
            )
        )
    assert job_lock(deps.job_id).acquire(blocking=False)
    job_lock(deps.job_id).release()
    assert deps.current()['status'] == status


def test_snapshot_reports_the_supervision_clock(tmp_path):
    state = JobState(tmp_path)
    Supervisor(state).turn()
    assert state.snapshot()['supervision']['turn'] == 1


def test_unpivot_has_one_to_many_source_accounting(tmp_path):
    state = JobState(tmp_path)
    state.put(
        'dataset',
        'wide',
        {'id': 'wide', 'kinds': {'id': 'identifier', 'oct': 'decimal', 'nov': 'decimal'}, 'instruction_revision': 0},
    )
    state.put(
        'rows',
        'wide',
        [
            {
                'id': 'record',
                'sources': ['source'],
                'values': {'id': '001', 'oct': '10', 'nov': '20'},
                'evidence': {},
                'issues': [],
            }
        ],
    )
    result = transform_dataset(state, 'wide', [Transformation(operation='unpivot', fields=['oct', 'nov'])])
    rows = state.get('rows', result['id'])
    assert [row['values'] for row in rows] == [
        {'id': '001', 'category': 'oct', 'value': '10'},
        {'id': '001', 'category': 'nov', 'value': '20'},
    ]
    assert all(row['sources'] == ['source'] for row in rows)


def test_saving_verified_requirements_is_useful_progress(tmp_path):
    _, _, deps = messy_case(tmp_path)
    path = establish_requirements(deps)
    result = {'id': identity(path), **deps.state.get('requirement', identity(path))}
    assert observe(deps, 'save_output_requirements', {'path': path}, result)
    assert not observe(deps, 'save_output_requirements', {'path': path}, result)


def test_subset_of_accepted_region_cannot_hide_unextracted_rows(tmp_path):
    _, _, deps = messy_case(tmp_path)
    state = deps.state
    inspect_source(state, 'inputs/nested/October.xlsx')
    key = next(key for key, r in state.all('region') if r['bounds'] == 'A14:C16')
    result = {}
    evidence = observe(deps, 'inspect_regions', {'path': 'inputs/nested/October.xlsx', 'region': key}, result)[0]['key']
    save_role(state, RegionRole(region=key, role='records', rationale='Current payments', evidence=[evidence]))
    with pytest.raises(ValueError, match='complete accepted region'):
        extract_dataset(
            state, 'inputs/nested/October.xlsx', 'Payments', 14, 15, [Column(name='id', column=1)], [14], region_id=key
        )


def test_native_chart_is_visible_as_an_unsupported_material_capability(tmp_path):
    from openpyxl.chart import BarChart, Reference

    _, _, deps = messy_case(tmp_path)
    source = deps.root / 'inputs/nested/October.xlsx'
    workbook = load_workbook(source)
    chart = BarChart()
    chart.add_data(Reference(workbook['Payments'], min_col=2, min_row=15, max_row=16))
    workbook['Payments'].add_chart(chart, 'J5')
    workbook.save(source)
    file_id, entry = deps.state.file_entry('inputs/nested/October.xlsx')
    deps.state.put('inventory', file_id, entry | {'sha256': hashlib.sha256(source.read_bytes()).hexdigest()})
    inspect_source(deps.state, 'inputs/nested/October.xlsx')
    assert any(v['status'] == 'unsupported' and v['artifact'] is None for _, v in deps.state.all('visual'))
    assert any('visual' in issue['reason'] for issue in deps.state.coverage()['issues'])


def test_region_job_runs_through_actual_runtime_and_current_validation(tmp_path):
    from fastapi.testclient import TestClient
    from pydantic_ai.messages import ModelRequest, ModelResponse, ToolCallPart, ToolReturnPart
    from pydantic_ai.models.function import FunctionModel

    from app.main import create_app

    from helpers import LOCAL, wait_for_job

    _, _, fixture = messy_case(tmp_path / 'fixtures', top=37)
    returns, calls = {}, []

    def respond(messages, info):
        for message in messages:
            if isinstance(message, ModelRequest):
                for part in message.parts:
                    if isinstance(part, ToolReturnPart) and isinstance(part.content, dict):
                        returns[part.tool_name] = part.content
        number = len(calls)
        actions = []
        if number == 0:
            actions = [
                ('inspect_source', {'path': 'inputs/October.xlsx'}),
                ('inspect_source', {'path': 'template/destination.xlsx'}),
                ('inspect_workbook', {'path': 'template/destination.xlsx'}),
            ]
        elif number == 1:
            actions = [
                ('inspect_regions', {'path': 'inputs/October.xlsx'}),
                ('create_output', {'filename': 'completed.xlsx'}),
            ]
        elif number == 2:
            for region in returns['inspect_regions']['items']:
                role = (
                    'records'
                    if region['bounds'] == 'A37:C39'
                    else 'outside_scope'
                    if region['bounds'].startswith('A47:')
                    else 'reference'
                )
                actions.append(
                    (
                        'save_interpretation',
                        {
                            'decision': {
                                'region': region['id'],
                                'role': role,
                                'rationale': 'Hand-authored synthetic fixture interpretation from inspected content',
                                'evidence': [returns['inspect_regions']['evidence_id']],
                            }
                        },
                    )
                )
            ev = returns['inspect_workbook']['evidence_id']
            actions.append(
                (
                    'save_output_requirements',
                    {
                        'specification': {
                            'objective': 'Prepare the two current payments',
                            'path': 'outputs/completed.xlsx',
                            'areas': [
                                {
                                    'id': 'data',
                                    'sheet': 'Output',
                                    'bounds': 'A2:C3',
                                    'role': 'required',
                                    'fields': [
                                        {'column': 1, 'meaning': 'ID', 'kind': 'identifier'},
                                        {'column': 2, 'meaning': 'Amount', 'kind': 'decimal', 'expected_total': '30'},
                                    ],
                                    'min_records': 2,
                                    'evidence': [ev],
                                    'rationale': 'Current payments',
                                },
                                {
                                    'id': 'headers',
                                    'sheet': 'Output',
                                    'bounds': 'A1:C1',
                                    'role': 'reference',
                                    'evidence': [ev],
                                    'rationale': 'Field labels',
                                },
                                {
                                    'id': 'formula',
                                    'sheet': 'Output',
                                    'bounds': 'D2:D2',
                                    'role': 'preserved',
                                    'evidence': [ev],
                                    'rationale': 'Template calculation',
                                },
                            ],
                        }
                    },
                )
            )
        elif number == 3:
            region = next(r['id'] for r in returns['inspect_regions']['items'] if r['bounds'] == 'A37:C39')
            actions = [
                (
                    'extract_region',
                    {
                        'region': region,
                        'start_row': 37,
                        'end_row': 39,
                        'columns': [
                            {'name': 'id', 'column': 1, 'kind': 'identifier'},
                            {'name': 'amount', 'column': 2, 'kind': 'decimal'},
                            {'name': 'code', 'column': 3, 'kind': 'text'},
                        ],
                        'header_rows': [37],
                    },
                )
            ]
        elif number == 4:
            actions = [
                (
                    'populate_template',
                    {
                        'dataset': returns['extract_region']['id'],
                        'path': 'outputs/completed.xlsx',
                        'destination': {
                            'sheet': 'Output',
                            'start_row': 2,
                            'end_row': 3,
                            'columns': [
                                {'field': 'id', 'column': 1},
                                {'field': 'amount', 'column': 2},
                                {'field': 'code', 'column': 3},
                            ],
                            'required': ['id', 'amount', 'code'],
                        },
                        'operation_id': 'payments',
                    },
                )
            ]
        elif number == 5:
            actions = [('validate_output', {'path': 'outputs/completed.xlsx'})]
        elif number == 6:
            actions = [('validate_job', {'path': 'outputs/completed.xlsx'})]
        elif number == 7:
            actions = [
                (
                    'update_work_plan',
                    {
                        'plan': {
                            'objective': 'Prepare October payments',
                            'focus': 'Validated and ready to deliver',
                            'items': [
                                {
                                    'id': 'payments',
                                    'objective': 'Prepare and validate current payments',
                                    'status': 'done',
                                    'evidence': [returns['extract_region']['id']],
                                }
                            ],
                            'next_actions': ['Deliver the validated output'],
                        }
                    },
                )
            ]
        else:
            raise AssertionError('Verified completed work must finalize without another model call')
        calls.append(actions)
        return ModelResponse(
            parts=[
                ToolCallPart(name, args, tool_call_id=f'{number}-{index}') for index, (name, args) in enumerate(actions)
            ]
        )

    app = create_app(Settings(data_dir=tmp_path / 'app', api_key='synthetic'), model_override=FunctionModel(respond))
    with TestClient(app, base_url=LOCAL) as client:
        uploaded = client.post(
            '/api/jobs',
            data={'module': 'payroll', 'period': '2026-10', 'instructions': 'Prepare October payments'},
            files=[
                ('template', ('destination.xlsx', (fixture.root / 'template/destination.xlsx').read_bytes())),
                ('files', ('October.xlsx', (fixture.root / 'inputs/nested/October.xlsx').read_bytes())),
            ],
        )
        assert uploaded.status_code == 201
        job = uploaded.json()
        assert app.state.store.state(job['id']).get('schema', 'accounting')['version'] == 2
        assert client.post(f'/api/jobs/{job["id"]}/start').status_code == 200
        completed = wait_for_job(client, job['id'], statuses=('completed', 'partial', 'failed', 'paused'))
        assert completed['status'] == 'completed', completed
        assert client.get(f'/api/jobs/{job["id"]}/plan').json()['plan']['objective'] == 'Prepare October payments'
        assert client.get(f'/api/jobs/{job["id"]}/regions').json()['total'] >= 6
        workbook = load_workbook(io.BytesIO(client.get(f'/api/jobs/{job["id"]}/artifacts/completed.xlsx').content))
        assert [workbook['Output'][c].value for c in ('A2', 'B2', 'A3', 'B3', 'D2')] == [
            '00101',
            10,
            '00102',
            20,
            '=SUM(B2:B3)',
        ]


def test_reference_extraction_does_not_create_payment_obligations(tmp_path):
    _, _, deps = messy_case(tmp_path)
    state = deps.state
    inspect_source(state, 'inputs/nested/October.xlsx')
    key = next(key for key, r in state.all('region') if r['bounds'] == 'G14:H16')
    evidence = observe(deps, 'inspect_regions', {'path': 'inputs/nested/October.xlsx', 'region': key}, {})[0]['key']
    save_role(
        state,
        RegionRole(
            region=key, role='reference', rationale='Code lookup, not payroll transactions', evidence=[evidence]
        ),
    )
    result = extract_dataset(
        state,
        'inputs/nested/October.xlsx',
        'Payments',
        14,
        16,
        [Column(name='code', column=7), Column(name='meaning', column=8)],
        [14],
        region_id=key,
    )
    assert result['rows'] == 2
    assert state.coverage()['records'] == {'pending': 0, 'included': 0, 'excluded': 0, 'unresolved': 0}


def test_narrative_notes_do_not_invalidate_completed_checks(tmp_path):
    state = JobState(tmp_path)
    state.put('validation', 'current', {'passed': True, 'revision': state.revision()})
    state.put('memory', 'summary', {'value': 'Ready for delivery'})
    assert state.get('validation', 'current')['revision'] == state.revision()


def test_new_user_instruction_gets_bounded_recovery_but_resume_alone_does_not(tmp_path):
    supervisor = Supervisor(JobState(tmp_path))
    for _ in range(7):
        decision, _ = supervisor.turn()
    assert decision == 'pause'
    assert supervisor.instructions(1)
    assert supervisor.turn()[0] == 'continue'
    assert not supervisor.instructions(1)
    decisions = [supervisor.turn()[0] for _ in range(6)]
    assert 'recover' in decisions and 'pause' in decisions
    assert not supervisor.instructions(1)
    assert supervisor.turn()[0] == 'pause'


def test_rebuild_hysteresis_never_suppresses_recovery_at_the_hard_budget(tmp_path):
    from pydantic_ai.messages import ModelRequest, ModelResponse, TextPart, UserPromptPart
    from pydantic_ai.models import ModelRequestContext, ModelRequestParameters
    from pydantic_ai.models.function import FunctionModel
    from pydantic_ai.usage import RunUsage

    from app.context import ContextPolicy
    from app.runtime import Runner

    store = Store(tmp_path)
    job = store.create('payroll', '2026-10', 'Prepare a supported output', '')
    deps = JobDeps(
        store, job['id'], Settings(data_dir=tmp_path, context_tokens=6000, response_tokens=100, context_chars=1000000)
    )
    deps.state.put('context', 'main:usage', {'rebuild_high_water': 100000})

    def no_summary(messages, info):
        raise RuntimeError('Synthetic summary outage')

    model = FunctionModel(no_summary)
    history = [ModelRequest(parts=[UserPromptPart('Original objective')])]
    for _ in range(20):
        history.extend(
            [
                ModelResponse(parts=[TextPart('Detailed investigation with lots of observations. ' * 300)]),
                ModelRequest(parts=[UserPromptPart('Continue from saved facts')]),
            ]
        )
    context = ModelRequestContext(
        model=model, messages=history, model_settings={}, model_request_parameters=ModelRequestParameters()
    )
    ctx = SimpleNamespace(deps=deps, messages=history.copy(), usage=RunUsage())
    result = asyncio.run(ContextPolicy(Runner(store, deps.settings, model)).before_model_request(ctx, context))
    assert len(result.messages) < len(history)
    usage = deps.state.get('context', 'main:usage')
    assert usage['estimated_tokens'] < usage['budget']


def test_worker_evidence_does_not_establish_owner_review(tmp_path):
    _, _, deps = messy_case(tmp_path)
    state = deps.state
    inspect_source(state, 'inputs/nested/October.xlsx')
    region = next(key for key, value in state.all('region') if value['bounds'] == 'A14:C16')
    worker = JobDeps(deps.store, deps.job_id, deps.settings, actor='scout', scope=('inputs/nested/October.xlsx',))
    evidence = observe(worker, 'inspect_regions', {'path': 'inputs/nested/October.xlsx', 'region': region}, {})[0][
        'key'
    ]
    with pytest.raises(ValueError, match='Inspect source content'):
        save_role(state, RegionRole(region=region, role='records', rationale='Worker handover', evidence=[evidence]))
    owner_evidence = observe(deps, 'inspect_regions', {'path': 'inputs/nested/October.xlsx', 'region': region}, {})[0][
        'key'
    ]
    assert owner_evidence != evidence
    assert (
        save_role(
            state,
            RegionRole(
                region=region, role='records', rationale='Owner inspected source evidence', evidence=[owner_evidence]
            ),
        )['role']
        == 'records'
    )


def test_targeted_reads_expose_deep_comments_merges_formats_and_declared_tables(tmp_path):
    from openpyxl.comments import Comment
    from openpyxl.worksheet.table import Table

    store, job, deps = messy_case(tmp_path, top=37)
    source = deps.root / 'inputs/nested/October.xlsx'
    workbook = load_workbook(source)
    sheet = workbook['Payments']
    sheet['C38'].comment = Comment('REG is the current payment code.', 'Synthetic reviewer')
    sheet['A38'].number_format = '@'
    sheet.merge_cells('J37:K37')
    sheet['J37'] = 'Reference heading'
    sheet.add_table(Table(displayName='CurrentPayments', ref='A37:C39'))
    workbook.save(source)
    file_id, entry = deps.state.file_entry('inputs/nested/October.xlsx')
    deps.state.put('inventory', file_id, entry | {'sha256': hashlib.sha256(source.read_bytes()).hexdigest()})
    discovered = inspect_source(deps.state, 'inputs/nested/October.xlsx')
    assert discovered['sheets'][0]['tables'][0]['range'] == 'A37:C39'
    read = run_tool(
        store,
        job['id'],
        'read_sheet',
        {'path': 'inputs/nested/October.xlsx', 'sheet': 'Payments', 'start_row': 37, 'end_row': 39, 'end_column': 3},
    )
    assert read['rows'][1][2]['note'] == 'REG is the current payment code.'
    assert read['rows'][1][0]['number_format'] == '@'
    assert 'J37:K37' in read['merged_ranges']


def test_revising_an_accepted_region_invalidates_mapping_and_exposes_leftover_cells(tmp_path):
    from app.datasets import require_current

    store, job, deps = messy_case(tmp_path)
    state = deps.state
    inspect_source(state, 'inputs/nested/October.xlsx')
    key = next(key for key, r in state.all('region') if r['bounds'] == 'A14:C16')
    evidence = observe(deps, 'inspect_regions', {'path': 'inputs/nested/October.xlsx', 'region': key}, {})[0]['key']
    save_role(state, RegionRole(region=key, role='records', rationale='Current payments', evidence=[evidence]))
    extracted = extract_dataset(
        state, 'inputs/nested/October.xlsx', 'Payments', 14, 16, [Column(name='id', column=1)], [14], region_id=key
    )
    revised = run_tool(
        store,
        job['id'],
        'revise_region',
        {
            'region': key,
            'bounds': 'A14:C15',
            'supersede': [key],
            'reason': 'New evidence changes the boundary; remaining cells require investigation',
        },
    )
    assert 'error' not in revised, revised
    with pytest.raises(ValueError, match='interpretation changed'):
        require_current(state, extracted['id'])
    assert any('Unclassified occupied source area' in issue['reason'] for issue in state.coverage()['issues'])


def test_comment_only_content_below_previews_is_searchable_and_has_scope(tmp_path):
    from openpyxl.comments import Comment

    _, _, deps = messy_case(tmp_path)
    source = deps.root / 'inputs/nested/October.xlsx'
    workbook = load_workbook(source)
    workbook['Payments']['J50'].comment = Comment(
        'Material instruction: verify the effective period.', 'Synthetic reviewer'
    )
    workbook.save(source)
    file_id, entry = deps.state.file_entry('inputs/nested/October.xlsx')
    deps.state.put('inventory', file_id, entry | {'sha256': hashlib.sha256(source.read_bytes()).hexdigest()})
    inspect_source(deps.state, 'inputs/nested/October.xlsx')
    found = search(deps.state, 'inputs/nested/October.xlsx', 'Material instruction')
    assert found['items'][0]['cell'] == 'J50'
    assert any(r['bounds'] == 'J50:J50' for _, r in deps.state.all('region'))
    assert not deps.state.coverage()['passed']
