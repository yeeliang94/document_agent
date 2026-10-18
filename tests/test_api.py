"""Application contract tests; no external model requests are performed."""

import io
import time
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic_ai.messages import ModelResponse, ToolCallPart
from pydantic_ai.models.function import FunctionModel

from app.config import Settings
from app.main import create_app

from helpers import LOCAL, create_job, wait_for_job, workbook_bytes


def test_upload_nested_inputs_persists_job_and_serves_frontend(tmp_path):
    settings = Settings(data_dir=tmp_path, api_key='')
    with TestClient(base_url=LOCAL, app=create_app(settings)) as client:
        assert client.get('/').status_code == 200
        job = create_job(client)
        assert job['status'] == 'ready'
        assert client.get(f'/api/jobs/{job["id"]}').json()['instructions'] == 'Fill Employees A2:D2.'
        assert list((tmp_path / 'jobs' / job['id'] / 'inputs').rglob('employees.txt'))
    with TestClient(base_url=LOCAL, app=create_app(settings)) as client:
        assert client.get(f'/api/jobs/{job["id"]}').json()['period'] == '2026-10'


def test_start_without_key_has_actionable_error(tmp_path):
    with TestClient(base_url=LOCAL, app=create_app(Settings(data_dir=tmp_path, api_key=''))) as client:
        job = create_job(client)
        response = client.post(f'/api/jobs/{job["id"]}/start')
        assert response.status_code == 503
        assert 'OPENAI_API_KEY' in response.text
        assert client.get(f'/api/jobs/{job["id"]}').json()['status'] == 'ready'


@pytest.mark.parametrize('filename', ['../escape.txt', '/escape.txt', 'C:/escape.txt', 'safe/../../escape.txt'])
def test_uploaded_path_traversal_is_rejected(tmp_path, filename):
    with TestClient(base_url=LOCAL, app=create_app(Settings(data_dir=tmp_path, api_key=''))) as client:
        response = client.post(
            '/api/jobs',
            data={'module': 'payroll', 'period': '2026-10'},
            files=[
                ('template', ('template.xlsx', workbook_bytes())),
                ('files', (filename, b'unsafe')),
            ],
        )
        assert response.status_code == 400, response.text
        assert not (tmp_path / 'escape.txt').exists()


def test_zip_preserves_nested_files_and_rejects_zip_slip(tmp_path):
    for entry, expected_status in [('nested/team/employees.txt', 201), ('../escape.txt', 400)]:
        buffer = io.BytesIO()
        with ZipFile(buffer, 'w') as archive:
            archive.writestr(entry, 'Synthetic employees')
        with TestClient(base_url=LOCAL, app=create_app(Settings(data_dir=tmp_path, api_key=''))) as client:
            response = client.post(
                '/api/jobs',
                data={'module': 'payroll', 'period': '2026-10'},
                files=[
                    ('template', ('template.xlsx', workbook_bytes())),
                    ('files', ('packet.zip', buffer.getvalue(), 'application/zip')),
                ],
            )
            assert response.status_code == expected_status, response.text
            if expected_status == 201:
                assert (tmp_path / 'jobs' / response.json()['id'] / 'inputs' / 'packet' / entry).exists()
    assert not (tmp_path / 'escape.txt').exists()


def payroll_model():
    calls = []
    changes = [
        {'sheet': 'Employees', 'cell': cell, 'value': value}
        for cell, value in [('A2', 'EMP001'), ('B2', 'Alex Example'), ('C2', 4200), ('D2', '2026-10-01')]
    ]
    steps = [
        ('list_files', {}),
        ('inspect_workbook', {'path': 'template/template.xlsx'}),
        ('read_document', {'path': 'inputs/HR/Reports/employees.txt'}),
        ('create_output', {'filename': 'payroll_completed.xlsx'}),
        ('write_cells', {'path': 'outputs/payroll_completed.xlsx', 'changes': changes}),
        ('validate_output', {'path': 'outputs/payroll_completed.xlsx'}),
        (
            'final_result',
            {'summary': 'Prepared one employee; formula preserved.', 'status': 'complete', 'unresolved': []},
        ),
    ]

    def respond(messages, info):
        step = steps[len(calls)]
        calls.append(step[0])
        return ModelResponse(parts=[ToolCallPart(step[0], step[1], tool_call_id=f'test-{len(calls)}')])

    return FunctionModel(respond), calls


def test_real_agent_tools_write_validate_download_and_log(tmp_path):
    model, calls = payroll_model()
    with TestClient(
        base_url=LOCAL, app=create_app(Settings(data_dir=tmp_path, api_key='test-key'), model_override=model)
    ) as client:
        job = create_job(client)
        response = client.post(f'/api/jobs/{job["id"]}/start')
        assert response.status_code == 200, response.text
        completed = wait_for_job(client, job['id'])
        assert completed['status'] == 'completed', completed
        assert calls == [
            'list_files',
            'inspect_workbook',
            'read_document',
            'create_output',
            'write_cells',
            'validate_output',
            'final_result',
        ]
        output = tmp_path / 'jobs' / job['id'] / 'outputs/payroll_completed.xlsx'
        workbook = load_workbook(output)
        assert workbook['Employees']['A2'].value == 'EMP001'
        assert workbook['Employees']['C2'].value == 4200
        assert workbook['Employees']['E2'].value == '=C2*12'
        assert workbook['Lookup'].sheet_state == 'hidden'
        original = load_workbook(tmp_path / 'jobs' / job['id'] / 'template/template.xlsx')
        assert original['Employees']['A2'].value is None
        events = client.get(f'/api/jobs/{job["id"]}/events').json()
        assert any(event['name'] == 'write_cells' for event in events)
        assert any(event['name'] == 'validate_output' and event.get('result', {}).get('passed') for event in events)


def test_answer_resumes_deferred_tool_and_keeps_answer_in_history(tmp_path):
    from pydantic_ai.messages import ToolReturnPart

    steps = [
        ('ask_user', {'question': 'Confirm salary currency?'}),
        ('create_output', {'filename': 'confirmed.xlsx'}),
        ('validate_output', {'path': 'outputs/confirmed.xlsx'}),
        ('final_result', {'summary': 'Currency confirmed.', 'status': 'complete', 'unresolved': []}),
    ]
    calls = []
    answers = []

    def respond(messages, info):
        answers.extend(
            part.content
            for message in messages
            for part in message.parts
            if isinstance(part, ToolReturnPart) and part.tool_name == 'ask_user'
        )
        name, args = steps[len(calls)]
        calls.append(name)
        return ModelResponse(parts=[ToolCallPart(name, args, tool_call_id=f'answer-{len(calls)}')])

    with TestClient(
        base_url=LOCAL,
        app=create_app(Settings(data_dir=tmp_path, api_key='test-key'), model_override=FunctionModel(respond)),
    ) as client:
        job = create_job(client)
        client.post(f'/api/jobs/{job["id"]}/start')
        waiting = wait_for_job(client, job['id'], statuses=('waiting', 'failed'))
        assert waiting['status'] == 'waiting', waiting
        assert waiting['question']['text'] == 'Confirm salary currency?'
        response = client.post(f'/api/jobs/{job["id"]}/answer', json={'answer': 'MYR'})
        assert response.status_code == 200, response.text
        completed = wait_for_job(client, job['id'])
        assert completed['status'] == 'completed', completed
        assert 'MYR' in answers
        history = (tmp_path / 'jobs' / job['id'] / 'working/history.json').read_text()
        assert 'MYR' in history


def test_pause_checkpoints_at_tool_boundary_and_resume_completes(tmp_path):
    import asyncio
    import threading

    entered = threading.Event()
    release = threading.Event()
    calls = []
    steps = [
        ('create_output', {'filename': 'resumed.xlsx'}),
        (
            'write_cells',
            {'path': 'outputs/resumed.xlsx', 'changes': [{'sheet': 'Employees', 'cell': 'A2', 'value': 'EMP001'}]},
        ),
        ('validate_output', {'path': 'outputs/resumed.xlsx'}),
        ('final_result', {'summary': 'Resumed output complete.', 'status': 'complete', 'unresolved': []}),
    ]

    async def respond(messages, info):
        if not calls:
            entered.set()
            while not release.is_set():
                await asyncio.sleep(0.01)
        name, args = steps[len(calls)]
        calls.append(name)
        return ModelResponse(parts=[ToolCallPart(name, args, tool_call_id=f'pause-{len(calls)}')])

    with TestClient(
        base_url=LOCAL,
        app=create_app(Settings(data_dir=tmp_path, api_key='test-key'), model_override=FunctionModel(respond)),
    ) as client:
        job = create_job(client)
        client.post(f'/api/jobs/{job["id"]}/start')
        assert entered.wait(timeout=3)
        response = client.post(f'/api/jobs/{job["id"]}/pause')
        assert response.status_code == 200, response.text
        release.set()
        paused = wait_for_job(client, job['id'], statuses=('paused', 'failed'))
        assert paused['status'] == 'paused', paused
        assert calls == ['create_output']
        assert (tmp_path / 'jobs' / job['id'] / 'working/history.json').is_file()
        response = client.post(f'/api/jobs/{job["id"]}/resume')
        assert response.status_code == 200, response.text
        completed = wait_for_job(client, job['id'])
        assert completed['status'] == 'completed', completed
        events = client.get(f'/api/jobs/{job["id"]}/events').json()
        assert sum(e['name'] == 'create_output' for e in events) == 1
        output = client.get(f'/api/jobs/{job["id"]}/artifacts/resumed.xlsx')
        assert output.status_code == 200
        assert load_workbook(io.BytesIO(output.content))['Employees']['A2'].value == 'EMP001'


def test_compaction_keeps_tool_pairs_and_archives_full_content(tmp_path):
    from pydantic_ai.messages import ModelMessagesTypeAdapter, ToolReturnPart

    calls = []
    source = ('Synthetic payroll source row and detailed context. ' * 300).encode()

    def respond(messages, info):
        # Each supplied tool return must retain its preceding model tool call.
        known = set()
        for message in messages:
            for part in message.parts:
                if isinstance(part, ToolCallPart):
                    known.add(part.tool_call_id)
                if isinstance(part, ToolReturnPart):
                    assert part.tool_call_id in known
        calls.append(len(calls))
        if len(calls) <= 8:
            return ModelResponse(
                parts=[
                    ToolCallPart('read_document', {'path': 'inputs/large.txt'}, tool_call_id=f'compact-{len(calls)}')
                ]
            )
        return ModelResponse(
            parts=[
                ToolCallPart(
                    'final_result',
                    {
                        'summary': 'Read sources; no output requested by test.',
                        'status': 'partial',
                        'unresolved': ['Test stops after compaction.'],
                    },
                    tool_call_id='compact-final',
                )
            ]
        )

    settings = Settings(data_dir=tmp_path, api_key='test-key', context_chars=3000)
    with TestClient(base_url=LOCAL, app=create_app(settings, model_override=FunctionModel(respond))) as client:
        job = create_job(client, files=[('files', ('large.txt', source, 'text/plain'))])
        client.post(f'/api/jobs/{job["id"]}/start')
        completed = wait_for_job(client, job['id'])
        assert completed['status'] == 'partial', completed
        events = client.get(f'/api/jobs/{job["id"]}/events').json()
        compactions = [event for event in events if event['name'] == 'context.compact']
        assert compactions
        for event in compactions:
            assert event['bytes_after'] < event['bytes_before']
            archived = tmp_path / 'jobs' / job['id'] / event['artifact']
            assert archived.is_file()
            ModelMessagesTypeAdapter.validate_json(archived.read_bytes())
        assert (
            'Synthetic payroll source row' in (tmp_path / 'jobs' / job['id'] / compactions[0]['artifact']).read_text()
        )


def test_cancel_during_final_model_response_remains_cancelled(tmp_path):
    import asyncio
    import threading

    entered = threading.Event()
    release = threading.Event()
    calls = []

    async def respond(messages, info):
        index = len(calls)
        calls.append(index)
        if index == 0:
            return ModelResponse(
                parts=[ToolCallPart('create_output', {'filename': 'cancelled.xlsx'}, tool_call_id='cancel-1')]
            )
        if index == 1:
            return ModelResponse(
                parts=[ToolCallPart('validate_output', {'path': 'outputs/cancelled.xlsx'}, tool_call_id='cancel-2')]
            )
        entered.set()
        while not release.is_set():
            await asyncio.sleep(0.01)
        return ModelResponse(
            parts=[
                ToolCallPart(
                    'final_result',
                    {'summary': 'Final response arrived after cancellation.', 'status': 'complete', 'unresolved': []},
                    tool_call_id='cancel-final',
                )
            ]
        )

    app = create_app(Settings(data_dir=tmp_path, api_key='test-key'), model_override=FunctionModel(respond))
    with TestClient(app, base_url=LOCAL) as client:
        job = create_job(client)
        client.post(f'/api/jobs/{job["id"]}/start')
        assert entered.wait(timeout=3)
        response = client.post(f'/api/jobs/{job["id"]}/cancel')
        assert response.status_code == 200, response.text
        release.set()
        deadline = time.monotonic() + 3
        while app.state.runner.active is not None and time.monotonic() < deadline:
            time.sleep(0.01)
        assert app.state.runner.active is None
        result = client.get(f'/api/jobs/{job["id"]}').json()
        assert result['status'] == 'cancelled', result
        output = client.get(f'/api/jobs/{job["id"]}/artifacts/cancelled.xlsx')
        assert output.status_code == 200


def test_multiple_deferred_questions_resume_all_call_ids(tmp_path):
    from pydantic_ai.messages import ToolReturnPart

    responses = 0
    answers = {}

    def respond(messages, info):
        nonlocal responses
        responses += 1
        if responses == 1:
            return ModelResponse(
                parts=[
                    ToolCallPart('ask_user', {'question': 'Confirm currency?'}, tool_call_id='question-currency'),
                    ToolCallPart('ask_user', {'question': 'Confirm payroll period?'}, tool_call_id='question-period'),
                ]
            )
        for message in messages:
            for part in message.parts:
                if isinstance(part, ToolReturnPart) and part.tool_name == 'ask_user':
                    answers[part.tool_call_id] = part.content
        return ModelResponse(
            parts=[
                ToolCallPart(
                    'final_result',
                    {
                        'summary': 'Both answers received.',
                        'status': 'partial',
                        'unresolved': ['Synthetic test has no output.'],
                    },
                    tool_call_id='questions-final',
                )
            ]
        )

    with TestClient(
        base_url=LOCAL,
        app=create_app(Settings(data_dir=tmp_path, api_key='test-key'), model_override=FunctionModel(respond)),
    ) as client:
        job = create_job(client)
        client.post(f'/api/jobs/{job["id"]}/start')
        waiting = wait_for_job(client, job['id'], statuses=('waiting', 'failed'))
        assert waiting['status'] == 'waiting', waiting
        question = waiting['question']['text']
        assert 'currency' in question and 'payroll period' in question
        response = client.post(f'/api/jobs/{job["id"]}/answer', json={'answer': 'MYR; October 2026'})
        assert response.status_code == 200
        completed = wait_for_job(client, job['id'])
        assert completed['status'] == 'partial', completed
        assert answers == {'question-currency': 'MYR; October 2026', 'question-period': 'MYR; October 2026'}


@pytest.mark.parametrize('action,stopped_status', [('pause', 'paused'), ('cancel', 'cancelled')])
def test_resume_after_stopping_on_deferred_question_restores_waiting_state(tmp_path, action, stopped_status):
    import asyncio
    import threading

    entered = threading.Event()
    release = threading.Event()
    calls = []

    async def respond(messages, info):
        if not calls:
            entered.set()
            while not release.is_set():
                await asyncio.sleep(0.01)
            calls.append('question')
            return ModelResponse(
                parts=[ToolCallPart('ask_user', {'question': 'Confirm currency?'}, tool_call_id='stopped-question')]
            )
        calls.append('result')
        return ModelResponse(
            parts=[
                ToolCallPart(
                    'final_result',
                    {'summary': 'Answer received.', 'status': 'partial', 'unresolved': ['Synthetic run stops here.']},
                    tool_call_id='stopped-final',
                )
            ]
        )

    app = create_app(Settings(data_dir=tmp_path, api_key='test-key'), model_override=FunctionModel(respond))
    with TestClient(app, base_url=LOCAL) as client:
        job = create_job(client)
        client.post(f'/api/jobs/{job["id"]}/start')
        assert entered.wait(timeout=3)
        assert client.post(f'/api/jobs/{job["id"]}/{action}').status_code == 200
        release.set()
        stopped = wait_for_job(client, job['id'], statuses=(stopped_status, 'failed'))
        assert stopped['status'] == stopped_status, stopped
        deadline = time.monotonic() + 3
        while app.state.runner.active is not None and time.monotonic() < deadline:
            time.sleep(0.01)
        response = client.post(f'/api/jobs/{job["id"]}/resume')
        assert response.status_code == 200, response.text
        waiting = wait_for_job(client, job['id'], statuses=('waiting', 'failed'))
        assert waiting['status'] == 'waiting', waiting
        assert waiting['question']['text'] == 'Confirm currency?'
        assert client.post(f'/api/jobs/{job["id"]}/answer', json={'answer': 'MYR'}).status_code == 200
        assert wait_for_job(client, job['id'])['status'] == 'partial'
        assert calls == ['question', 'result']


def test_config_publishes_the_status_groups_the_api_enforces(tmp_path):
    from app import jobstatus

    with TestClient(base_url=LOCAL, app=create_app(Settings(data_dir=tmp_path, api_key=''))) as client:
        statuses = client.get('/api/config').json()['statuses']
    assert statuses == jobstatus.as_json()
    assert 'running' in statuses['active'] and 'partial' in statuses['startable']
