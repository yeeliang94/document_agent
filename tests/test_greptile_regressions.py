"""Greptile regressions for sparse extraction, protected areas and provenance recovery."""

import hashlib
from datetime import datetime, timezone
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.datasets import Column, Destination, DestinationColumn, extract_dataset, populate_template, recover_operation
from app.files import add_file, atomic_bytes
from app.inventory import inspect_source
from app.progress import observe
from app.regions import RegionRole, save_role
from app.requirements import OutputArea, OutputSpecification, check_write_scope, save_requirements
from app.state import identity
from app.store import Store
from app.tools import JobDeps, assign_cells

from goal_helpers import workbook_payload
from helpers import run_tool


def source_case(tmp_path, rows):
    store = Store(tmp_path)
    job = store.create('payroll', '2026-10', 'Copy supported employee records', '')
    source = Workbook()
    source.active.title = 'Source'
    source.active.append(['Employee ID', 'Amount', 'Note'])
    for row in rows:
        source.active.append(row)
    template = Workbook()
    template.active.title = 'Output'
    template.active.append(['Employee ID', 'Amount', 'Note'])
    root, manifest = store.path(job['id']), []
    add_file(root, 'inputs', 'source.xlsx', workbook_payload(source), manifest)
    add_file(root, 'template', 'template.xlsx', workbook_payload(template), manifest)
    store.update(job['id'], files=manifest)
    deps = JobDeps(store, job['id'])
    deps.state.put('schema', 'accounting', {'version': 2})
    inspect_source(deps.state, 'inputs/source.xlsx')
    region, _ = next((key, item) for key, item in deps.state.all('region') if item['path'] == 'inputs/source.xlsx')
    evidence = observe(deps, 'inspect_regions', {'path': 'inputs/source.xlsx', 'region': region}, {})[0]['key']
    save_role(deps.state, RegionRole(region=region, role='records', rationale='Employee records', evidence=[evidence]))
    path = 'outputs/output.xlsx'
    atomic_bytes(root / path, (root / 'template/template.xlsx').read_bytes())
    inspect_source(deps.state, 'template/template.xlsx')
    inspected = {}
    observe(deps, 'read_sheet', {'path': 'template/template.xlsx', 'sheet': 'Employees'}, inspected)
    save_requirements(
        deps,
        OutputSpecification(
            objective='Copy supported employees',
            path=path,
            areas=[
                OutputArea(
                    id='inputs',
                    sheet='Output',
                    bounds='A2:C10',
                    role='optional',
                    evidence=[inspected['evidence_id']],
                    rationale='Writable employee rows',
                ),
                OutputArea(
                    id='header',
                    sheet='Output',
                    bounds='A1:C1',
                    role='reference',
                    evidence=[inspected['evidence_id']],
                    rationale='Preserve field labels',
                ),
            ],
        ),
    )
    return store, job, deps, path, region


def extract(deps, region, rows, columns=None):
    return extract_dataset(
        deps.state,
        'inputs/source.xlsx',
        'Source',
        1,
        rows + 1,
        columns or [Column(name='id', column=1, kind='identifier'), Column(name='amount', column=2, kind='decimal')],
        [1],
        region_id=region,
    )


def test_sparse_optional_fields_retain_values_coordinates_and_coverage(tmp_path):
    _, _, deps, _, region = source_case(tmp_path, [['00101', None, 'Optional amount'], ['00102', 20]])
    dataset = extract(deps, region, 2)
    assert dataset['rows'] == 2 and dataset['exceptions'] == 0
    records = deps.state.get('rows', dataset['id'])
    assert records[0]['values'] == {'id': '00101', 'amount': None}
    assert records[0]['evidence']['amount']['cell'] == 'B2'
    assert deps.state.get('observation', records[0]['id'])['cells'] == ['A2', 'B2']
    assert deps.state.coverage()['records']['pending'] == 2


@pytest.mark.parametrize('content', ['amount', 'note', 'formula'])
def test_unmapped_occupied_rows_remain_unresolved_but_blank_rows_do_not(tmp_path, content):
    extra = {'amount': [None, 20], 'note': [None, None, 'Investigate employee'], 'formula': [None, '=10+10']}[content]
    _, _, deps, path, region = source_case(tmp_path, [['00101', 10], [None, None, None], extra])
    dataset = extract(deps, region, 3, [Column(name='id', column=1, kind='identifier')])
    assert dataset['rows'] == 2
    records = deps.state.get('rows', dataset['id'])
    assert records[1]['values'] == {'id': None}
    assert records[1]['issues']
    assert {item['row'] for _, item in deps.state.all('observation')} == {2, 4}
    written = populate_template(
        deps.state,
        dataset['id'],
        path,
        Destination(sheet='Output', start_row=2, end_row=3, columns=[DestinationColumn(field='id', column=1)]),
        'ids-only',
    )
    assert written['records_written'] == 1 and written['unresolved'] == 1
    assert not deps.state.coverage()['passed']
    assert deps.state.coverage()['records']['pending'] == 1


@pytest.mark.parametrize('role', ['preserved', 'reference', 'unresolved'])
@pytest.mark.parametrize('method', ['cells', 'bulk'])
def test_protected_overlap_blocks_both_writers_including_interior_rows(tmp_path, role, method):
    store, job, deps, path, region = source_case(tmp_path, [['00101', 10], ['00102', 20], ['00103', 30]])
    dataset = extract(deps, region, 3)
    specification = deps.state.get('requirement', identity(path))
    protection = specification['areas'][0] | {'id': 'protected', 'bounds': 'B3:B3', 'role': role}
    # Persisted specifications can overlap; protection must win regardless of area order.
    specification['areas'].append(protection)
    save_requirements(deps, OutputSpecification.model_validate(specification))
    original = (deps.root / path).read_bytes()
    if method == 'cells':
        result = run_tool(
            store,
            job['id'],
            'write_cells',
            {'path': path, 'changes': [{'sheet': 'Output', 'cell': 'B3', 'value': 999}]},
        )
        assert 'protected' in result['error'].lower()
    else:
        with pytest.raises(ValueError, match='protected'):
            populate_template(
                deps.state,
                dataset['id'],
                path,
                Destination(
                    sheet='Output', start_row=2, end_row=4, columns=[DestinationColumn(field='amount', column=2)]
                ),
                'protected-overwrite',
            )
    assert (deps.root / path).read_bytes() == original
    assert not deps.state.all('operation')
    check_write_scope(deps.state, path, 'Output', 2, 4, [1, 3])


@pytest.fixture
def frozen_excel(monkeypatch):
    fixed = datetime(2026, 10, 19, tzinfo=timezone.utc)
    monkeypatch.setattr(
        'openpyxl.writer.excel.datetime',
        SimpleNamespace(datetime=SimpleNamespace(now=lambda **kwargs: fixed), timezone=timezone),
    )
    monkeypatch.setattr('zipfile.time.localtime', lambda *args: fixed.timetuple())


@pytest.mark.parametrize('new_instructions', [False, True])
def test_identical_byte_cell_recovery_restores_committed_source_links(
    tmp_path, monkeypatch, frozen_excel, new_instructions
):
    store, job, deps, path, region = source_case(tmp_path, [['00101', 10], ['00101', 20]])
    dataset = extract(deps, region, 2)
    first, second = [row['id'] for row in deps.state.get('rows', dataset['id'])]
    workbook = load_workbook(deps.root / path)
    workbook['Output']['A2'] = '00101'
    workbook.save(deps.root / path)
    workbook.close()
    assign_cells(deps.state, path, [{'sheet': 'Output', 'cell': 'A2', 'value': '00101', 'record': first}])
    before = (deps.root / path).read_bytes()
    monkeypatch.setenv('DI_OPERATION_WORKER', '1')

    def interrupted(*args):
        raise OSError('Interrupted before source links were saved')

    with monkeypatch.context() as crash:
        crash.setattr('app.tools.assign_cells', interrupted)
        result = run_tool(
            store,
            job['id'],
            'write_cells',
            {'path': path, 'changes': [{'sheet': 'Output', 'cell': 'A2', 'value': '00101', 'record': second}]},
        )
    assert 'Interrupted' in result['error']
    operation, journal = next((key, item) for key, item in deps.state.all('operation') if item['status'] == 'prepared')
    assert journal['before_sha256'] == journal['sha256']
    assert (deps.root / path).read_bytes() == before
    if new_instructions:
        deps.state.instruction('Revisit all records')
        deps.state.apply_instructions([1], 'Changed period', reset_all=True)
    assert recover_operation(deps.state, operation)['status'] == 'completed'
    assignment = deps.state.get('assignment', identity(path, 'Output', 2))
    assert assignment['sources'] == [second]
    assert assignment['provenance']['A']['instruction_revision'] == 0
    assert recover_operation(deps.state, operation)['status'] == 'completed'
    assert (deps.root / path).read_bytes() == before


def test_uncommitted_cell_recovery_preserves_prior_source_links(tmp_path):
    _, _, deps, path, region = source_case(tmp_path, [['00101', 10]])
    dataset = extract(deps, region, 1)
    source = deps.state.get('rows', dataset['id'])[0]['id']
    assign_cells(deps.state, path, [{'sheet': 'Output', 'cell': 'A2', 'value': '00101', 'record': source}])
    before = deps.state.all('assignment')
    deps.state.put(
        'operation',
        'uncommitted',
        {
            'kind': 'cells',
            'status': 'prepared',
            'result': {'path': path},
            'before_sha256': hashlib.sha256((deps.root / path).read_bytes()).hexdigest(),
            'sha256': hashlib.sha256(b'never written').hexdigest(),
            'changes': [],
            'instruction_revision': 0,
        },
    )
    assert recover_operation(deps.state, 'uncommitted')['status'] == 'aborted'
    assert deps.state.all('assignment') == before
