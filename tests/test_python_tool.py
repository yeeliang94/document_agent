"""The opt-in run_python tool: sandbox escapes, staged commits, dataset accounting and a scripted run."""

import asyncio
import hashlib
import io
import json
import socket
import sys
from dataclasses import replace

import pytest
from openpyxl import load_workbook
from pydantic_ai.models.function import FunctionModel

from app import python_tool
from app.config import Settings
from app.datasets import Column, extract_dataset
from app.files import atomic_bytes
from app.inventory import inspect_source
from app.python_tool import run_python
from app.runtime import Runner
from app.store import Store
from app.tools import job_lock

from goal_helpers import messy_case

pytestmark = pytest.mark.skipif(sys.platform != 'darwin', reason='The prototype sandbox uses macOS sandbox-exec')

SOURCE = 'inputs/nested/October.xlsx'


@pytest.fixture(params=['both layers', 'OS sandbox only'])
def layers(request, monkeypatch):
    """Every escape must fail even with the in-process audit hook switched off."""
    if request.param == 'OS sandbox only':
        monkeypatch.setattr(python_tool, 'BOOT', python_tool.BOOT.replace('sys.addaudithook(hook)', 'pass'))
    return request.param


def run(deps, code, datasets=()):
    return asyncio.run(run_python(deps, code, list(datasets)))


def attempts(*snippets):
    """A script that tries each snippet and prints what happened, so effects can be checked from outside."""
    return '\n'.join(
        f'try:\n    {snippet}\n    print("allowed", {snippet!r})\nexcept Exception as error:\n'
        f'    print("blocked", {snippet!r}, type(error).__name__)'
        for snippet in snippets
    )


def payments_dataset(deps):
    state = deps.state
    inspect_source(state, SOURCE)
    region = next(key for key, r in state.all('region') if r['bounds'] == 'A14:C16')
    from app.progress import observe
    from app.regions import RegionRole, save_role

    evidence = observe(deps, 'inspect_regions', {'path': SOURCE, 'region': region}, {'id': region})[0]['key']
    save_role(state, RegionRole(region=region, role='records', rationale='Current payments', evidence=[evidence]))
    columns = [
        Column(name='id', column=1, kind='identifier'),
        Column(name='amount', column=2, kind='decimal'),
        Column(name='code', column=3, kind='text'),
    ]
    return extract_dataset(state, SOURCE, 'Payments', 14, 16, columns, [14], region_id=region)


def test_tool_is_registered_only_when_enabled(tmp_path):
    model = FunctionModel(lambda messages, info: None)
    tools = {}
    for enabled in (False, True):
        settings = Settings(data_dir=tmp_path, python_tool=enabled)
        agent = Runner(Store(tmp_path), settings, model_override=model).build_agent('test')
        tools[enabled] = set(agent._function_toolset.tools)
    assert 'run_python' not in tools[False]
    assert tools[True] - tools[False] == {'run_python'}


def test_network_is_unreachable(tmp_path, layers):
    _, _, deps = messy_case(tmp_path)
    listener = socket.create_server(('127.0.0.1', 0))
    port = listener.getsockname()[1]
    result = run(
        deps,
        'import socket, urllib.request\n'
        + attempts(
            f"socket.create_connection(('127.0.0.1', {port}), timeout=1)",
            "socket.getaddrinfo('example.com', 443)",
            "urllib.request.urlopen('http://127.0.0.1:%d' % " + str(port) + ', timeout=1)',
            'socket.socket(socket.AF_UNIX).connect("/var/run/mDNSResponder")',
        ),
    )
    # A connection that got through would be waiting in the listen backlog.
    listener.settimeout(0.2)
    with pytest.raises(TimeoutError):
        listener.accept()
    listener.close()
    assert result['ok'], result
    assert 'allowed' not in result['stdout'], result['stdout']
    assert result['stdout'].count('blocked') == 4


def test_paths_outside_the_job_cannot_be_read_or_written(tmp_path, layers):
    _, _, deps = messy_case(tmp_path / 'store')
    secret = tmp_path / 'secret.txt'
    secret.write_text('payroll of another company')
    elsewhere = tmp_path / 'elsewhere'
    elsewhere.mkdir()
    other_job = deps.store.create('payroll', '2026-10', 'Another job', '')
    other = deps.store.path(other_job['id']) / 'working'
    result = run(
        deps,
        'import os, shutil\n'
        + attempts(
            f'print(open({str(secret)!r}).read())',
            f'open({str(elsewhere / "planted.txt")!r}, "w").write("x")',
            f'open({str(other / "planted.txt")!r}, "w").write("x")',
            f'os.listdir({str(tmp_path)!r})',
            f'os.symlink({str(secret)!r}, SCRATCH + "/link"); print(open(SCRATCH + "/link").read())',
            'print(open(SCRATCH + "/../../../../../../secret.txt").read())',
            "print(open(os.path.expanduser('~/.ssh/id_rsa')).read())",
            f'shutil.copy({str(secret)!r}, SCRATCH + "/copy.txt")',
        ),
    )
    assert result['ok'], result
    assert 'payroll of another company' not in result['stdout']
    assert result['stdout'].count('blocked') == 8, result['stdout']
    assert not (elsewhere / 'planted.txt').exists()
    assert not (other / 'planted.txt').exists()
    assert not (deps.root / 'working/python/scratch/copy.txt').exists()


def test_inputs_and_harness_state_are_read_only(tmp_path, layers):
    _, _, deps = messy_case(tmp_path)
    before = {
        path: hashlib.sha256((deps.root / path).read_bytes()).hexdigest()
        for path in (SOURCE, 'template/destination.xlsx', 'working/state.sqlite')
    }
    result = run(
        deps,
        'import os\n'
        + attempts(
            "open(INPUTS + '/nested/October.xlsx', 'ab').write(b'x')",
            "open(INPUTS + '/nested/October.xlsx', 'wb').write(b'')",
            "os.remove(INPUTS + '/nested/October.xlsx')",
            "os.rename(INPUTS + '/nested/October.xlsx', SCRATCH + '/moved.xlsx')",
            "os.link(INPUTS + '/nested/October.xlsx', SCRATCH + '/hard.xlsx')",
            "open(INPUTS + '/nested/new.xlsx', 'w').write('x')",
            "open(TEMPLATE + '/destination.xlsx', 'ab').write(b'x')",
            "open(WORKING + '/state.sqlite', 'ab').write(b'x')",
            "open(WORKING + '/validation.json', 'w').write('{\"passed\": true}')",
            "os.chmod(INPUTS + '/nested/October.xlsx', 0o777)",
        ),
    )
    assert result['ok'], result
    assert 'allowed' not in result['stdout'], result['stdout']
    assert not (deps.root / 'inputs/nested/new.xlsx').exists()
    assert not (deps.root / 'working/validation.json').exists()
    after = {path: hashlib.sha256((deps.root / path).read_bytes()).hexdigest() for path in before}
    assert after == before
    # The intake hash check still accepts the untouched source.
    assert inspect_source(deps.state, SOURCE)


def test_processes_cannot_be_started(tmp_path, layers):
    _, _, deps = messy_case(tmp_path)
    marker = deps.root / 'working/python/scratch/spawned'
    result = run(
        deps,
        'import os, subprocess\n'
        + attempts(
            "os.system('touch ' + SCRATCH + '/spawned')",
            "subprocess.run(['/usr/bin/touch', SCRATCH + '/spawned'])",
            'os.fork()',
            "os.posix_spawn('/usr/bin/touch', ['touch', SCRATCH + '/spawned'], {})",
            "os.execv('/usr/bin/touch', ['touch', SCRATCH + '/spawned'])",
            "__import__('ctypes').CDLL(None)",
        ),
    )
    assert not marker.exists()
    if layers == 'both layers':
        assert result['ok'] and 'allowed' not in result['stdout'], result


def test_outputs_are_staged_and_committed_only_on_success(tmp_path):
    _, _, deps = messy_case(tmp_path)
    output = deps.root / 'outputs/output.xlsx'
    atomic_bytes(output, (deps.root / 'template/destination.xlsx').read_bytes())
    original = output.read_bytes()
    edit = (
        "import openpyxl\nwb = openpyxl.load_workbook(OUTPUTS + '/output.xlsx')\n"
        "wb['Output']['A2'] = 'EDITED'\nwb.save(OUTPUTS + '/output.xlsx')\n"
    )
    failed = run(deps, edit + "raise ValueError('changed my mind')")
    assert not failed['ok'] and not failed['saved']
    assert 'changed my mind' in failed['stderr'] and '<run_python>' in failed['stderr']
    assert output.read_bytes() == original
    new_file = run(deps, "open(OUTPUTS + '/extra.csv', 'w').write('a')")
    assert not new_file['saved'] and 'create_output' in new_file['error']
    assert not (deps.root / 'outputs/extra.csv').exists()
    saved = run(deps, edit)
    assert saved['ok'] and saved['outputs_changed'] == ['outputs/output.xlsx'], saved
    assert load_workbook(output)['Output']['A2'].value == 'EDITED'
    assert load_workbook(output)['Output']['D2'].value == '=SUM(B2:B3)'
    assert json.loads((deps.root / saved['audit']).read_text())['code'] == edit
    assert not list((deps.root / 'working/python/runs').iterdir())


def test_deadline_kills_the_script_and_saves_nothing(tmp_path):
    _, _, deps = messy_case(tmp_path)
    deps.settings = replace(deps.settings, tool_timeout=1)
    output = deps.root / 'outputs/output.xlsx'
    atomic_bytes(output, (deps.root / 'template/destination.xlsx').read_bytes())
    original = output.read_bytes()
    with pytest.raises(TimeoutError, match='nothing was saved'):
        run(deps, "open(OUTPUTS + '/output.xlsx', 'wb').write(b'partial')\nwhile True:\n    pass")
    assert output.read_bytes() == original
    assert deps.current()['status'] != 'pausing'
    assert job_lock(deps.job_id).acquire(blocking=False)
    job_lock(deps.job_id).release()


def test_output_returned_to_the_model_is_bounded(tmp_path):
    _, _, deps = messy_case(tmp_path)
    limit = deps.settings.tool_result_bytes // 4
    result = run(deps, "import sys\nprint('x' * 5_000_000)\nsys.stderr.write('e' * 100_000 + 'END')")
    assert result['ok']
    assert len(result['stdout']) == limit and len(result['stderr']) == limit
    assert result['stderr'].endswith('END')
    assert '5000001' in result['stdout_truncated']


def test_saved_datasets_keep_source_accounting(tmp_path):
    _, _, deps = messy_case(tmp_path)
    dataset = payments_dataset(deps)
    rows = deps.state.get('rows', dataset['id'])
    load = f'rows = DATASETS[{dataset["id"]!r}]["rows"]\n'
    dropped = run(deps, load + 'save_dataset(rows[:1])', [dataset['id']])
    assert not dropped['saved'] and 'neither in the result nor in excluded' in dropped['error']
    invented = run(
        deps, load + "save_dataset([{'values': {'id': 'X'}, 'sources': ['made-up']}] + rows)", [dataset['id']]
    )
    assert not invented['saved'] and 'sources' in invented['error']
    excluded = run(
        deps,
        load + "save_dataset(rows[:1], excluded={s: 'Paid separately' for s in rows[1]['sources']})",
        [dataset['id']],
    )
    assert excluded['saved'] and excluded['dataset']['rows'] == 1, excluded
    derived = excluded['dataset']['id']
    assert deps.state.get('exclusions', derived) == [[source, 'Paid separately'] for source in rows[1]['sources']]
    assert deps.state.get('mapping', derived)['operations'][0]['operation'] == 'python'


def test_scripted_job_uses_python_between_extraction_and_population(tmp_path):
    from fastapi.testclient import TestClient
    from pydantic_ai.messages import ModelRequest, ModelResponse, ToolCallPart, ToolReturnPart

    from app.main import create_app

    from helpers import LOCAL, wait_for_job

    _, _, fixture = messy_case(tmp_path / 'fixtures', top=37)
    returns, calls = {}, []
    code = (
        'from decimal import Decimal\n'
        'rows = []\n'
        'for row in DATASETS[dataset]["rows"]:\n'
        '    values = dict(row["values"], amount=str(Decimal(row["values"]["amount"]) * 1))\n'
        '    rows.append(dict(row, values=values))\n'
        'save_dataset(rows)\n'
        'print("rows", len(rows))\n'
    )

    def respond(messages, info):
        for message in messages:
            if isinstance(message, ModelRequest):
                for part in message.parts:
                    if isinstance(part, ToolReturnPart) and isinstance(part.content, dict):
                        returns[part.tool_name] = part.content
        number = len(calls)
        if number == 0:
            actions = [
                ('inspect_source', {'path': 'inputs/October.xlsx'}),
                ('inspect_source', {'path': 'template/destination.xlsx'}),
                ('read_sheet', {'path': 'template/destination.xlsx', 'sheet': 'Output'}),
            ]
        elif number == 1:
            actions = [
                ('inspect_regions', {'path': 'inputs/October.xlsx'}),
                ('create_output', {'filename': 'output.xlsx'}),
            ]
        elif number == 2:
            actions = [
                (
                    'save_interpretation',
                    {
                        'decision': {
                            'region': region['id'],
                            'role': 'records'
                            if region['bounds'] == 'A37:C39'
                            else 'outside_scope'
                            if region['bounds'].startswith('A47:')
                            else 'reference',
                            'rationale': 'Hand-authored synthetic fixture interpretation from inspected content',
                            'evidence': [returns['inspect_regions']['evidence_id']],
                        }
                    },
                )
                for region in returns['inspect_regions']['items']
            ]
            actions.append(
                (
                    'save_output_requirements',
                    {
                        'specification': {
                            'objective': 'Prepare the two current payments',
                            'path': 'outputs/output.xlsx',
                            'areas': [
                                {
                                    'id': 'data',
                                    'sheet': 'Output',
                                    'bounds': 'A2:C3',
                                    'role': 'required',
                                    'fields': [
                                        {'column': 1, 'meaning': 'ID', 'kind': 'identifier'},
                                        {'column': 2, 'meaning': 'Amount', 'kind': 'decimal', 'expected_total': '30'},
                                    ],
                                    'min_records': 2,
                                    'evidence': [returns['read_sheet']['evidence_id']],
                                    'rationale': 'Current payments',
                                },
                                {
                                    'id': 'headers',
                                    'sheet': 'Output',
                                    'bounds': 'A1:C1',
                                    'role': 'reference',
                                    'evidence': [returns['read_sheet']['evidence_id']],
                                    'rationale': 'Field labels',
                                },
                                {
                                    'id': 'formula',
                                    'sheet': 'Output',
                                    'bounds': 'D2:D2',
                                    'role': 'preserved',
                                    'evidence': [returns['read_sheet']['evidence_id']],
                                    'rationale': 'Template calculation',
                                },
                            ],
                        }
                    },
                )
            )
        elif number == 3:
            region = next(r['id'] for r in returns['inspect_regions']['items'] if r['bounds'] == 'A37:C39')
            columns = [
                {'name': 'id', 'column': 1, 'kind': 'identifier'},
                {'name': 'amount', 'column': 2, 'kind': 'decimal'},
                {'name': 'code', 'column': 3, 'kind': 'text'},
            ]
            actions = [
                (
                    'extract_region',
                    {'region': region, 'start_row': 37, 'end_row': 39, 'columns': columns, 'header_rows': [37]},
                )
            ]
        elif number == 4:
            dataset = returns['extract_region']['id']
            actions = [('run_python', {'code': f'dataset = {dataset!r}\n' + code, 'datasets': [dataset]})]
        elif number == 5:
            assert returns['run_python']['saved'], returns['run_python']
            assert returns['run_python']['stdout'] == 'rows 2\n'
            actions = [
                (
                    'populate_template',
                    {
                        'dataset': returns['run_python']['dataset']['id'],
                        'path': 'outputs/output.xlsx',
                        'destination': {
                            'sheet': 'Output',
                            'start_row': 2,
                            'end_row': 3,
                            'columns': [
                                {'field': 'id', 'column': 1},
                                {'field': 'amount', 'column': 2},
                                {'field': 'code', 'column': 3},
                            ],
                            'required': ['id', 'amount', 'code'],
                        },
                        'operation_id': 'payments',
                    },
                )
            ]
        elif number == 6:
            actions = [('validate_output', {'path': 'outputs/output.xlsx'})]
        elif number == 7:
            actions = [('validate_output', {'path': 'outputs/output.xlsx'})]
        elif number == 8:
            assert returns['validate_output']['passed'], returns['validate_output']
            actions = [
                (
                    'update_work_plan',
                    {
                        'plan': {
                            'objective': 'Prepare October payments',
                            'focus': 'Validated and ready to deliver',
                            'items': [
                                {
                                    'id': 'payments',
                                    'objective': 'Prepare and validate current payments',
                                    'status': 'done',
                                    'evidence': [returns['run_python']['dataset']['id']],
                                }
                            ],
                            'next_actions': ['Deliver the validated output'],
                        }
                    },
                )
            ]
        elif number == 9:
            actions = [('final_result', {'summary': 'Prepared October payments.', 'status': 'complete'})]
        else:
            raise AssertionError('The run should have ended with its final result')
        calls.append(actions)
        return ModelResponse(
            parts=[ToolCallPart(name, args, tool_call_id=f'{number}-{i}') for i, (name, args) in enumerate(actions)]
        )

    settings = Settings(data_dir=tmp_path / 'app', api_key='synthetic', python_tool=True)
    app = create_app(settings, model_override=FunctionModel(respond))
    with TestClient(app, base_url=LOCAL) as client:
        uploaded = client.post(
            '/api/jobs',
            data={'module': 'payroll', 'period': '2026-10', 'instructions': 'Prepare October payments'},
            files=[
                ('template', ('destination.xlsx', (fixture.root / 'template/destination.xlsx').read_bytes())),
                ('files', ('October.xlsx', (fixture.root / 'inputs/nested/October.xlsx').read_bytes())),
            ],
        )
        assert uploaded.status_code == 201
        job = uploaded.json()
        assert client.post(f'/api/jobs/{job["id"]}/start').status_code == 200
        completed = wait_for_job(client, job['id'], statuses=('completed', 'partial', 'failed', 'paused'))
        assert completed['status'] == 'completed', completed
        workbook = load_workbook(io.BytesIO(client.get(f'/api/jobs/{job["id"]}/artifacts/output.xlsx').content))
        assert [workbook['Output'][c].value for c in ('A2', 'B2', 'A3', 'B3', 'D2')] == [
            '00101',
            10,
            '00102',
            20,
            '=SUM(B2:B3)',
        ]


def test_application_folder_and_other_jobs_cannot_be_read(tmp_path, layers):
    from app.config import ROOT

    _, _, deps = messy_case(tmp_path)
    other_job = deps.store.create('payroll', '2026-10', 'Another job', '')
    other = deps.store.path(other_job['id']) / 'working/notes.txt'
    other.write_text('another job')
    reads = python_tool.python_reads(deps.settings.data_dir)
    assert not any(str(ROOT).startswith(p) or str(tmp_path.resolve()).startswith(p) for p in reads)
    result = run(
        deps,
        attempts(f'print(open({str(ROOT / "requirements.txt")!r}).read())', f'print(open({str(other)!r}).read())'),
    )
    assert result['ok'], result
    assert 'another job' not in result['stdout'] and 'openpyxl' not in result['stdout']
    assert result['stdout'].count('blocked') == 2, result['stdout']


def test_links_in_staged_outputs_are_refused(tmp_path):
    _, _, deps = messy_case(tmp_path / 'store')
    secret = tmp_path / 'secret.txt'
    secret.write_text('payroll of another company')
    output = deps.root / 'outputs/notes.txt'
    atomic_bytes(output, b'draft')
    # With the audit hook disabled from inside the script, only the commit stands between a link and its target.
    unhook = 'import __main__\n__main__.READ = __main__.WRITE = ("",)\nimport os\n'
    for attack in (
        f'os.remove(OUTPUTS + "/notes.txt"); os.symlink({str(secret)!r}, OUTPUTS + "/notes.txt")',
        f'os.rename(OUTPUTS, OUTPUTS + "-old"); os.symlink({str(tmp_path)!r}, OUTPUTS)',
    ):
        result = run(deps, unhook + attack)
        assert result['returncode'] == 0 and not result['saved'], result
        assert output.read_bytes() == b'draft'
    assert 'payroll' not in json.dumps(deps.store.events(deps.job_id))


def test_python_edits_outside_input_areas_are_refused(tmp_path):
    from goal_helpers import establish_requirements

    _, _, deps = messy_case(tmp_path)
    path = establish_requirements(deps)
    output = deps.root / path
    original = output.read_bytes()
    edit = (
        "import openpyxl\nwb = openpyxl.load_workbook(OUTPUTS + '/output.xlsx')\n{}\nwb.save(OUTPUTS + '/output.xlsx')"
    )
    for change in ("wb['Output']['A1'] = 'Renamed'", "wb['Lookup']['A1'] = 'Invented'"):
        refused = run(deps, edit.format(change))
        assert not refused['saved'] and 'outside its input areas' in refused['error'], refused
        assert output.read_bytes() == original
    allowed = run(deps, edit.format("wb['Output']['A2'] = '00101'"))
    assert allowed['saved'], allowed


def test_cancelled_commit_finishes_before_the_lock_is_released(tmp_path, monkeypatch):
    import threading
    import time

    _, _, deps = messy_case(tmp_path)
    started, finished = threading.Event(), []

    def slow_commit(*args):
        started.set()
        time.sleep(0.3)
        finished.append(True)
        return {}

    monkeypatch.setattr(python_tool, '_commit', slow_commit)

    async def cancel_during_commit():
        task = asyncio.create_task(run_python(deps, 'pass', []))
        while not started.is_set():
            await asyncio.sleep(0.01)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        return list(finished)

    assert asyncio.run(cancel_during_commit()) == [True]


def test_derived_rows_keep_source_issues_unless_resolved(tmp_path):
    _, _, deps = messy_case(tmp_path)
    dataset = payments_dataset(deps)
    rows = deps.state.get('rows', dataset['id'])
    rows[0]['issues'] = ['Formula cache missing']
    deps.state.put('rows', dataset['id'], rows)
    rebuild = (
        f"rows = [{{'values': r['values'], 'sources': r['sources']}} for r in DATASETS[{dataset['id']!r}]['rows']]\n"
    )
    kept = run(deps, rebuild + 'save_dataset(rows)', [dataset['id']])
    assert kept['saved'] and kept['dataset']['exceptions'] == 1, kept
    assert deps.state.get('rows', kept['dataset']['id'])[0]['issues'] == ['Formula cache missing']
    invented = run(
        deps, rebuild + "rows[0]['resolved'] = {'Something else': 'ok'}\nsave_dataset(rows)", [dataset['id']]
    )
    assert not invented['saved'] and 'resolved' in invented['error']
    resolved = run(
        deps,
        rebuild + "rows[0]['resolved'] = {'Formula cache missing': 'Recomputed from B14'}\nsave_dataset(rows)",
        [dataset['id']],
    )
    assert resolved['saved'] and resolved['dataset']['exceptions'] == 0, resolved
    saved = deps.state.get('rows', resolved['dataset']['id'])[0]
    assert saved['evidence']['resolved_issues'] == {'Formula cache missing': 'Recomputed from B14'}
