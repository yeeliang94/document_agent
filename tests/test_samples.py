from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image

from scripts.create_sample import create_sample


def test_synthetic_packet_is_nested_and_template_has_preservable_features(tmp_path):
    create_sample(tmp_path)
    with ZipFile(tmp_path / 'payroll_packet.zip') as archive:
        names = archive.namelist()
        assert 'October 2026/HR/Reports/employee_report.xlsx' in names
        assert 'October 2026/Supporting documents/salary_confirmation.png' in names
    with Image.open(tmp_path / 'payroll_packet/October 2026/Supporting documents/salary_confirmation.png') as image:
        assert image.width == 1200
    template = load_workbook(tmp_path / 'output_template.xlsx')
    assert template['Lookup'].sheet_state == 'hidden'
    assert template['Employees']['E2'].value == '=C2*12'
    assert template['Employees']['E3'].value == '=C3*12'
    assert template['Employees']['A2'].value is None
    source = load_workbook(tmp_path / 'payroll_packet/October 2026/HR/Reports/employee_report.xlsx')
    assert list(source.active.values)[1:] == [
        ('EMP001', 'Alex Example', 4200, '2026-10-01'),
        ('EMP002', 'Sam Sample', 5800, '2026-10-01'),
    ]
